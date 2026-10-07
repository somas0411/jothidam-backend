"""
excel_generator.py — HoroscopeGen Excel workbook.

Sheets: Summary, Vedic, Divisional Charts, KP, ALP, Dasa, Notes. When the
AI block from enrich.py is passed in, the flat "AI_" sheets follow Notes:
plain tables for AI readers (values only, no merged cells), in landscape.

Charts are drawn with real cells (merged cells, borders and diagonal
borders), never pasted images, on a grid of narrow columns: every chart is
12 x 12 small cells, two charts side by side. Tables on the chart sheets use
merged spans of that grid; the Dasa sheet uses ordinary columns so it can be
sorted and filtered. Dates are real Excel dates. Every sheet is set up for
A4 portrait, one page wide.
"""
import io
from datetime import date, datetime, time
from functools import lru_cache

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

import charts
import i18n
from astro_engine import fmt_dms, fmt_sign_dms

# ── PALETTE (print-friendly: white paper, dark ink, one accent) ───────────────
INK = '2B1B0E'
MUTED = '7A6A58'
BAND = '5A2A0C'          # title bands
HEAD = '8A4B14'          # table headers
ACCENT = 'C8811A'
RULE = 'C9B79C'          # thin borders
ALT = 'FBF5EA'           # zebra rows
SOFT = 'F4E7CF'          # label cells, chart centre
LAGNA = 'FFE2A8'         # lagna cell
CURRENT = 'FFF1B8'       # current period
WHITE = 'FFFFFF'

FONT = 'Calibri'
GRID_W = 4.0             # width of one narrow grid column
CELL_H = 24.0            # height of one chart row (roughly square cells)
ROW_H = 17.0
PURPOSE_H = 16.0          # height of the purpose line under a divisional chart's title
COL0 = 2                 # first grid column (B); column A is a margin
GRID_N = 25              # 12 + 1 gap + 12
RIGHT0 = COL0 + 13       # first column of the right-hand chart

DATE_FMT = 'dd-mm-yyyy'
TIME_FMT = 'hh:mm:ss'

_thin = Side(style='thin', color=RULE)
_med = Side(style='medium', color=INK)
_diag = Side(style='thin', color=INK)


@lru_cache(maxsize=None)
def _fill(color):
    return PatternFill('solid', fgColor=color)


@lru_cache(maxsize=None)
def _font(size=9, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


@lru_cache(maxsize=None)
def _align(h='left', v='center', wrap=False, shrink=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, shrink_to_fit=shrink, indent=indent)


_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _sheet_name(key, lang):
    name = i18n.LABELS['en'][key] if lang == 'bi' else i18n.L(key, lang)
    for ch in '[]:*?/\\':
        name = name.replace(ch, ' ')
    return name[:31]


def _d(dt):
    """Naive datetime -> date for a real Excel date cell."""
    return dt.date() if isinstance(dt, datetime) else dt


def _t(dt):
    return time(dt.hour, dt.minute, dt.second) if isinstance(dt, datetime) else dt


class Grid:
    """Thin helper around a worksheet laid out on the narrow column grid."""

    def __init__(self, wb, title, lang, first=False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title
        self.lang = lang
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 1.5
        for c in range(COL0, COL0 + GRID_N):
            ws.column_dimensions[get_column_letter(c)].width = GRID_W
        self.wrap = lang == 'bi'

    # -- primitives -----------------------------------------------------------
    def put(self, r, c, value, span=1, rows=1, size=9, bold=False, color=INK, fill=None,
            h='left', v='center', wrap=None, border=True, fmt=None, italic=False, shrink=False):
        ws = self.ws
        if span > 1 or rows > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r + rows - 1, end_column=c + span - 1)
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = _font(size, bold, color, italic)
        do_wrap = self.wrap if wrap is None else wrap
        cell.alignment = _align(h, v, do_wrap, shrink and not do_wrap, 1 if h == 'left' else 0)
        if fmt:
            cell.number_format = fmt
        for rr in range(r, r + rows):
            for cc in range(c, c + span):
                x = ws.cell(row=rr, column=cc)
                if fill:
                    x.fill = _fill(fill)
                if border:
                    x.border = _BOX
        return cell

    def height(self, r, h):
        self.ws.row_dimensions[r].height = h

    def title(self, r, text, sub=None):
        self.height(r, 30)
        self.put(r, COL0, text, span=GRID_N, size=15, bold=True, color=WHITE, fill=BAND, h='center',
                 wrap=False, border=False)
        if sub:
            self.height(r + 1, 18)
            self.put(r + 1, COL0, sub, span=GRID_N, size=9, color=MUTED, fill=SOFT, h='center',
                     wrap=False, border=False, italic=True)
            return r + 3
        return r + 2

    def page_break(self, r):
        """Start a new printed page at row r."""
        self.ws.row_breaks.append(Break(id=r - 1))

    def section(self, r, text, c=COL0, span=GRID_N, keep_height=False):
        if not keep_height:
            self.height(r, 20)
        self.put(r, c, text, span=span, size=10.5, bold=True, color=WHITE, fill=HEAD, wrap=False, border=False)
        return r + 1

    def pairs(self, r, c, items, label_span=5, value_span=7, keep_height=False):
        """
        Label / value rows inside a 12-column half. items: (label, value[, fmt]).
        keep_height: the rows are shared with a chart, so leave their height alone.
        """
        for it in items:
            label, value = it[0], it[1]
            fmt = it[2] if len(it) > 2 else None
            if not keep_height:
                self.height(r, ROW_H + (9 if self.wrap else 0))
            self.put(r, c, label, span=label_span, bold=True, color=MUTED, fill=SOFT)
            self.put(r, c + label_span, value, span=value_span, fmt=fmt, shrink=True)
            r += 1
        return r

    def table(self, r, cols, rows, current=None, c=COL0, row_h=None, tall=None):
        """
        cols: [(header, span, align)], rows: list of value lists.
        A value may be (value, number_format). current: set of row indexes to highlight.
        """
        self.height(r, 26 if self.lang == 'en' else 38)
        cc = c
        for head, span, _ in cols:
            self.put(r, cc, head, span=span, size=7 if span == 1 else 8.5, bold=True, color=WHITE, fill=HEAD,
                     h='center', wrap=span > 1)
            cc += span
        r += 1
        for i, row in enumerate(rows):
            is_cur = bool(current) and i in current
            fill = CURRENT if is_cur else (ALT if i % 2 else WHITE)
            is_tall = bool(tall) and i in tall
            self.height(r, 28 if is_tall else (row_h or (ROW_H + (10 if self.wrap else 0))))
            cc = c
            for (head, span, align), val in zip(cols, row):
                fmt = None
                if isinstance(val, tuple):
                    val, fmt = val
                self.put(r, cc, val, span=span, fill=fill, h=align, bold=is_cur, fmt=fmt,
                         wrap=True if is_tall else None, shrink=True)
                cc += span
            r += 1
        return r

    # -- charts ---------------------------------------------------------------
    def chart(self, r, c, spec, style='south', purpose=None):
        """
        Draw one 12x12-cell chart with its title row, and a purpose line under
        the title when given. Returns the next free row.
        """
        self.height(r, 20)
        self.put(r, c, spec['title'], span=12, size=10, bold=True, color=WHITE, fill=HEAD, h='center',
                 wrap=False, border=False)
        top = r + 1
        if purpose is not None:
            self.height(top, PURPOSE_H if self.lang != 'bi' else PURPOSE_H * 2 - 4)
            self.put(top, c, purpose, span=12, size=8, color=MUTED, fill=SOFT, h='center', italic=True,
                     wrap=self.lang == 'bi', border=False, shrink=True)
            top += 1
        for rr in range(top, top + 12):
            self.height(rr, CELL_H)
        if style == 'north':
            self._north(top, c, spec)
        else:
            self._south(top, c, spec)
        self._outline(top, c)
        return top + 13

    def _outline(self, top, c):
        ws = self.ws
        for i in range(12):
            for (rr, cc, side) in ((top, c + i, 'top'), (top + 11, c + i, 'bottom'),
                                   (top + i, c, 'left'), (top + i, c + 11, 'right')):
                cell = ws.cell(row=rr, column=cc)
                b = cell.border
                kw = dict(left=b.left, right=b.right, top=b.top, bottom=b.bottom,
                          diagonal=b.diagonal, diagonalUp=b.diagonalUp, diagonalDown=b.diagonalDown)
                kw[side] = _med
                cell.border = Border(**kw)

    def _south(self, top, c, spec):
        ws = self.ws
        for cell in spec['signs']:
            gr, gc = charts.SOUTH_POS[cell['sign']]
            r0, c0 = top + gr * 3, c + gc * 3
            fill = LAGNA if cell['lagna'] else WHITE
            tag = cell['tag']
            self.put(r0, c0, f"{cell['label']}   {tag}".rstrip(), span=3, size=7, color=MUTED, fill=fill,
                     wrap=False, border=False, bold=cell['lagna'])
            self.put(r0 + 1, c0, ' '.join(cell['planets']), span=3, rows=2, size=9, bold=True, fill=fill,
                     h='center', v='top', wrap=True, border=False)
            ink = Side(style='thin', color=INK)
            for i in range(3):
                for (rr, cc, side) in ((r0, c0 + i, 'top'), (r0 + 2, c0 + i, 'bottom'),
                                       (r0 + i, c0, 'left'), (r0 + i, c0 + 2, 'right')):
                    x = ws.cell(row=rr, column=cc)
                    b = x.border
                    kw = dict(left=b.left, right=b.right, top=b.top, bottom=b.bottom)
                    kw[side] = ink
                    x.border = Border(**kw)
        self.put(top + 3, c + 3, '\n'.join(spec['center']), span=6, rows=6, size=10, bold=True,
                 fill=SOFT, h='center', v='center', wrap=True, border=False)

    # North Indian layout on the 12x12 small-cell grid:
    # house -> (planet rect, sign-number rect), each rect = (r0, c0, r1, c1)
    NORTH_CELLS = {
        1:  ((2, 4, 3, 7),   (4, 5, 4, 6)),
        2:  ((0, 1, 0, 4),   (1, 2, 1, 3)),
        3:  ((1, 0, 4, 0),   (2, 1, 3, 1)),
        4:  ((4, 2, 7, 3),   (5, 4, 6, 4)),
        5:  ((7, 0, 10, 0),  (8, 1, 9, 1)),
        6:  ((11, 1, 11, 4), (10, 2, 10, 3)),
        7:  ((8, 4, 9, 7),   (7, 5, 7, 6)),
        8:  ((11, 7, 11, 10), (10, 8, 10, 9)),
        9:  ((7, 11, 10, 11), (8, 10, 9, 10)),
        10: ((4, 8, 7, 9),   (5, 7, 6, 7)),
        11: ((1, 11, 4, 11), (2, 10, 3, 10)),
        12: ((0, 7, 0, 10),  (1, 8, 1, 9)),
    }

    def _north(self, top, c, spec):
        ws = self.ws
        # Diagonals: each 3x3 block carries one diagonal, drawn cell by cell.
        for br in range(4):
            for bc in range(4):
                down = (br + bc) % 2 == 0
                for i in range(3):
                    rr = top + br * 3 + i
                    cc = c + bc * 3 + (i if down else 2 - i)
                    ws.cell(row=rr, column=cc).border = Border(diagonal=_diag, diagonalDown=down,
                                                               diagonalUp=not down)
        for h in spec['houses']:
            (pr0, pc0, pr1, pc1), (nr0, nc0, nr1, nc1) = self.NORTH_CELLS[h['house']]
            tall = pc0 == pc1          # side triangles: stack the planets
            text = ('\n' if tall else ' ').join(h['planets'])
            crowded = len(h['planets']) > 4          # e.g. Hora, where two signs hold every point
            self.put(top + pr0, c + pc0, text, span=pc1 - pc0 + 1, rows=pr1 - pr0 + 1,
                     size=(7 if crowded else 8) if tall else (7.5 if crowded else 9),
                     bold=True, h='center', v='center', wrap=True, border=False)
            self.put(top + nr0, c + nc0, h['sign_num'], span=nc1 - nc0 + 1, rows=nr1 - nr0 + 1,
                     size=7.5, color=ACCENT if h['house'] == 1 else MUTED, bold=h['house'] == 1,
                     h='center', v='center', wrap=False, border=False)


# ── PAGE SETUP ────────────────────────────────────────────────────────────────

PRINT_W = 525.0     # usable A4 width in points with the margins below
PRINT_H = 745.0     # usable A4 height in points


def _page_setup(ws, name, last_col, title_rows=None, one_page=False):
    """
    A4 portrait, one page wide. The scale is computed from the column widths
    rather than using fit-to-page, because Excel ignores manual page breaks
    in fit-to-page mode. one_page also fits the sheet to a single page tall.
    """
    width_px = 0
    for c in range(2, last_col + 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 8.43
        width_px += int(w * 7 + 5)
    scale = PRINT_W / (width_px * 0.75) * 100
    if one_page:
        height = sum((ws.row_dimensions[r].height or 15) for r in range(1, ws.max_row + 1))
        scale = min(scale, PRINT_H / height * 100)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.scale = max(40, min(100, int(scale)))
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    ws.oddFooter.left.text = 'horoscopegen.in'
    ws.oddFooter.center.text = '&P / &N'
    ws.oddFooter.right.text = name
    for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
        part.size = 8
    ws.print_area = f'B1:{get_column_letter(last_col)}{ws.max_row}'
    if title_rows:
        ws.print_title_rows = title_rows


# ── VALUE HELPERS ─────────────────────────────────────────────────────────────

def _state(p, lang):
    bits = []
    if p.get('retro'):
        bits.append(i18n.L('retro', lang))
    if p.get('combust'):
        bits.append(i18n.L('combust', lang))
    if p.get('dignity'):
        bits.append(i18n.word(p['dignity'], lang))
    return ', '.join(bits) if bits else '—'


def _status(row, now, lang):
    if row.get('current'):
        return i18n.L('current', lang)
    return i18n.L('completed', lang) if row['end'] <= now else i18n.L('upcoming', lang)


def _balance_text(bal, lang):
    return (f"{i18n.planet(bal['lord'], lang)}: {bal['y']} {i18n.Lv('y', lang)} "
            f"{bal['m']} {i18n.Lv('m', lang)} {bal['d']} {i18n.Lv('d', lang)}")


def _houses(nums):
    return ', '.join(str(n) for n in nums) if nums else '—'


def _points(v):
    """Rows of the planet table: lagna, the nine planets, then Mandi when available."""
    return v['planets'] + ([v['mandi']] if v.get('mandi') else [])


def _mandi_note(m, v):
    """Notes text explaining how Mandi was obtained for this chart."""
    base = ('Weekday ghati table for a 30-ghati day - Sunday 26, Monday 22, Tuesday 18, Wednesday 14, '
            'Thursday 10, Friday 6, Saturday 2 ghatis after sunrise; a night birth counts from sunset and '
            'uses the value of the 5th weekday from the birth weekday. The value is scaled to the actual '
            'length of the day or night, and Mandi is the lagna rising at that moment. Gulika is not shown.')
    md = v.get('mandi')
    if not md:
        return base + ' Mandi could not be computed for this chart: the Sun does not rise or set at this place on this date.'
    return (base + f" This chart: {'day' if md['is_day'] else 'night'} birth, {md['weekday']}, "
            f"{md['ghatis']} ghatis after {'sunrise' if md['is_day'] else 'sunset'} "
            f"({md['period_start'].strftime('%d-%m-%Y %H:%M:%S')}), rising at "
            f"{md['rise_time'].strftime('%d-%m-%Y %H:%M:%S')}.")


def _age(start, birth):
    return round((start - birth).days / 365.25, 2)


# ── SHEETS ────────────────────────────────────────────────────────────────────

def _summary(wb, res, specs, lang):
    g = Grid(wb, _sheet_name('summary', lang), lang, first=True)
    m, v, kp, alp = res['meta'], res['vedic'], res['kp'], res['alp']
    L = lambda k: i18n.L(k, lang)
    pc = v['panchangam']
    lagna = v['planets'][0]
    cur = v['dasa']['current']

    r = g.title(1, f"{L('brand')}  ·  {L('report_title')}",
                f"{m['name']}  ·  {L('subtitle')}  ·  horoscopegen.in")

    g.section(r, L('birth_details'))
    r += 1
    left = [(L('name'), m['name']),
            (L('dob'), m['local_dt'].date(), DATE_FMT),
            (L('tob'), _t(m['local_dt']), TIME_FMT),
            (L('pob'), m['pob']),
            (L('lat'), f"{m['lat']:.4f}°"),
            (L('lon'), f"{m['lon']:.4f}°")]
    right = [(L('tz'), f"{m['tz']} ({m['utc_offset_str']})"),
             (L('ayanamsha'), f"{m['ayanamsha_name']}  {v['ayanamsha_dms']}"),
             (L('kp_ayanamsha'), kp['ayanamsha_dms']),
             (L('nodes'), i18n.word('True' if m['node'] == 'true' else 'Mean', lang)),
             (L('sunrise'), _t(pc['sunrise']) if pc['sunrise'] else '—', TIME_FMT),
             (L('sunset'), _t(pc['sunset']) if pc['sunset'] else '—', TIME_FMT)]
    if m.get('gender'):
        left.append((L('gender'), i18n.word(m['gender'].title(), lang)))
    r = max(g.pairs(r, COL0, left), g.pairs(r, RIGHT0, right)) + 1

    g.section(r, L('core'))
    r += 1
    left = [(L('lagna'), f"{i18n.rasi(lagna['sign'], lang)}  {lagna['dms']}"),
            (L('janma_rasi'), i18n.rasi(v['moon_sign'], lang)),
            (L('nakshatra'), f"{i18n.nak(v['moon_nak'], lang)} - {i18n.Lv('pada', lang)} {v['moon_pada']}"),
            (L('nak_lord'), i18n.planet(v['dasa']['birth_star_lord'], lang)),
            (L('tithi'), f"{i18n.word(pc['paksha'], lang)} {i18n.tithi(pc['tithi'], lang)}"),
            (L('vara'), i18n.weekday(pc['vara'], lang))]
    md = v.get('mandi')
    if md:
        left.append((i18n.planet('Mandi', lang),
                     f"{i18n.rasi(md['sign'], lang)}  {md['dms']}  ({md['rise_time'].strftime('%H:%M:%S')})"))
    al = alp['lagna']
    right = [(L('dasa_balance'), _balance_text(v['dasa']['balance'], lang))]
    for key, lab in (('dasa', 'dasa'), ('bhukti', 'bhukti'), ('antaram', 'antara')):
        if lab in cur:
            right.append((i18n.Lj(['current', key], lang),
                          f"{i18n.planet(cur[lab]['lord'], lang)}  ({i18n.Lv('until', lang)} "
                          f"{cur[lab]['end'].strftime('%d-%m-%Y')})"))
    right.append((L('alp_lagna_now'), f"{i18n.rasi(al['sign'], lang)}  {al['dms']}"))
    right.append((L('as_of'), m['generated_at'].date(), DATE_FMT))
    r = max(g.pairs(r, COL0, left), g.pairs(r, RIGHT0, right)) + 1

    style = m['chart_style']
    g.chart(r, COL0, specs['d1'], style)
    r = g.chart(r, RIGHT0, specs['d9'], style)
    g.chart(r, COL0, specs['kp'], style)
    r = g.chart(r, RIGHT0, specs['alp'], style)

    g.put(r, COL0, L('disclaimer'), span=GRID_N, size=8, color=MUTED, h='center', italic=True, border=False,
          wrap=False)
    g.ws.freeze_panes = 'A3'
    _page_setup(g.ws, m['name'], COL0 + GRID_N - 1, one_page=True)


def _vedic(wb, res, specs, lang):
    g = Grid(wb, _sheet_name('vedic', lang), lang)
    m, v = res['meta'], res['vedic']
    L = lambda k: i18n.L(k, lang)
    pc = v['panchangam']
    style = m['chart_style']

    r = g.title(1, f"{L('vedic')}  ·  {m['name']}",
                f"{L('ayanamsha')}: {m['ayanamsha_name']} {v['ayanamsha_dms']}")
    g.chart(r, COL0, specs['d1'], style)
    r = g.chart(r, RIGHT0, specs['d9'], style)

    top = r
    g.chart(top, COL0, specs['bhava'], style)
    rr = g.section(top, L('panchangam'), c=RIGHT0, span=12, keep_height=True)
    rr = g.pairs(rr, RIGHT0, keep_height=True, items=[
        (L('tithi'), f"{i18n.word(pc['paksha'], lang)} {i18n.tithi(pc['tithi'], lang)}"),
        (L('vara'), i18n.weekday(pc['vara'], lang)),
        (L('nakshatra'), f"{i18n.nak(pc['nak'], lang)} - {i18n.Lv('pada', lang)} {pc['pada']}"),
        (L('yoga'), pc['yoga']),
        (L('karana'), pc['karana']),
        (L('sunrise'), _t(pc['sunrise']) if pc['sunrise'] else '—', TIME_FMT),
        (L('sunset'), _t(pc['sunset']) if pc['sunset'] else '—', TIME_FMT),
        (L('mandi_rise'), _t(v['mandi']['rise_time']) if v.get('mandi') else '—', TIME_FMT),
        (L('dasa_balance'), _balance_text(v['dasa']['balance'], lang)),
    ])
    r = top + 14

    g.page_break(r)
    r = g.section(r, L('planet_positions'))
    cols = [(L('planet'), 3, 'left'), (L('rasi'), 3, 'left'), (L('degree'), 3, 'center'),
            (L('nakshatra'), 4, 'left'), (L('pada'), 1, 'center'), (L('rasi_lord'), 2, 'left'),
            (L('star_lord'), 2, 'left'), (L('house'), 1, 'center'), (L('bhava'), 1, 'center'),
            (L('navamsa'), 2, 'left'), (f"{L('retro')} / {L('combust')} / {L('dignity')}", 3, 'left')]
    rows = [[i18n.planet(p['name'], lang), i18n.rasi(p['sign'], lang), p['dms'], i18n.nak(p['nak'], lang),
             p['pada'], i18n.planet(p['sign_lord'], lang), i18n.planet(p['star_lord'], lang),
             p['house'], p['bhava'], i18n.rasi(p['navamsa'], lang),
             _state(p, lang)] for p in _points(v)]
    tall = {i for i, p in enumerate(_points(v))
            if sum(bool(p.get(k)) for k in ('retro', 'combust', 'dignity')) > 1}
    r = g.table(r, cols, rows, tall=tall) + 1

    r = g.section(r, L('bhava_table'))
    by_bhava = {}
    for p in _points(v):
        by_bhava.setdefault(p['bhava'], []).append(i18n.planet(p['name'], lang))
    cols = [(L('bhava'), 2, 'center'), (L('bhava_start'), 6, 'left'), (L('bhava_mid'), 6, 'left'),
            (L('planet'), 11, 'left')]
    rows = [[b['house'], f"{i18n.rasi(b['start_sign'], lang)}  {b['start_dms']}",
             f"{i18n.rasi(b['madhya_sign'], lang)}  {b['madhya_dms']}",
             ', '.join(by_bhava.get(b['house'], [])) or '—'] for b in v['bhavas']]
    g.table(r, cols, rows)
    g.ws.freeze_panes = 'A3'
    _page_setup(g.ws, m['name'], COL0 + GRID_N - 1)


def _divisional(wb, res, lang):
    """The sixteen divisional charts: a table of signs, then the charts two to a row."""
    g = Grid(wb, _sheet_name('divisional', lang), lang)
    m = res['meta']
    L = lambda k: i18n.L(k, lang)
    vspecs = charts.build_varga_specs(res, lang)
    keys, rows = charts.varga_table(res, lang)

    r = g.title(1, f"{L('divisional')}  ·  {m['name']}", L('varga_time_note'))
    if lang == 'bi':                     # two languages: let the note take two lines
        g.height(2, 30)
        g.ws.cell(row=2, column=COL0).alignment = _align('center', 'center', True, False, 0)

    r = g.section(r, L('varga_table'))
    cols = [(L('planet'), 5, 'left')] + [(k, 1, 'center') for k in keys] + [(L('vargottama'), 4, 'center')]
    yes = i18n.word('Yes', lang)
    table_rows = [[row['label']] + row['signs'] + [yes if row['vargottama'] else '—'] for row in rows]
    hdr = r
    r = g.table(r, cols, table_rows)
    # sign names are short but the columns are one grid cell wide: use a small font
    for rr in range(hdr + 1, r):
        for cc in range(COL0 + 5, COL0 + 5 + len(keys)):
            cell = g.ws.cell(row=rr, column=cc)
            cell.font = _font(7.5)
            cell.alignment = _align('center', 'center', False, True, 0)
    g.height(r, 16)
    g.put(r, COL0, L('vargottama_note'), span=GRID_N, size=8, color=MUTED, italic=True, border=False, wrap=False)
    r += 2

    style = m['chart_style']
    for i in range(0, len(vspecs), 2):
        # First page: the table and one row of charts; after that two rows of charts per page.
        if i >= 2 and (i // 2) % 2 == 1:
            g.page_break(r)
        g.chart(r, COL0, vspecs[i], style, purpose=vspecs[i]['purpose'])
        r = g.chart(r, RIGHT0, vspecs[i + 1], style, purpose=vspecs[i + 1]['purpose'])
    g.ws.freeze_panes = 'A3'
    _page_setup(g.ws, m['name'], COL0 + GRID_N - 1)


def _kp(wb, res, specs, lang):
    g = Grid(wb, _sheet_name('kp', lang), lang)
    m, kp = res['meta'], res['kp']
    L = lambda k: i18n.L(k, lang)
    P = lambda n: i18n.planet(n, lang)
    now = m['generated_at']

    r = g.title(1, f"{L('kp')}  ·  {m['name']}",
                f"{L('kp_ayanamsha')}: {kp['ayanamsha_dms']}  ·  {L('house_system')}: {i18n.word(kp['house_system'], lang)}")
    top = r
    g.chart(top, COL0, specs['kp'], m['chart_style'])
    rp = kp['ruling_planets']
    rr = g.section(top, L('ruling_planets'), c=RIGHT0, span=12, keep_height=True)
    rr = g.pairs(rr, RIGHT0, [(L(k), P(rp[k])) for k in
                              ('day_lord', 'moon_sign_lord', 'moon_star_lord', 'moon_sub_lord',
                               'lagna_sign_lord', 'lagna_star_lord', 'lagna_sub_lord')],
                 label_span=6, value_span=6, keep_height=True)
    r = top + 14

    r = g.section(r, L('node_agents'))
    cols = [(L('node'), 4, 'left'), (L('rasi_lord'), 5, 'left'), (L('star_lord'), 5, 'left'),
            (L('conjoined'), 11, 'left')]
    r = g.table(r, cols, [[P(a['node']), P(a['sign_lord']), P(a['star_lord']),
                           i18n.planets(a['conjoined'], lang)] for a in kp['node_agents']]) + 1

    lord_cols = [(L('rasi_lord'), 3, 'left'), (L('star_lord'), 3, 'left'),
                 (L('sub_lord'), 3, 'left'), (L('subsub_lord'), 3, 'left')]
    r = g.section(r, L('cusps'))
    cols = [(L('cusp'), 2, 'center'), (L('rasi'), 3, 'left'), (L('degree'), 3, 'center'),
            (L('nakshatra'), 5, 'left')] + lord_cols
    rows = [[c['house'], i18n.rasi(c['sign'], lang), c['dms'], i18n.nak(c['nak'], lang),
             P(c['sign_lord']), P(c['star_lord']), P(c['sub_lord']), P(c['subsub_lord'])] for c in kp['cusps']]
    r = g.table(r, cols, rows) + 1

    g.page_break(r)
    r = g.section(r, L('kp_planets'))
    cols = [(L('planet'), 3, 'left'), (L('rasi'), 3, 'left'), (L('degree'), 3, 'center'),
            (L('nakshatra'), 4, 'left'), (L('rasi_lord'), 2, 'left'), (L('star_lord'), 2, 'left'),
            (L('sub_lord'), 2, 'left'), (L('subsub_lord'), 2, 'left'), (L('house'), 2, 'center'),
            (L('retro'), 2, 'center')]
    rows = [[P(p['name']), i18n.rasi(p['sign'], lang), p['dms'], i18n.nak(p['nak'], lang),
             P(p['sign_lord']), P(p['star_lord']), P(p['sub_lord']), P(p['subsub_lord']), p['house'],
             i18n.retro_mark(lang).strip('()') if p['retro'] else '—'] for p in kp['planets']]
    r = g.table(r, cols, rows) + 1

    r = g.section(r, L('planet_sig'))
    cols = [(L('planet'), 3, 'left'), (L('star_lord'), 3, 'left'), (L('sub_lord'), 3, 'left'),
            (L('level1'), 4, 'center'), (L('level2'), 4, 'center'), (L('level3'), 4, 'center'),
            (L('level4'), 4, 'center')]
    rows = [[P(s['planet']), P(s['star_lord']), P(s['sub_lord']), _houses(s['l1']), _houses(s['l2']),
             _houses(s['l3']), _houses(s['l4'])] for s in kp['planet_significators']]
    r = g.table(r, cols, rows) + 1
    g.height(r - len(rows) - 2, 44)

    r = g.section(r, L('house_sig'))
    cols = [(L('house'), 2, 'center'), (L('sig_a'), 7, 'left'), (L('sig_b'), 6, 'left'),
            (L('sig_c'), 7, 'left'), (L('sig_d'), 3, 'left')]
    rows = [[s['house'], i18n.planets(s['a'], lang), i18n.planets(s['b'], lang), i18n.planets(s['c'], lang),
             i18n.planets(s['d'], lang)] for s in kp['house_significators']]
    hdr = r
    r = g.table(r, cols, rows, row_h=30 if lang == 'bi' else 22) + 1
    g.height(hdr, 38)
    for rr in range(hdr + 1, hdr + 1 + len(rows)):
        for cc in range(COL0, COL0 + GRID_N):
            cell = g.ws.cell(row=rr, column=cc)
            cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical='center',
                                       wrap_text=True, indent=cell.alignment.indent)

    bal = kp['dasa']['balance']
    g.page_break(r)
    r = g.section(r, f"{L('kp_dasa')}  ·  {L('dasa_balance')}: {_balance_text(bal, lang)}")
    cols = [(L('dasa'), 5, 'left'), (L('bhukti'), 5, 'left'), (L('start'), 4, 'center'),
            (L('end'), 4, 'center'), (L('age'), 3, 'center'), (L('status'), 4, 'center')]
    rows, current = [], set()
    for d in kp['dasa']['dasas']:
        for b in d['sub']:
            if b['current']:
                current.add(len(rows))
            rows.append([P(d['lord']), P(b['lord']), (_d(b['start']), DATE_FMT), (_d(b['end']), DATE_FMT),
                         _age(b['start'], m['local_dt']), _status(b, now, lang)])
    g.table(r, cols, rows, current=current)
    g.ws.freeze_panes = 'A3'
    _page_setup(g.ws, m['name'], COL0 + GRID_N - 1)


def _alp(wb, res, specs, lang):
    g = Grid(wb, _sheet_name('alp', lang), lang)
    m, v, alp = res['meta'], res['vedic'], res['alp']
    L = lambda k: i18n.L(k, lang)
    P = lambda n: i18n.planet(n, lang)
    now = m['generated_at']
    al, lagna = alp['lagna'], v['planets'][0]

    r = g.title(1, f"{L('alp_title')}  ·  {m['name']}", alp['rate'])
    top = r
    g.chart(top, COL0, specs['alp'], m['chart_style'])
    rr = g.section(top, L('alp_lagna_now'), c=RIGHT0, span=12, keep_height=True)
    rr = g.pairs(rr, RIGHT0, keep_height=True, items=[
        (L('alp_birth_lagna'), f"{i18n.rasi(lagna['sign'], lang)}  {lagna['dms']}"),
        (L('as_of'), now.date(), DATE_FMT),
        (L('age'), f"{alp['age_years']:.2f} {i18n.Lv('years', lang)}"),
        (L('alp_lagna_now'), f"{i18n.rasi(al['sign'], lang)}  {al['dms']}"),
        (L('nakshatra'), f"{i18n.nak(al['nak'], lang)} - {i18n.Lv('pada', lang)} {al['pada']}"),
        (L('rasi_lord'), P(al['sign_lord'])),
        (L('star_lord'), P(al['star_lord'])),
        (L('sub_lord'), P(al['sub_lord'])),
    ])
    r = top + 14

    r = g.section(r, L('alp_sign_periods'))
    cols = [(L('from'), 4, 'center'), (L('to'), 4, 'center'), (L('age_from'), 3, 'center'),
            (L('age_to'), 3, 'center'), (L('rasi'), 4, 'left'), (L('rasi_lord'), 4, 'left'),
            (L('status'), 3, 'center')]
    rows, current = [], set()
    for s in alp['sign_periods']:
        if s['current']:
            current.add(len(rows))
        rows.append([(_d(s['start']), DATE_FMT), (_d(s['end']), DATE_FMT), s['age_from'], s['age_to'],
                     i18n.rasi(s['sign'], lang), P(s['sign_lord']), _status(s, now, lang)])
    r = g.table(r, cols, rows, current=current) + 1

    g.page_break(r)
    r = g.section(r, L('alp_pada_periods'))
    cols = [(L('from'), 3, 'center'), (L('to'), 3, 'center'), (L('age_from'), 2, 'center'),
            (L('age_to'), 2, 'center'), (L('rasi'), 3, 'left'), (L('nakshatra'), 4, 'left'),
            (L('pada'), 1, 'center'), (L('rasi_lord'), 2, 'left'), (L('star_lord'), 2, 'left'),
            (L('status'), 3, 'center')]
    rows, current = [], set()
    for s in alp['pada_periods']:
        if s['current']:
            current.add(len(rows))
        rows.append([(_d(s['start']), DATE_FMT), (_d(s['end']), DATE_FMT), s['age_from'], s['age_to'],
                     i18n.rasi(s['sign'], lang), i18n.nak(s['nak'], lang), s['pada'],
                     P(s['sign_lord']), P(s['star_lord']), _status(s, now, lang)])
    g.table(r, cols, rows, current=current)
    g.ws.freeze_panes = 'A3'
    _page_setup(g.ws, m['name'], COL0 + GRID_N - 1)


def _dasa(wb, res, lang):
    """Vimshottari dasa / bhukti / antaram in ordinary columns (sortable, filterable)."""
    ws = wb.create_sheet(_sheet_name('dasa_sheet', lang))
    ws.sheet_view.showGridLines = False
    m, v = res['meta'], res['vedic']
    L = lambda k: i18n.L(k, lang)
    P = lambda n: i18n.planet(n, lang)
    now = m['generated_at']
    wide = 24 if lang == 'bi' else 15
    widths = [1.5, wide, wide, wide, 13, 13, 9, 9, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=9)
    ws.row_dimensions[1].height = 30
    c = ws.cell(row=1, column=2, value=f"{L('vim_dasa')}  ·  {m['name']}")
    c.font, c.fill = _font(15, True, WHITE), _fill(BAND)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=9)
    c = ws.cell(row=2, column=2, value=f"{L('dasa_balance')}: {_balance_text(v['dasa']['balance'], lang)}"
                                       f"  ·  {L('ayanamsha')}: {m['ayanamsha_name']}")
    c.font, c.fill = _font(9, False, MUTED, True), _fill(SOFT)
    c.alignment = Alignment(horizontal='center', vertical='center')

    heads = [L('dasa'), L('bhukti'), L('antaram'), L('start'), L('end'),
             i18n.Lj(['age', 'start'], lang), L('years'), L('status')]
    ws.row_dimensions[3].height = 30 if lang == 'bi' else 22
    for i, h in enumerate(heads, start=2):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.fill, c.border = _font(9, True, WHITE), _fill(HEAD), border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    r = 4
    for d in v['dasa']['dasas']:
        for b in d['sub']:
            for a in b['sub']:
                cur = a['current']
                fill = _fill(CURRENT if cur else (ALT if r % 2 else WHITE))
                vals = [P(d['lord']), P(b['lord']), P(a['lord']), _d(a['start']), _d(a['end']),
                        _age(a['start'], m['local_dt']), round(a['years'], 3), _status(a, now, lang)]
                for i, val in enumerate(vals, start=2):
                    c = ws.cell(row=r, column=i, value=val)
                    c.font, c.fill, c.border = _font(9, cur), fill, border
                    c.alignment = _align('left' if i <= 4 else 'center', 'center', False, False,
                                         1 if i <= 4 else 0)
                    if i in (5, 6):
                        c.number_format = DATE_FMT
                    elif i == 7:
                        c.number_format = '0.00'
                    elif i == 8:
                        c.number_format = '0.000'
                r += 1
    ws.auto_filter.ref = f'B3:I{r - 1}'
    ws.freeze_panes = 'A4'
    _page_setup(ws, m['name'], 9, title_rows='1:3')


def _notes(wb, res, lang, ai=None):
    ws = wb.create_sheet(_sheet_name('notes', lang))
    ws.sheet_view.showGridLines = False
    m, v, kp = res['meta'], res['vedic'], res['kp']
    L = lambda k: i18n.L(k, lang)
    ws.column_dimensions['A'].width = 1.5
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 78
    border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    def band(r, text, fill=HEAD, size=10.5, h=20):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = h
        c = ws.cell(row=r, column=2, value=text)
        c.font, c.fill = _font(size, True, WHITE), _fill(fill)
        c.alignment = Alignment(horizontal='left' if size < 14 else 'center', vertical='center', indent=1)
        return r + 1

    def row(r, k, val):
        a = ws.cell(row=r, column=2, value=k)
        a.font, a.fill, a.border = _font(9, True, MUTED), _fill(SOFT), border
        a.alignment = Alignment(vertical='top', wrap_text=True, indent=1)
        b = ws.cell(row=r, column=3, value=val)
        b.font, b.border = _font(9), border
        b.alignment = Alignment(vertical='top', wrap_text=True, indent=1)
        lines = max(1, -(-len(str(val)) // 95), -(-len(str(k)) // 34))
        ws.row_dimensions[r].height = 14 * lines + 4
        return r + 1

    r = band(1, f"{L('notes')}  ·  {m['name']}", fill=BAND, size=15, h=30) + 1
    r = band(r, L('how_calculated'))
    node = 'True node' if m['node'] == 'true' else 'Mean node'
    for k, val in [
        ('Ephemeris', f"{m['engine']}. Positions are geocentric, apparent, sidereal."),
        ('Birth moment', f"Local time {m['local_dt'].strftime('%d-%m-%Y %H:%M:%S')} at {m['tz']} "
                         f"({m['utc_offset_str']}) = UT {m['ut_dt'].strftime('%d-%m-%Y %H:%M:%S')}; "
                         f"Julian day (UT) {m['jd_ut']}."),
        ('Place', f"{m['pob']} — latitude {m['lat']:.4f}°, longitude {m['lon']:.4f}°. "
                  "If these are not the birth place, the lagna and house cusps will be wrong."),
        ('Vedic ayanamsha', f"{m['ayanamsha_name']}: {v['ayanamsha_dms']} at birth."),
        ('KP ayanamsha', f"Krishnamurti: {kp['ayanamsha_dms']} at birth. The KP sheet uses this value throughout, "
                         "so KP positions differ slightly from the Vedic sheet."),
        ('Rahu / Ketu', f"{node}. Ketu is exactly opposite Rahu."),
        ('Vedic houses', 'House = whole sign counted from the lagna sign. Bhava = Sripati: the Porphyry cusps '
                         'are the bhava centres and each bhava begins midway between two centres.'),
        ('KP houses', f"{kp['house_system']} cusps; a planet belongs to the house whose cusp it has passed. "
                      'Star, sub and sub-sub lords divide each nakshatra in Vimshottari proportion.'),
        ('KP significators', 'Planet levels 1-4: house occupied by its star lord; house it occupies; houses '
                             'owned by its star lord; houses it owns. House levels A-D: planets in the star of '
                             'occupants; occupants; planets in the star of the owner; the owner. '
                             'Rahu and Ketu own no houses and also act for their sign lord and for planets '
                             'in the same sign (listed separately).'),
        ('Vimshottari', f"Year of {m['year_days']} days. Balance at birth from the Moon's position in its "
                        'nakshatra. Periods that ended before birth are not listed.'),
        ('Panchangam', 'Tithi, yoga and karana at the birth moment. Sunrise and sunset are for the visible '
                       'upper limb with standard refraction. The weekday runs from sunrise to sunrise.'),
        ('Mandi', _mandi_note(m, v)),
        ('Combustion', 'Orbs from the Sun: Moon 12°, Mars 17°, Mercury 14° (12° retrograde), Jupiter 11°, '
                       'Venus 10° (8° retrograde), Saturn 15°.'),
        ('Dignity', 'Only exaltation, debilitation and own sign are marked.'),
        ('ALP', 'Akshaya Lagna Paddhati: the birth lagna advances 30° every 10 years (3° a year), so one '
                f"nakshatra pada lasts 1 year 1 month 10 days. Year of {m['year_days']} days."),
        ('Generated', m['generated_at'].strftime('%d-%m-%Y %H:%M') + f" ({m['utc_offset_str']})"),
    ]:
        r = row(r, k, val)

    r = band(r + 1, L('divisional'))
    r = row(r, L('tob'), L('varga_time_note') if lang in ('ta', 'hi', 'mr', 'bi') else i18n.LABELS['en']['varga_time_note'])
    for vg in v['vargas']:
        r = row(r, i18n.varga_name(vg['key'], lang), vg['rule'])
    r = row(r, L('vargottama'), i18n.LABELS['en']['vargottama_note'])

    r = band(r + 1, L('abbreviations'))
    names = ['Lagna', 'Sun', 'Moon', 'Mars', 'Mercury', 'Jupiter', 'Venus', 'Saturn', 'Rahu', 'Ketu', 'Mandi', 'ALP']
    r = row(r, L('planet'), '   '.join(
        f"{i18n.abbr(n, lang)} = {i18n.planet(n, lang) if n != 'ALP' else L('alp_lagna_now')}" for n in names))
    r = row(r, L('retro'), f"{i18n.retro_mark(lang)} after a planet = retrograde")
    r = row(r, L('rasi'), '   '.join(f"{i + 1} = {i18n.rasi(i, lang)}" for i in range(12)))
    r = row(r, L('kp_chart'), 'Roman numerals I-XII mark the sign in which each house cusp falls.')
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=3)
    c = ws.cell(row=r + 1, column=2, value=L('disclaimer'))
    c.font = _font(8, False, MUTED, True)
    c.alignment = Alignment(horizontal='center')
    if ai:
        # Added below everything that was already on the sheet, so no existing cell moves.
        r = band(r + 3, 'AI sheets: rules and variants')
        r = row(r, 'About', 'The sheets named AI_ repeat this horoscope as flat tables for AI readers and add '
                            'derived data. They are always in English. AI_ReadMe explains every column. '
                            'The rules below say how each derived value was obtained and which variant was used.')
        for item in ai['rules']:
            r = row(r, item['rule'], item['variant'])
    _page_setup(ws, m['name'], 3)


# ── AI SHEETS ─────────────────────────────────────────────────────────────────
# Flat tables for AI readers: header in row 1, one table per sheet from A1,
# values only, no merged cells, no formulas, no colour that carries meaning.

AI_DATETIME_FMT = 'yyyy-mm-dd hh:mm:ss'
AI_DATE_FMT = 'yyyy-mm-dd'
AI_DEG_FMT = '0.000000'
AI_WRAP = {('AI_ReadMe', 'text'): 120, ('AI_Facts', 'note'): 70, ('AI_Yogas', 'rule'): 75,
           ('AI_Yogas', 'note'): 70}
AI_DEG_COLUMNS = ('deg_in_sign', 'deg_in_chart', 'lagna_deg_in_chart', 'speed_deg_per_day')
AI_MAX_WIDTH = 60


def _ai_sheets(wb, ai):
    """One worksheet per table of the AI block, in the order given by ai['sheets']."""
    head_font = _font(10, True)
    wrap = Alignment(vertical='top', wrap_text=True)
    top = Alignment(vertical='top')
    for sheet in ai['sheets']:
        ws = wb.create_sheet(sheet)
        cols = [c['name'] for c in ai['columns'][sheet]]
        rows = ai['tables'][sheet]
        ws.append(cols)
        for c in ws[1]:
            c.font = head_font
        widths = [len(name) for name in cols]
        dated = set()                               # columns holding at least one date or datetime
        for row in rows:
            values = []
            for i, name in enumerate(cols):
                val = row.get(name)
                if isinstance(val, datetime):
                    dated.add(i)
                    size = 19
                elif isinstance(val, date):
                    dated.add(i)
                    size = 10
                elif val is None:
                    size = 0
                else:
                    size = len(str(val))
                if size > widths[i]:
                    widths[i] = size
                values.append(val)
            ws.append(values)
        last = len(rows) + 1
        tall = any(s == sheet for s, _ in AI_WRAP)       # rows grow with wrapped text: keep every cell at the top
        for i, name in enumerate(cols):
            letter = get_column_letter(i + 1)
            degrees = name.endswith('_deg') or name in AI_DEG_COLUMNS
            wrap_width = AI_WRAP.get((sheet, name))
            if degrees or wrap_width or tall or i in dated:
                # The format follows the value in each cell, not the column: AI_Facts
                # keeps dates, numbers and text in one "value" column.
                for (cell,) in ws.iter_rows(min_row=2, max_row=last, min_col=i + 1, max_col=i + 1):
                    val = cell.value
                    if isinstance(val, datetime):
                        cell.number_format = AI_DATETIME_FMT
                    elif isinstance(val, date):
                        cell.number_format = AI_DATE_FMT
                    elif degrees and isinstance(val, float):
                        cell.number_format = AI_DEG_FMT
                    if wrap_width:
                        cell.alignment = wrap
                    elif tall:
                        cell.alignment = top
            ws.column_dimensions[letter].width = wrap_width or min(max(widths[i], 6) + 2, AI_MAX_WIDTH)
        ws.freeze_panes = 'A2'
        # If printed: landscape, all columns on one page width, as many pages down as needed.
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = 'landscape'
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:1'


# ── ENTRY POINT ───────────────────────────────────────────────────────────────

def generate_excel(res, lang=None, ai=None):
    """
    Build the workbook from compute() output. Returns bytes.
    ai: the block returned by enrich.enrich(res); when given, the AI_ sheets
    are added after Notes and the Notes sheet gains a section on their rules.
    """
    lang = lang or res['meta'].get('lang', 'en')
    if lang not in i18n.LANGS:
        lang = 'en'
    specs = charts.build_specs(res, lang)
    wb = openpyxl.Workbook()
    _summary(wb, res, specs, lang)
    _vedic(wb, res, specs, lang)
    _divisional(wb, res, lang)
    _kp(wb, res, specs, lang)
    _alp(wb, res, specs, lang)
    _dasa(wb, res, lang)
    _notes(wb, res, lang, ai)
    if ai:
        _ai_sheets(wb, ai)
    wb.properties.title = f"HoroscopeGen - {res['meta']['name']}"
    wb.properties.creator = 'horoscopegen.in'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
