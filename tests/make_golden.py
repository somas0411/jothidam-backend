"""
Fingerprint of the output that must not change when the AI sheets are added:
the seven report sheets of the workbook and the /api/horoscope response.

    python tests/make_golden.py > tests/golden_main.json

tests/golden_main.json was produced by running this file on `main` at commit
bd112bd, before enrich.py existed. tests/test_enrich.py compares today's
output with it, so "the old sheets are unchanged" is checked against what
main really produced and not against the code under test.

Run it again only after an intended change to the seven sheets or to the
API, or after upgrading openpyxl or Flask (which can change formats).
"""
import functools
import hashlib
import io
import json
import sys
from datetime import date, datetime, time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

NOW = datetime(2026, 10, 7, 17, 0)
CHART_S = dict(name='Chart S', dob='1984-11-04', tob='00:20', pob='Chennai', lat=13.0827, lon=80.2707)
API_BODY = {'name': 'Chart S', 'dob': '1984-11-04', 'tob': '00:20', 'pob': 'Chennai',
            'lat': 13.0827, 'lon': 80.2707, 'lang': 'en', 'chartStyle': 'south', 'ayanamsha': 'lahiri'}
LANGS = ('en', 'ta')


def _value(v):
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    return v


def _side(s):
    return (s.style, s.color.rgb if s is not None and s.color is not None else None) if s is not None else None


def cell_print(c):
    """Everything about a cell a reader could see: value, format, font, fill, alignment, borders."""
    f, a, b = c.font, c.alignment, c.border
    return [_value(c.value), c.number_format, f.name, f.sz, bool(f.b), bool(f.i),
            f.color.rgb if f.color is not None else None,
            c.fill.fill_type, c.fill.fgColor.rgb if c.fill.fill_type else None,
            a.horizontal, a.vertical, bool(a.wrap_text), bool(a.shrink_to_fit), a.indent,
            _side(b.left), _side(b.right), _side(b.top), _side(b.bottom), _side(b.diagonal),
            bool(b.diagonalUp), bool(b.diagonalDown)]


def _digest(obj):
    return hashlib.sha1(json.dumps(obj, ensure_ascii=False, sort_keys=True, default=str).encode('utf-8')).hexdigest()[:16]


def sheet_print(ws, rows=None):
    """
    Fingerprint of a worksheet. rows limits it to the first rows of the sheet
    (used for Notes, which may gain rows below what main had).
    """
    last = rows or ws.max_row
    cols = ws.max_column
    row_digests = [_digest([cell_print(c) for c in row])
                   for row in ws.iter_rows(min_row=1, max_row=last, min_col=1, max_col=cols)]
    merged = sorted(str(r) for r in ws.merged_cells.ranges if r.max_row <= last)
    heights = {r: ws.row_dimensions[r].height for r in range(1, last + 1) if ws.row_dimensions[r].height is not None}
    widths = {k: d.width for k, d in sorted(ws.column_dimensions.items()) if d.width is not None}
    return {'title': ws.title, 'rows': last, 'columns': cols, 'row_digests': row_digests,
            'merged': _digest(merged), 'merged_count': len(merged), 'heights': _digest(heights),
            'widths': _digest(widths), 'freeze': ws.freeze_panes,
            'breaks': sorted(b.id for b in ws.row_breaks.brk if b.id <= last)}


def workbook_print(data, limit=None):
    """Fingerprints of the first seven sheets. limit: {sheet index: rows} to compare fewer rows."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return [sheet_print(ws, (limit or {}).get(i)) for i, ws in enumerate(wb.worksheets[:7])]


def api_print():
    """SHA-256 of the /api/horoscope response body for chart S at a fixed report moment."""
    import app as app_module
    import astro_engine as A
    original = app_module.compute
    app_module.compute = functools.partial(A.compute, now=NOW)
    try:
        rv = app_module.app.test_client().post('/api/horoscope', json=API_BODY)
    finally:
        app_module.compute = original
    assert rv.status_code == 200
    return {'sha256': hashlib.sha256(rv.data).hexdigest(), 'bytes': len(rv.data)}


def main():
    import astro_engine as A
    from excel_generator import generate_excel
    out = {'now': NOW.isoformat(), 'chart': CHART_S, 'workbooks': {}, 'api_horoscope': api_print()}
    for lang in LANGS:
        res = A.compute(lang=lang, now=NOW, **CHART_S)
        out['workbooks'][lang] = workbook_print(generate_excel(res, lang))
    json.dump(out, sys.stdout, ensure_ascii=False, indent=1)
    sys.stdout.write('\n')


if __name__ == '__main__':
    main()
