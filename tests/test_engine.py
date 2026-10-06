"""
Tests for the calculation engine and the report generators.

Reference values that do not come from this code:
  * KP sub-lord table for Ashwini (the standard published 249-sub table).
  * Vimshottari proportions (Venus-Venus = 3 years 4 months, 120-year cycle).
  * Navamsa starting signs for fire / earth / air / water signs.
  * Indian Standard Time history from the tz database.
The sample chart is checked against the Swiss Ephemeris directly.
"""
import io
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
import swisseph as swe

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import astro_engine as A          # noqa: E402
import charts                     # noqa: E402
import i18n                       # noqa: E402

NOW = datetime(2026, 10, 3, 20, 0)
ARCMIN = 1 / 60.0


@pytest.fixture(scope='module')
def res():
    return A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090, now=NOW)


def dms(d, m=0, s=0):
    return d + m / 60 + s / 3600


# ── pure functions ────────────────────────────────────────────────────────────

def test_rasi_nakshatra_pada():
    assert A.get_rasi(0) == 0 and A.get_rasi(29.999) == 0 and A.get_rasi(30) == 1 and A.get_rasi(359.9) == 11
    assert A.get_nak(0) == 0 and A.get_nak(13.34) == 1 and A.get_nak(359.9) == 26
    assert [A.get_pada(x) for x in (0, 3.34, 6.67, 10.01)] == [1, 2, 3, 4]
    # exact boundaries belong to the next division
    assert A.get_nak(40 / 3) == 1 and A.get_pada(10 / 3) == 2 and A.get_nak(360 - 1e-9) == 26


def test_navamsa_starting_signs():
    assert A.navamsa_sign(0) == 0            # Mesha (fire) starts from Mesha
    assert A.navamsa_sign(dms(3, 21)) == 1
    assert A.navamsa_sign(30) == 9           # Rishabha (earth) starts from Makara
    assert A.navamsa_sign(60) == 6           # Mithuna (air) starts from Thula
    assert A.navamsa_sign(90) == 3           # Kataka (water) starts from Kataka
    assert A.navamsa_sign(359.99) == 11


def test_dms_rounding_carries():
    assert A.fmt_sign_dms(dms(25, 58, 15)) == '25°58\'15"'
    assert A.fmt_sign_dms(29.9999999) == '00°00\'00"'      # rounds up into the next sign, never 29°59'60"
    assert A.fmt_dms(23.9999999) == '24°00\'00"'


@pytest.mark.parametrize('lon,sub', [
    (dms(0, 0, 1), 'Ketu'), (dms(0, 46, 39), 'Ketu'), (dms(0, 46, 41), 'Venus'),
    (dms(2, 59, 59), 'Venus'), (dms(3, 0, 1), 'Sun'), (dms(3, 40, 1), 'Moon'),
    (dms(4, 46, 41), 'Mars'), (dms(5, 33, 21), 'Rahu'), (dms(7, 33, 21), 'Jupiter'),
    (dms(9, 20, 1), 'Saturn'), (dms(11, 26, 41), 'Mercury'), (dms(13, 19, 59), 'Mercury'),
])
def test_kp_sub_lords_match_published_table(lon, sub):
    sign_lord, star_lord, sub_lord, _ = A.kp_lords(lon)
    assert (sign_lord, star_lord, sub_lord) == ('Mars', 'Ketu', sub)


def test_kp_sub_table_has_249_divisions():
    seen, prev = 0, None
    steps = 360 * 3600 // 20
    for i in range(steps):
        lon = (i * 20 + 10) / 3600.0
        key = (A.get_rasi(lon),) + A.kp_lords(lon)[1:3]
        if key != prev:
            seen += 1
            prev = key
    assert seen == 249


def test_sub_sub_lord_starts_with_sub_lord():
    assert A.kp_lords(dms(0, 46, 41))[3] == 'Venus'
    assert A.kp_lords(0.0001)[3] == 'Ketu'


# ── time zone ─────────────────────────────────────────────────────────────────

def test_ist_and_wartime_offset():
    off, name = A.resolve_utc_offset(datetime(1998, 3, 6, 1, 37), 28.6, 77.2)
    assert name == 'Asia/Kolkata' and off == 5.5
    off, _ = A.resolve_utc_offset(datetime(1943, 6, 1, 12, 0), 28.6, 77.2)
    assert off == 6.5            # war time, 1942-1945


def test_manual_offset_wins():
    off, label = A.resolve_utc_offset(datetime(1998, 3, 6), 28.6, 77.2, tz_name='Asia/Kolkata', utc_offset=-5)
    assert off == -5 and label == 'Manual offset'


def test_unknown_place_is_an_error_not_a_guess(monkeypatch):
    monkeypatch.setattr(A, 'geocode', lambda q, limit=1: [])
    with pytest.raises(A.InputError):
        A.compute('X', '1998-03-06', '01:37', 'Nowhereville')


def test_bad_input():
    with pytest.raises(A.InputError):
        A.compute('X', '06/03/1998', '01:37', 'Delhi', lat=28.6, lon=77.2)
    with pytest.raises(A.InputError):
        A.compute('X', '1998-03-06', '01:37', 'Delhi', lat=128.6, lon=77.2)


# ── sample chart against the Swiss Ephemeris ──────────────────────────────────

def test_ayanamsha_values():
    swe.set_sid_mode(swe.SIDM_LAHIRI)
    lahiri_2000 = swe.get_ayanamsa_ut(swe.julday(2000, 1, 1, 0))
    assert abs(lahiri_2000 - dms(23, 51)) < 0.02          # Lahiri is 23°51' at 2000


def test_positions_match_swiss_ephemeris(res):
    jd = swe.julday(1998, 3, 5, 20 + 7 / 60)               # 01:37 IST = 20:07 UT the day before
    assert abs(res['meta']['jd_ut'] - jd) < 1e-5
    swe.set_sid_mode(swe.SIDM_LAHIRI)
    flags = swe.FLG_MOSEPH | swe.FLG_SIDEREAL | swe.FLG_SPEED
    got = {p['name']: p for p in res['vedic']['planets']}
    for name, pid in A.SWE_IDS.items():
        want = swe.calc_ut(jd, pid, flags)[0][0]
        assert A.angle_diff(got[name]['lon'], want) < ARCMIN, name
    asc = swe.houses_ex(jd, 28.6139, 77.2090, b'P', swe.FLG_SIDEREAL)[1][0]
    assert A.angle_diff(got['Lagna']['lon'], asc) < ARCMIN
    assert A.angle_diff(got['Ketu']['lon'], got['Rahu']['lon'] + 180) < 1e-6


def test_sample_chart_key_facts(res):
    v = res['vedic']
    got = {p['name']: p for p in v['planets']}
    assert A.RASIS[v['lagna_sign']] == 'Vrischika'
    assert A.RASIS[v['moon_sign']] == 'Rishabha'
    assert A.NAKS[v['moon_nak']] == 'Mrigashira' and v['moon_pada'] == 2
    assert A.RASIS[got['Venus']['sign']] == 'Makara'
    assert A.RASIS[got['Mercury']['sign']] == 'Meena' and got['Mercury']['dignity'] == 'Debilitated'
    assert got['Moon']['dignity'] == 'Exalted'
    assert got['Sun']['house'] == 4 and got['Moon']['house'] == 7
    assert not got['Rahu']['retro'] and not got['Lagna']['retro']


def test_panchangam(res):
    pc = res['vedic']['panchangam']
    assert pc['paksha'] == 'Shukla' and pc['tithi'] == 'Ashtami'
    # 06-03-1998 was a Friday, but 01:37 is before sunrise, so the Vedic day is Thursday.
    assert date(1998, 3, 6).weekday() == 4
    assert pc['vara'] == 'Thursday' and pc['vara_lord'] == 'Jupiter'
    assert time(6, 30) < pc['sunrise'].time() < time(6, 55)
    assert time(18, 10) < pc['sunset'].time() < time(18, 35)


# ── Vimshottari ───────────────────────────────────────────────────────────────

def test_vimshottari_structure(res):
    d = res['vedic']['dasa']
    birth = res['meta']['local_dt']
    assert d['balance']['lord'] == 'Mars' and 4.9 < d['balance']['years'] < 5.1
    lords = [x['lord'] for x in d['dasas']]
    assert lords == ['Mars', 'Rahu', 'Jupiter', 'Saturn', 'Mercury', 'Ketu', 'Venus', 'Sun', 'Moon']
    assert d['dasas'][0]['start'] == birth
    # contiguous at every level, nothing before birth
    for dasa in d['dasas']:
        assert dasa['sub'][0]['start'] == dasa['start'] and dasa['sub'][-1]['end'] == dasa['end']
        for a, b in zip(dasa['sub'], dasa['sub'][1:]):
            assert a['end'] == b['start']
        for bh in dasa['sub']:
            assert bh['sub'][0]['start'] == bh['start'] and bh['sub'][-1]['end'] == bh['end']
    # the first dasa lists only the bhuktis still to run, not the whole dasa from its first bhukti
    assert d['dasas'][0]['sub'][0]['lord'] != 'Mars'
    assert len(d['dasas'][0]['sub']) < 9 and all(len(x['sub']) == 9 for x in d['dasas'][1:])
    # full dasas have their full length
    rahu = d['dasas'][1]
    assert abs((rahu['end'] - rahu['start']).days - 18 * 365.25) <= 1


def test_vimshottari_proportions():
    # Moon at 0° of Bharani: Venus dasa starts exactly at birth.
    birth = datetime(2000, 1, 1)
    d = A.vimshottari(A.NAK_SIZE, birth, now=birth)
    assert d['balance']['lord'] == 'Venus' and abs(d['balance']['years'] - 20) < 1e-9
    vv = d['dasas'][0]['sub'][0]
    assert vv['lord'] == 'Venus'
    assert abs((vv['end'] - vv['start']).total_seconds() / 86400 - (20 * 20 / 120) * 365.25) < 1e-6
    total = (d['dasas'][-1]['end'] - d['dasas'][0]['start']).total_seconds() / 86400
    assert abs(total - 120 * 365.25) < 1e-6


def test_exactly_one_current_period(res):
    d = res['vedic']['dasa']
    assert sum(x['current'] for x in d['dasas']) == 1
    cur = d['current']
    assert cur['dasa']['start'] <= NOW < cur['dasa']['end']
    assert cur['bhukti']['start'] <= NOW < cur['bhukti']['end']
    assert cur['antara']['start'] <= NOW < cur['antara']['end']


# ── KP ────────────────────────────────────────────────────────────────────────

def test_kp_cusps_and_houses(res):
    kp = res['kp']
    assert kp['house_system'] == 'Placidus' and len(kp['cusps']) == 12
    c = [x['lon'] for x in kp['cusps']]
    for i in range(6):                                   # opposite cusps are 180° apart
        assert A.angle_diff(c[i] + 180, c[i + 6]) < 1e-6
    for p in kp['planets']:
        h = p['house']
        assert A.in_arc(p['lon'], c[h - 1], c[h % 12])
    # KP ayanamsha is smaller than Lahiri by about 6 arc-minutes
    assert 0.05 < res['vedic']['ayanamsha'] - kp['ayanamsha'] < 0.15


def test_kp_significators_follow_the_rules(res):
    kp = res['kp']
    by = {p['name']: p for p in kp['planets']}
    owned = {pl: [x['house'] for x in kp['cusps'] if x['sign_lord'] == pl] for pl in A.PLANETS}
    for s in kp['planet_significators']:
        p = by[s['planet']]
        assert s['l1'] == [by[p['star_lord']]['house']]
        assert s['l2'] == [p['house']]
        assert s['l3'] == owned[p['star_lord']] and s['l4'] == owned[p['name']]
    assert owned['Rahu'] == [] and owned['Ketu'] == []
    assert sorted(h for v in owned.values() for h in v) == list(range(1, 13))
    for s in kp['house_significators']:
        occ = [p['name'] for p in kp['planets'] if p['house'] == s['house']]
        assert s['b'] == occ and s['d'] == [kp['cusps'][s['house'] - 1]['sign_lord']]
        assert s['a'] == [p['name'] for p in kp['planets'] if p['star_lord'] in occ]


def test_ruling_planets(res):
    rp = res['kp']['ruling_planets']
    moon = next(p for p in res['kp']['planets'] if p['name'] == 'Moon')
    assert rp['day_lord'] == 'Jupiter'
    assert (rp['moon_sign_lord'], rp['moon_star_lord']) == (moon['sign_lord'], moon['star_lord'])
    assert rp['lagna_sign_lord'] == res['kp']['cusps'][0]['sign_lord']


# ── ALP ───────────────────────────────────────────────────────────────────────

def test_alp_moves_thirty_degrees_in_ten_years():
    birth = datetime(2000, 1, 1)
    alp = A._build_alp(100.0, birth, now=birth + timedelta(days=10 * 365.25))
    assert abs(alp['lagna']['lon'] - 130.0) < 1e-6
    assert abs(alp['age_years'] - 10) < 0.01


def test_alp_periods(res):
    alp = res['alp']
    lagna = res['vedic']['planets'][0]
    pp, sp = alp['pada_periods'], alp['sign_periods']
    assert pp[0]['start'] == res['meta']['local_dt'] and pp[0]['sign'] == lagna['sign']
    assert (pp[0]['nak'], pp[0]['pada']) == (lagna['nak'], lagna['pada'])
    for a, b in zip(pp, pp[1:]):
        assert a['end'] == b['start']
    full = pp[1]                                        # a complete pada = 10/9 year
    assert abs((full['end'] - full['start']).total_seconds() / 86400 - (10 / 9) * 365.25) < 1e-3
    assert abs((sp[1]['end'] - sp[1]['start']).total_seconds() / 86400 - 10 * 365.25) < 1e-3
    assert sum(x['current'] for x in pp) == 1 and sum(x['current'] for x in sp) == 1
    cur = next(x for x in pp if x['current'])
    assert (cur['sign'], cur['nak'], cur['pada']) == (alp['lagna']['sign'], alp['lagna']['nak'], alp['lagna']['pada'])
    assert abs((pp[-1]['end'] - pp[0]['start']).days - 120 * 365.25) <= 1


# ── Mandi ─────────────────────────────────────────────────────────────────────

DAYS = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']


def test_mandi_ghati_tables():
    assert [A.mandi_ghatis(d, True) for d in DAYS] == [26, 22, 18, 14, 10, 6, 2]
    assert [A.mandi_ghatis(d, False) for d in DAYS] == [10, 6, 2, 26, 22, 18, 14]


def test_mandi_time_on_an_equal_day_and_night():
    sunrise, sunset = datetime(2000, 1, 2, 6, 0), datetime(2000, 1, 2, 18, 0)      # a Sunday
    next_rise = datetime(2000, 1, 3, 6, 0)
    assert A.mandi_time(sunrise, sunset, A.mandi_ghatis('Sunday', True)) == datetime(2000, 1, 2, 16, 24)
    assert A.mandi_time(sunset, next_rise, A.mandi_ghatis('Sunday', False)) == datetime(2000, 1, 2, 22, 0)


def test_mandi_scales_with_day_length():
    # a 14-hour day: Saturday's 2 ghatis = 2/30 of 14 h = 56 minutes after sunrise
    rise = A.mandi_time(datetime(2000, 6, 3, 5, 0), datetime(2000, 6, 3, 19, 0), 2)
    assert rise == datetime(2000, 6, 3, 5, 56)


def _swe_lagna(local_moment, lat, lon, offset=5.5, mode=swe.SIDM_LAHIRI):
    ut = local_moment - timedelta(hours=offset)
    jd = swe.julday(ut.year, ut.month, ut.day, ut.hour + ut.minute / 60 + (ut.second + ut.microsecond / 1e6) / 3600)
    swe.set_sid_mode(mode)
    return swe.houses_ex(jd, lat, lon, b'P', swe.FLG_SIDEREAL)[1][0]


def test_mandi_before_sunrise_uses_previous_weekday_and_sunset(res):
    md = res['vedic']['mandi']
    pc = res['vedic']['panchangam']
    # 01:37 on Friday 06-03-1998 is before sunrise: Thursday night, counted from Thursday's sunset
    assert md['weekday'] == 'Thursday' == pc['vara'] and md['is_day'] is False and md['ghatis'] == 22
    assert md['period_start'].date() == date(1998, 3, 5) and time(18, 10) < md['period_start'].time() < time(18, 35)
    assert md['period_end'] == pc['sunrise']
    assert md['period_start'] < res['meta']['local_dt'] < md['period_end']
    night = md['period_end'] - md['period_start']
    assert md['rise_time'] == md['period_start'] + night * (22 / 30)
    assert A.angle_diff(md['lon'], _swe_lagna(md['rise_time'], 28.6139, 77.2090)) < ARCMIN
    assert md['house'] == ((md['sign'] - res['vedic']['lagna_sign']) % 12) + 1
    assert md['navamsa'] == A.navamsa_sign(md['lon'])
    assert md['retro'] is False and md['combust'] is False and md['dignity'] == ''


@pytest.mark.parametrize('tob,is_day,weekday,ghatis', [
    ('10:30', True, 'Friday', 6),          # day birth on Friday 06-03-1998
    ('21:15', False, 'Friday', 18),        # after sunset the same day: Friday night
])
def test_mandi_day_and_evening_births(tob, is_day, weekday, ghatis):
    r = A.compute('X', '1998-03-06', tob, 'Delhi', lat=28.6139, lon=77.2090, now=NOW)
    md, pc = r['vedic']['mandi'], r['vedic']['panchangam']
    assert (md['is_day'], md['weekday'], md['ghatis']) == (is_day, weekday, ghatis)
    if is_day:
        assert (md['period_start'], md['period_end']) == (pc['sunrise'], pc['sunset'])
    else:
        assert md['period_start'] == pc['sunset'] and md['period_end'].date() == date(1998, 3, 7)
    assert md['rise_time'] == A.mandi_time(md['period_start'], md['period_end'], ghatis)
    assert A.angle_diff(md['lon'], _swe_lagna(md['rise_time'], 28.6139, 77.2090)) < ARCMIN


def test_mandi_follows_the_chosen_ayanamsha():
    a = A.compute('X', '1998-03-06', '10:30', 'Delhi', lat=28.6139, lon=77.2090, now=NOW)
    b = A.compute('X', '1998-03-06', '10:30', 'Delhi', lat=28.6139, lon=77.2090, now=NOW, ayanamsha='raman')
    shift = a['vedic']['ayanamsha'] - b['vedic']['ayanamsha']
    assert abs(((b['vedic']['mandi']['lon'] - a['vedic']['mandi']['lon'] + 180) % 360 - 180) - shift) < 1e-4
    assert a['vedic']['mandi']['rise_time'] == b['vedic']['mandi']['rise_time']


def test_mandi_does_not_change_other_results(res):
    v, kp, alp = res['vedic'], res['kp'], res['alp']
    assert [p['name'] for p in v['planets']] == ['Lagna'] + A.PLANETS       # still lagna + 9 planets
    assert [p['name'] for p in kp['planets']] == A.PLANETS
    assert all(p['name'] != 'Mandi' for p in alp['planets'])
    assert all('Mandi' not in (s['a'] + s['b'] + s['c'] + s['d']) for s in kp['house_significators'])
    assert v['dasa']['balance']['lord'] == 'Mars'


def test_mandi_left_out_where_the_sun_does_not_set():
    # Tromso, Norway, at midsummer: the Sun stays up, so there is no Mandi
    r = A.compute('X', '2001-06-21', '12:00', 'Tromso', lat=69.65, lon=18.96, tz='Europe/Oslo', now=NOW)
    assert r['vedic']['mandi'] is None
    specs = charts.build_specs(r, 'en')
    assert sum(len(s['planets']) for s in specs['d1']['signs']) == 10
    from excel_generator import generate_excel
    from pdf_generator import generate_pdf
    assert generate_excel(r, 'en')[:2] == b'PK' and generate_pdf(r, 'en')[:5] == b'%PDF-'


def test_mandi_in_charts(res):
    specs = charts.build_specs(res, 'en')
    md = res['vedic']['mandi']
    assert 'Md' in specs['d1']['signs'][md['sign']]['planets']
    assert 'Md' in specs['d9']['signs'][md['navamsa']]['planets']
    assert 'Md' in specs['bhava']['houses'][md['bhava'] - 1]['planets']
    assert 'Md' in specs['alp']['signs'][md['sign']]['planets']
    assert all('Md' not in s['planets'] for s in specs['kp']['signs'])
    assert all('Md' not in h['planets'] for h in specs['kp']['houses'])
    ta = charts.build_specs(res, 'ta')
    assert 'மா' in ta['d1']['signs'][md['sign']]['planets']


# ── divisional charts ─────────────────────────────────────────────────────────

S = {n: i for i, n in enumerate(A.RASIS)}          # sign name -> index
MESHA, RISHABHA, MITHUNA, KATAKA = 0.0, 30.0, 60.0, 90.0     # start of an odd/movable, even/fixed, dual, water sign


def vs(sign_start, deg, key):
    return A.RASIS[A.varga_sign(sign_start + deg, key)]


def test_varga_list():
    assert A.VARGA_KEYS == ['D1', 'D2', 'D3', 'D4', 'D7', 'D9', 'D10', 'D12', 'D16', 'D20', 'D24', 'D27',
                            'D30', 'D40', 'D45', 'D60']
    for lang in ('en', 'ta', 'hi'):
        assert set(i18n.VARGA_NAMES[lang]) == set(A.VARGA_KEYS) == set(i18n.VARGA_PURPOSE[lang])


@pytest.mark.parametrize('key,start,first,last', [
    # first and last part of an odd sign (Mesha) and an even sign (Rishabha)
    ('D1', MESHA, 'Mesha', 'Mesha'), ('D1', RISHABHA, 'Rishabha', 'Rishabha'),
    ('D2', MESHA, 'Simha', 'Kataka'), ('D2', RISHABHA, 'Kataka', 'Simha'),
    ('D3', MESHA, 'Mesha', 'Dhanu'), ('D3', RISHABHA, 'Rishabha', 'Makara'),
    ('D4', MESHA, 'Mesha', 'Makara'), ('D4', RISHABHA, 'Rishabha', 'Kumbha'),
    ('D7', MESHA, 'Mesha', 'Thula'), ('D7', RISHABHA, 'Vrischika', 'Rishabha'),
    ('D10', MESHA, 'Mesha', 'Makara'), ('D10', RISHABHA, 'Makara', 'Thula'),
    ('D12', MESHA, 'Mesha', 'Meena'), ('D12', RISHABHA, 'Rishabha', 'Mesha'),
    ('D24', MESHA, 'Simha', 'Kataka'), ('D24', RISHABHA, 'Kataka', 'Mithuna'),
    ('D40', MESHA, 'Mesha', 'Kataka'), ('D40', RISHABHA, 'Thula', 'Makara'),
    ('D60', MESHA, 'Mesha', 'Meena'), ('D60', RISHABHA, 'Rishabha', 'Mesha'),
    # fire / earth / air / water signs
    ('D9', MESHA, 'Mesha', 'Dhanu'), ('D9', RISHABHA, 'Makara', 'Kanya'),
    ('D9', MITHUNA, 'Thula', 'Mithuna'), ('D9', KATAKA, 'Kataka', 'Meena'),
    ('D27', MESHA, 'Mesha', 'Mithuna'), ('D27', RISHABHA, 'Kataka', 'Kanya'),
    ('D27', MITHUNA, 'Thula', 'Dhanu'), ('D27', KATAKA, 'Makara', 'Meena'),
    # movable (Mesha) / fixed (Rishabha) / dual (Mithuna) signs
    ('D16', MESHA, 'Mesha', 'Kataka'), ('D16', RISHABHA, 'Simha', 'Vrischika'), ('D16', MITHUNA, 'Dhanu', 'Meena'),
    ('D20', MESHA, 'Mesha', 'Vrischika'), ('D20', RISHABHA, 'Dhanu', 'Kataka'), ('D20', MITHUNA, 'Simha', 'Meena'),
    ('D45', MESHA, 'Mesha', 'Dhanu'), ('D45', RISHABHA, 'Simha', 'Mesha'), ('D45', MITHUNA, 'Dhanu', 'Simha'),
])
def test_varga_first_and_last_part(key, start, first, last):
    assert vs(start, 0.001, key) == first
    assert vs(start, 29.999, key) == last


def test_trimsamsa_boundaries():
    odd = [(0.0, 'Mesha'), (4.999, 'Mesha'), (5.0, 'Kumbha'), (9.999, 'Kumbha'), (10.0, 'Dhanu'),
           (17.999, 'Dhanu'), (18.0, 'Mithuna'), (24.999, 'Mithuna'), (25.0, 'Thula'), (29.999, 'Thula')]
    even = [(0.0, 'Rishabha'), (4.999, 'Rishabha'), (5.0, 'Kanya'), (11.999, 'Kanya'), (12.0, 'Meena'),
            (19.999, 'Meena'), (20.0, 'Makara'), (24.999, 'Makara'), (25.0, 'Vrischika'), (29.999, 'Vrischika')]
    for start in (MESHA, 120.0, 300.0):                  # Mesha, Simha, Kumbha: odd signs
        assert [vs(start, d, 'D30') for d, _ in odd] == [s for _, s in odd]
    for start in (RISHABHA, 150.0, 330.0):               # Rishabha, Kanya, Meena: even signs
        assert [vs(start, d, 'D30') for d, _ in even] == [s for _, s in even]


def test_varga_exact_edges_fall_in_the_next_part():
    assert vs(MESHA, 15.0, 'D2') == 'Kataka' and vs(MESHA, 10.0, 'D3') == 'Simha'
    assert vs(MESHA, 7.5, 'D4') == 'Kataka' and vs(MESHA, 3.0, 'D10') == 'Rishabha'
    assert vs(MESHA, 2.5, 'D12') == 'Rishabha' and vs(MESHA, 0.5, 'D60') == 'Rishabha'
    assert vs(MESHA, 30 / 9, 'D9') == 'Rishabha' and vs(RISHABHA, 0.0, 'D9') == 'Makara'
    assert A.varga_sign(360.0, 'D60') == S['Mesha'] and A.varga_sign(-0.25, 'D60') == S['Kumbha']


@pytest.mark.parametrize('key,div', [(k, d) for k, d, _, _ in A.VARGAS])
def test_varga_covers_the_zodiac_without_gaps(key, div):
    # scan the zodiac in 15-arc-second steps and count the parts met
    parts_per_sign = 5 if key == 'D30' else div
    seen, prev, order = set(), None, []
    for i in range(360 * 240):
        lon = (i + 0.5) / 240.0
        cell = (A.get_rasi(lon), A.varga_part(lon, key))
        assert 0 <= A.varga_sign(lon, key) <= 11
        if cell != prev:
            assert cell not in seen                       # a part is one unbroken stretch
            seen.add(cell)
            order.append(cell)
            prev = cell
    assert len(seen) == 12 * parts_per_sign
    for sign in range(12):
        assert [p for s, p in order if s == sign] == list(range(parts_per_sign))


def test_d9_matches_navamsa_everywhere(res):
    for i in range(360 * 60):
        lon = (i + 0.5) / 60.0
        assert A.varga_sign(lon, 'D9') == A.navamsa_sign(lon)
    d9 = next(x for x in res['vedic']['vargas'] if x['key'] == 'D9')
    for p in res['vedic']['planets'] + [res['vedic']['mandi']]:
        assert d9['signs'][p['name']] == p['navamsa']


def test_vargas_for_the_sample_chart(res):
    v = res['vedic']
    names = ['Lagna'] + A.PLANETS + ['Mandi']
    assert [x['key'] for x in v['vargas']] == A.VARGA_KEYS
    for x in v['vargas']:
        assert list(x['signs']) == names                  # every point exactly once
        assert x['lagna_sign'] == x['signs']['Lagna']
    by = {x['key']: x['signs'] for x in v['vargas']}
    assert all(by['D1'][p['name']] == p['sign'] for p in v['planets'])
    # the Sun at Kumbha 21°13'06", worked by hand from the rules
    sun = {k: A.RASIS[by[k]['Sun']] for k in A.VARGA_KEYS}
    assert sun == {'D1': 'Kumbha', 'D2': 'Kataka', 'D3': 'Thula', 'D4': 'Simha', 'D7': 'Mithuna', 'D9': 'Mesha',
                   'D10': 'Kanya', 'D12': 'Thula', 'D16': 'Kataka', 'D20': 'Kumbha', 'D24': 'Dhanu',
                   'D27': 'Rishabha', 'D30': 'Mithuna', 'D40': 'Simha', 'D45': 'Meena', 'D60': 'Simha'}
    assert set(by['D2'].values()) <= {S['Kataka'], S['Simha']}


def test_varga_specs_and_table(res):
    specs = charts.build_varga_specs(res, 'en')
    assert [s['key'] for s in specs] == A.VARGA_KEYS
    for s in specs:
        assert sum(len(c['planets']) for c in s['signs']) == 11 == sum(len(h['planets']) for h in s['houses'])
        assert s['purpose'] and s['title'].startswith(s['key'] + ' ')
        assert s['signs'][s['lagna_sign']]['lagna'] and 'La' in s['signs'][s['lagna_sign']]['planets']
    assert specs[4]['purpose'] == 'Children and progeny'
    keys, rows = charts.varga_table(res, 'en')
    assert keys == A.VARGA_KEYS and len(rows) == 11 and all(len(r['signs']) == 16 for r in rows)
    assert [r['name'] for r in rows if r['vargottama']] == ['Rahu', 'Ketu']
    ta = charts.build_varga_specs(res, 'ta')
    assert ta[5]['title'] == 'D9 நவாம்சம்' and ta[1]['purpose'] == 'செல்வம், பொருளாதாரம்'


def test_vargas_do_not_change_other_results(res):
    v = res['vedic']
    assert [p['name'] for p in v['planets']] == ['Lagna'] + A.PLANETS
    assert v['dasa']['balance']['lord'] == 'Mars' and v['mandi']['ghatis'] == 22
    assert len(charts.build_specs(res, 'en')) == 5


def test_excel_divisional_sheet():
    import openpyxl
    from excel_generator import generate_excel, COL0, RIGHT0
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090, now=NOW)
    wb = openpyxl.load_workbook(io.BytesIO(generate_excel(r, 'en')))
    assert wb.sheetnames == ['Summary', 'Vedic', 'Divisional Charts', 'KP', 'ALP', 'Dasa', 'Notes']
    ws = wb['Divisional Charts']
    head = next(c.row for c in ws['B'] if c.value == 'Planet')
    heads = [c.value for c in ws[head] if c.value is not None]
    assert heads[1:17] == A.VARGA_KEYS and heads[17] == 'Vargottama'
    sun = next(c.row for c in ws['B'] if c.value == 'Sun')
    vals = [c.value for c in ws[sun] if c.value is not None]
    assert vals[1:4] == ['Kum', 'Kat', 'Thu'] and vals[16] == 'Sim' and vals[17] == '—'
    rahu = next(c.row for c in ws['B'] if c.value == 'Rahu')
    assert [c.value for c in ws[rahu] if c.value is not None][17] == 'Yes'
    # 16 charts, each with its purpose line under the title
    titles = {}
    for col in (COL0, RIGHT0):
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row, col).value
            if isinstance(val, str) and val.split(' ')[0] in A.VARGA_KEYS and val != 'D1':
                if ws.cell(row + 1, col).value in i18n.VARGA_PURPOSE['en'].values():
                    titles[val.split(' ')[0]] = (row, col)
    assert sorted(titles, key=A.VARGA_KEYS.index) == A.VARGA_KEYS
    assert ws.cell(titles['D10'][0] + 1, titles['D10'][1]).value == 'Career, status and achievements'

    def block(key, sign):
        row, col = titles[key]
        gr, gc = charts.SOUTH_POS[sign]
        return ws.cell(row + 2 + gr * 3 + 1, col + gc * 3).value or ''
    assert 'Su' in block('D3', S['Thula']).split()           # Sun: Thula in D3
    assert 'Su' in block('D10', S['Kanya']).split()          # Sun: Kanya in D10
    assert 'Md' in block('D60', S['Mithuna']).split()        # Mandi: Mithuna in D60
    assert len(ws.row_breaks.brk) == 4                       # charts are not split across pages


@pytest.mark.parametrize('lang', ['en', 'ta'])
def test_pdf_has_divisional_section(lang, tmp_path):
    import shutil
    import subprocess
    if not shutil.which('pdftotext'):
        pytest.skip('pdftotext not installed')
    from pdf_generator import generate_pdf
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090, lang=lang, now=NOW)
    f = tmp_path / 'r.pdf'
    f.write_bytes(generate_pdf(r, lang))
    text = subprocess.run(['pdftotext', '-layout', str(f), '-'], capture_output=True, text=True).stdout
    for key in A.VARGA_KEYS:
        assert sum(1 for line in text.splitlines() if key + ' ' in line) >= 1, key
    if lang == 'en':
        for purpose in i18n.VARGA_PURPOSE['en'].values():
            assert purpose in text


# ── charts ────────────────────────────────────────────────────────────────────

def test_chart_specs(res):
    specs = charts.build_specs(res, 'en')
    assert set(specs) == {'d1', 'd9', 'bhava', 'kp', 'alp'}
    d1 = specs['d1']
    lagna_sign = res['vedic']['lagna_sign']
    cell = d1['signs'][lagna_sign]
    assert cell['lagna'] and 'La' in cell['planets'] and cell['tag'] == '1'
    # lagna + 9 planets + Mandi
    assert sum(len(s['planets']) for s in d1['signs']) == 11
    assert sum(len(h['planets']) for h in d1['houses']) == 11
    assert d1['houses'][0]['sign_num'] == lagna_sign + 1 and 'La' in d1['houses'][0]['planets']
    assert sum(len(h['planets']) for h in specs['kp']['houses']) == 9
    tags = ' '.join(s['tag'] for s in specs['kp']['signs']).split()
    assert sorted(tags) == sorted(charts.ROMAN)
    assert 'AL' in specs['alp']['signs'][res['alp']['lagna_sign']]['planets']


def test_south_layout_is_the_standard_one():
    assert charts.SOUTH_POS[11] == (0, 0) and charts.SOUTH_POS[0] == (0, 1)     # Meena top-left, then Mesha
    assert charts.SOUTH_POS[3] == (1, 3) and charts.SOUTH_POS[8] == (3, 0)
    assert len(set(charts.SOUTH_POS.values())) == 12


# ── reports ───────────────────────────────────────────────────────────────────

@pytest.mark.parametrize('lang,style', [('en', 'south'), ('ta', 'south'), ('bi', 'north'), ('hi', 'north'),
                                        ('te', 'south'), ('kn', 'south'), ('ml', 'south'), ('mr', 'south'),
                                        ('bn', 'south')])
def test_excel_workbook(lang, style):
    import openpyxl
    from excel_generator import generate_excel
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090,
                  lang=lang, chart_style=style, now=NOW)
    wb = openpyxl.load_workbook(io.BytesIO(generate_excel(r, lang)))
    assert len(wb.sheetnames) == 7
    for ws in wb:
        assert ws.max_row > 5
        assert str(ws.page_setup.paperSize) == '9' and ws.page_setup.orientation == 'portrait'    # 9 = A4
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not c.value.startswith('#') and 'None' not in c.value and 'TODO' not in c.value
    dasa = wb.worksheets[5]
    assert isinstance(dasa['E4'].value, datetime) and dasa['E4'].number_format == 'dd-mm-yyyy'
    assert dasa.freeze_panes == 'A4' and dasa.auto_filter.ref.startswith('B3')
    rows = dasa.max_row - 3
    assert 600 < rows <= 729


def test_excel_chart_places_planets_by_sign():
    import openpyxl
    from excel_generator import generate_excel, COL0
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090, now=NOW)
    ws = openpyxl.load_workbook(io.BytesIO(generate_excel(r, 'en'))).worksheets[1]      # Vedic sheet
    top = next(c.row for c in ws['B'] if c.value == 'Rasi Chart (D1)') + 1
    def block(sign):
        gr, gc = charts.SOUTH_POS[sign]
        r0, c0 = top + gr * 3, COL0 + gc * 3
        return ws.cell(r0, c0).value, ws.cell(r0 + 1, c0).value
    label, planets = block(7)                    # Vrischika: lagna
    assert label.startswith('Vri') and label.endswith('1') and planets == 'La'
    assert block(10)[1].split() == ['Su', 'Ju', 'Ke']     # Kumbha
    assert block(1)[1] == 'Mo' and block(9)[1] == 'Ve'    # Rishabha, Makara
    assert block(11)[0].startswith('Mee')                 # Meena is the top-left cell
    md = r['vedic']['mandi']
    assert A.RASIS[md['sign']] == 'Dhanu' and block(8)[1] == 'Md'      # Mandi sits in its own sign
    # Planet Positions table: Mandi is the row after Ketu, with a dash for retro / combust / dignity
    ketu = next(c.row for c in ws['B'] if c.value == 'Ketu')
    assert ws.cell(ketu + 1, 2).value == 'Mandi'
    row = [c.value for c in ws[ketu + 1] if c.value is not None]
    assert row[0] == 'Mandi' and row[1] == 'Dhanu' and row[2] == md['dms'] and row[-1] == '—'
    assert (top, COL0) == (top + charts.SOUTH_POS[11][0] * 3, COL0 + charts.SOUTH_POS[11][1] * 3)


@pytest.mark.parametrize('lang', ['en', 'ta', 'bi', 'hi', 'te', 'kn', 'ml', 'mr', 'bn'])
def test_pdf_report(lang, recwarn):
    from pdf_generator import generate_pdf
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090,
                  lang=lang, chart_style='north' if lang in ('bi', 'hi') else 'south', now=NOW)
    data = generate_pdf(r, lang)
    assert data[:5] == b'%PDF-' and len(data) > 40000
    missing = [str(w.message) for w in recwarn if 'missing' in str(w.message).lower()]
    assert not missing, missing


@pytest.mark.parametrize('lang', ['ta', 'hi', 'te', 'bi'])
def test_pdf_numbers_survive_next_to_indian_script(lang, tmp_path):
    """
    Digits and degree values printed after a cell in an Indian script must
    still read as digits (they once came out as stray letters).
    """
    import shutil
    import subprocess
    if not shutil.which('pdftotext'):
        pytest.skip('pdftotext not installed')
    from pdf_generator import generate_pdf
    r = A.compute('Sample Native', '1998-03-06', '01:37', 'Delhi', lat=28.6139, lon=77.2090, lang=lang, now=NOW)
    f = tmp_path / 'r.pdf'
    f.write_bytes(generate_pdf(r, lang))
    text = subprocess.run(['pdftotext', '-layout', str(f), '-'], capture_output=True, text=True).stdout
    flat = text.replace(' ', '')
    for p in r['vedic']['planets'] + [r['vedic']['mandi']] + r['kp']['cusps']:
        assert p['dms'].replace(' ', '') in flat, (lang, p['name'], p['dms'])
    for d in r['vedic']['dasa']['dasas']:
        assert d['end'].strftime('%d-%m-%Y') in flat, (lang, d['lord'])
    for s in r['alp']['pada_periods'][:20]:
        assert s['end'].strftime('%d-%m-%Y') in flat


def test_labels_complete():
    keys = set(i18n.LABELS['en'])
    for lang in ('ta', 'hi'):
        assert set(i18n.LABELS[lang]) == keys, lang
    for lang in ('en', 'ta', 'hi', 'te', 'kn', 'ml', 'mr', 'bn'):
        assert len(i18n.RASIS[lang]) == 12 and len(i18n.PLANETS[lang]) == 11
        # every chart abbreviation, Mandi and the ALP marker included, is distinct
        assert len(set(i18n.ABBR[lang][p] for p in A.PLANETS + ['Lagna', 'Mandi', 'ALP'])) == 12, lang


# ── API ───────────────────────────────────────────────────────────────────────

@pytest.fixture(scope='module')
def client():
    import app as app_module
    return app_module.app.test_client()


BODY = {'name': 'Sample Native', 'dob': '1998-03-06', 'tob': '01:37', 'pob': 'Delhi',
        'lat': 28.6139, 'lon': 77.2090, 'lang': 'en', 'chartStyle': 'south', 'ayanamsha': 'lahiri'}


def test_api_horoscope(client):
    rv = client.post('/api/horoscope', json=BODY)
    assert rv.status_code == 200
    j = rv.get_json()
    assert set(j) >= {'meta', 'vedic', 'kp', 'alp', 'charts'}
    assert j['meta']['utc_offset_str'] == 'UTC+05:30' and j['meta']['tz'] == 'Asia/Kolkata'
    assert j['vedic']['planets'][0]['name'] == 'Lagna'
    md = j['vedic']['mandi']
    assert md['name'] == 'Mandi' and md['rise_time'].startswith('1998-03-06T03:24') and md['ghatis'] == 22
    assert j['names']['planets']['Mandi'] == 'Mandi' and j['names']['abbr']['Mandi'] == 'Md'
    assert len(j['planets']) == 10                         # legacy list is unchanged
    assert [x['key'] for x in j['vedic']['vargas']] == A.VARGA_KEYS
    assert len(j['varga_charts']) == 16 and j['varga_charts'][1]['purpose'] == 'Wealth and resources'
    assert j['varga_table']['keys'] == A.VARGA_KEYS and len(j['varga_table']['rows']) == 11


def test_api_ayanamsha_option_is_used(client):
    a = client.post('/api/horoscope', json=BODY).get_json()
    b = client.post('/api/horoscope', json=dict(BODY, ayanamsha='raman')).get_json()
    assert abs((a['vedic']['ayanamsha'] - b['vedic']['ayanamsha']) - 1.446) < 0.01
    assert a['kp']['ayanamsha'] == b['kp']['ayanamsha']


def test_api_downloads(client):
    rv = client.post('/api/download/excel', json=BODY)
    assert rv.status_code == 200 and rv.data[:2] == b'PK'
    assert 'HoroscopeGen_Sample_Native.xlsx' in rv.headers['Content-Disposition']
    rv = client.post('/api/download/pdf', json=BODY)
    assert rv.status_code == 200 and rv.data[:5] == b'%PDF-'
    assert 'HoroscopeGen_Sample_Native.pdf' in rv.headers['Content-Disposition']


def test_api_errors(client):
    assert client.post('/api/horoscope', json={'name': 'X'}).status_code == 400
    assert client.post('/api/horoscope', json=dict(BODY, dob='bad')).status_code == 400
    assert client.post('/api/horoscope', json=dict(BODY, lat='abc')).status_code == 400
    assert client.get('/health').get_json() == {'status': 'ok'}
    assert client.get('/api/verify-payment').status_code == 404
