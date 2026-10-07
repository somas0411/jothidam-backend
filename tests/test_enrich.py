"""
Tests for enrich.py (derived facts for AI readers), the AI_ sheets of the
workbook and the JSON routes.

Reference values that do not come from the code under test:
  * Ashtakavarga totals 48, 49, 39, 54, 56, 52, 39 and 337 (fixed for every chart).
  * Chart S (04-11-1984 00:20 Chennai): Vimsopaka figures, lagna margins,
    yogas and transit dates worked out in an earlier session and by hand.
  * tests/golden_main.json: fingerprints of the seven report sheets and of
    the /api/horoscope response as produced by `main` before this change
    (see tests/make_golden.py).
  * A transit search written here with plain one-day steps.
"""
import functools
import hashlib
import io
import json
import random
import sys
from datetime import date, datetime, timedelta
from itertools import combinations
from pathlib import Path

import pytest
import swisseph as swe

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import astro_engine as A          # noqa: E402
import enrich as E                # noqa: E402
import make_golden as G           # noqa: E402

NOW = G.NOW                                           # 07-10-2026 17:00
S = dict(name='Chart S', dob='1984-11-04', tob='00:20', pob='Chennai', lat=13.0827, lon=80.2707)
CHART_A = dict(name='Chart A', dob='1990-08-15', tob='10:30', pob='Chennai', lat=13.0827, lon=80.2707)
CHART_B = dict(name='Chart B', dob='1985-12-25', tob='21:45', pob='Madurai', lat=9.9252, lon=78.1198)
GOLDEN = json.loads((Path(__file__).parent / 'golden_main.json').read_text(encoding='utf-8'))
R = {n: i for i, n in enumerate(A.RASIS)}
IST = timedelta(hours=5, minutes=30)


def chart(lagna, **signs):
    """A minimal compute() result with points placed by sign name (15 degrees in), for rule tests."""
    place = dict(Sun='Mesha', Moon='Mesha', Mars='Mesha', Mercury='Mesha', Jupiter='Mesha', Venus='Mesha',
                 Saturn='Mesha', Rahu='Mithuna', Ketu='Dhanu')
    place.update(signs)
    lons = {'Lagna': R[lagna] * 30 + 15.0}
    for name, where in place.items():
        sign, deg = (where, 15.0) if isinstance(where, str) else where
        lons[name] = R[sign] * 30 + deg
    lagna_sign = R[lagna]
    planets = []
    for name in ['Lagna'] + A.PLANETS:
        p = A._point(name, lons[name], 1.0 if name in A.PLANETS[:7] else 0.0)
        p.update(house=(p['sign'] - lagna_sign) % 12 + 1, bhava=(p['sign'] - lagna_sign) % 12 + 1,
                 navamsa=A.navamsa_sign(lons[name]), dignity=A.dignity_of(name, p['sign']),
                 combust=A.is_combust(name, lons[name], 1.0, lons['Sun']))
        planets.append(p)
    birth = datetime(2000, 1, 1, 12, 0)
    cusps = []
    for i in range(12):
        row = A._point(f'Cusp {i + 1}', (lagna_sign + i) * 30 + 15.0)
        row['house'] = i + 1
        cusps.append(row)
    kp_planets = [dict(p, house=p['house']) for p in planets[1:]]
    return {
        'meta': {'local_dt': birth, 'ut_dt': birth - IST, 'generated_at': NOW, 'lat': 13.0, 'lon': 80.0,
                 'ayanamsha_key': 'lahiri', 'node': 'mean', 'gender': ''},
        'vedic': {'planets': planets, 'lagna_sign': lagna_sign, 'mandi': None, 'vargas': A.build_vargas(planets)},
        'kp': {'planets': kp_planets, 'cusps': cusps, 'house_system': 'Placidus',
               'planet_significators': [{'planet': p['name'], 'all': [p['house']]} for p in kp_planets]},
    }


def ctx(lagna, **signs):
    return E.Ctx(chart(lagna, **signs))


def yogas(lagna, **signs):
    return {y['yoga']: y for y in E._yogas(ctx(lagna, **signs))}


@pytest.fixture(scope='module')
def res():
    return A.compute(now=NOW, **S)


@pytest.fixture(scope='module')
def ai(res):
    return E.enrich(res)


@pytest.fixture(scope='module')
def T(ai):
    return ai['tables']


@pytest.fixture(scope='module')
def workbooks():
    """The workbook with AI sheets for chart S in English and Tamil: {lang: (bytes, openpyxl workbook)}."""
    import openpyxl
    from excel_generator import generate_excel
    out = {}
    for lang in ('en', 'ta'):
        r = A.compute(lang=lang, now=NOW, **S)
        data = generate_excel(r, lang, ai=E.enrich(r))
        out[lang] = (data, openpyxl.load_workbook(io.BytesIO(data)))
    return out


def by(rows, key):
    return {r[key]: r for r in rows}


# ── rule tables ───────────────────────────────────────────────────────────────

def test_ashtakavarga_tables_add_up():
    for planet, donors in E.BAV_PLACES.items():
        assert list(donors) == E.SEVEN + ['Lagna']
        assert sum(len(v) for v in donors.values()) == E.BAV_TOTALS[planet]
        assert all(1 <= p <= 12 and len(set(v)) == len(v) for v in donors.values() for p in v)
    assert sum(E.BAV_TOTALS.values()) == E.SAV_TOTAL == 337
    assert [E.BAV_TOTALS[p] for p in E.SEVEN] == [48, 49, 39, 54, 56, 52, 39]


def test_ashtakavarga_totals_on_reference_and_random_charts():
    rnd = random.Random(1984)
    births = [S, CHART_A, CHART_B]
    for i in range(50):
        d = datetime(1930, 1, 1) + timedelta(minutes=rnd.randrange(0, 95 * 525960))
        births.append(dict(name=f'R{i}', dob=d.strftime('%Y-%m-%d'), tob=d.strftime('%H:%M'), pob='x',
                           lat=rnd.uniform(-50, 60), lon=rnd.uniform(-150, 150), utc_offset=0))
    for b in births:
        r = A.compute(now=NOW, **b)
        sign = {p['name']: p['sign'] for p in r['vedic']['planets']}
        bav, sav = E.ashtakavarga(sign)
        assert [sum(bav[p]) for p in E.SEVEN] == [48, 49, 39, 54, 56, 52, 39], b['name']
        assert sum(sav) == 337 and all(0 <= x <= 8 for p in E.SEVEN for x in bav[p])
        assert sav == [sum(bav[p][s] for p in E.SEVEN) for s in range(12)]


def test_ashtakavarga_counts_places_from_each_donor():
    # everything in Mesha: a bindu falls in the n-th sign once for every donor that lists place n
    sign = {n: 0 for n in E.SEVEN + ['Lagna']}
    bav, _ = E.ashtakavarga(sign)
    for planet in E.SEVEN:
        expect = [sum(1 for places in E.BAV_PLACES[planet].values() if s + 1 in places) for s in range(12)]
        assert bav[planet] == expect
    # moving one donor moves only its own bindus, by the same number of signs
    moved, _ = E.ashtakavarga(dict(sign, Saturn=3))
    assert sum(moved['Sun']) == 48 and moved['Sun'] != bav['Sun']
    for s in range(12):
        given_by_saturn = 1 if s + 1 in E.BAV_PLACES['Sun']['Saturn'] else 0
        after = 1 if (s - 3) % 12 + 1 in E.BAV_PLACES['Sun']['Saturn'] else 0
        assert moved['Sun'][s] == bav['Sun'][s] - given_by_saturn + after


def test_ashtakavarga_chart_s(T):
    rows = T['AI_Ashtakavarga']
    assert len(rows) == 8 * 13
    totals = {r['planet']: r['bindus'] for r in rows if r['row_type'] == 'total'}
    assert totals == {'Sun': 48, 'Moon': 49, 'Mars': 39, 'Mercury': 54, 'Jupiter': 56, 'Venus': 52, 'Saturn': 39,
                      'Sarvashtakavarga': 337}
    # checked against PyJHora's routine run with the same (BPHS) tables
    sav = [r['bindus'] for r in rows if r['planet'] == 'Sarvashtakavarga' and r['row_type'] == 'sign']
    assert sav == [34, 28, 21, 36, 32, 33, 33, 21, 30, 23, 19, 27]
    sun = [r['bindus'] for r in rows if r['planet'] == 'Sun' and r['row_type'] == 'sign']
    assert sun == [7, 5, 4, 5, 5, 3, 7, 3, 3, 4, 0, 2]
    kataka = next(r for r in rows if r['planet'] == 'Sun' and r['sign'] == 'Kataka')
    assert kataka['house'] == 1 and kataka['sign_no'] == 4               # Kataka lagna
    houses = by(T['AI_Houses'], 'house')
    assert [houses[h]['sav_bindus'] for h in range(1, 13)] == sav[3:] + sav[:3]


def test_natural_and_compound_friendship():
    assert E.natural_relation('Sun', 'Moon') == 'Friend' and E.natural_relation('Sun', 'Mercury') == 'Neutral'
    assert E.natural_relation('Sun', 'Saturn') == 'Enemy' and E.natural_relation('Moon', 'Saturn') == 'Neutral'
    assert E.natural_relation('Mercury', 'Moon') == 'Enemy' and E.natural_relation('Moon', 'Mercury') == 'Friend'
    assert all(len(set(f) | set(n) | set(e)) == 6 and len(f) + len(n) + len(e) == 6 for f, n, e in E.NATURAL.values())
    assert not E.NATURAL['Moon'][2]                                        # the Moon has no enemies
    # temporary: 2, 3, 4, 10, 11, 12 from a planet are friends
    assert [E.temporary_relation(0, s) for s in range(12)] == (
        ['Enemy', 'Friend', 'Friend', 'Friend'] + ['Enemy'] * 5 + ['Friend'] * 3)
    assert E.COMPOUND == {('Friend', 'Friend'): 'Great friend', ('Neutral', 'Friend'): 'Friend',
                          ('Enemy', 'Friend'): 'Neutral', ('Friend', 'Enemy'): 'Neutral',
                          ('Neutral', 'Enemy'): 'Enemy', ('Enemy', 'Enemy'): 'Great enemy'}


def test_dignity_ranges_in_d1_and_by_sign_elsewhere():
    comp = {p: {q: 'Neutral' for q in E.SEVEN if q != p} for p in E.SEVEN}
    d = lambda p, sign, deg, key='D1': E.dignity_in(p, R[sign], key, deg, comp)
    assert d('Moon', 'Rishabha', 2.99) == ('Exalted', False) and d('Moon', 'Rishabha', 3.0) == ('Moolatrikona', False)
    assert d('Mercury', 'Kanya', 14.99) == ('Exalted', False) and d('Mercury', 'Kanya', 15.0) == ('Moolatrikona', False)
    assert d('Mercury', 'Kanya', 19.99) == ('Moolatrikona', False) and d('Mercury', 'Kanya', 20.0) == ('Own sign', False)
    assert d('Sun', 'Simha', 19.99) == ('Moolatrikona', False) and d('Sun', 'Simha', 20.0) == ('Own sign', False)
    assert d('Mars', 'Mesha', 11.99)[0] == 'Moolatrikona' and d('Mars', 'Mesha', 12.0)[0] == 'Own sign'
    assert d('Jupiter', 'Dhanu', 9.99)[0] == 'Moolatrikona' and d('Jupiter', 'Dhanu', 10.0)[0] == 'Own sign'
    assert d('Venus', 'Thula', 14.99)[0] == 'Moolatrikona' and d('Venus', 'Thula', 15.0)[0] == 'Own sign'
    assert d('Saturn', 'Kumbha', 19.99)[0] == 'Moolatrikona' and d('Saturn', 'Kumbha', 20.0)[0] == 'Own sign'
    assert d('Saturn', 'Thula', 5.0) == ('Exalted', False) and d('Sun', 'Mesha', 29.0) == ('Exalted', False)
    # a debilitated planet carries the relationship with its sign lord, plus the flag
    assert d('Sun', 'Thula', 5.0) == ('Neutral', True) and d('Saturn', 'Mesha', 5.0) == ('Neutral', True)
    # other charts have no degrees: whole signs, no moolatrikona
    assert d('Moon', 'Rishabha', 20.0, 'D9') == ('Exalted', False)
    assert d('Mercury', 'Kanya', 25.0, 'D9') == ('Exalted', False) and d('Sun', 'Simha', 5.0, 'D10') == ('Own sign', False)
    for p, s in A.EXALT_SIGN.items():
        assert E.dignity_in(p, (s + 6) % 12, 'D9', 0.0, comp)[1] is True


@pytest.mark.parametrize('lagna,expect', [
    ('Kataka', dict(Moon='Benefic', Sun='Neutral', Mercury='Malefic', Venus='Malefic', Mars='Yogakaraka',
                    Jupiter='Benefic', Saturn='Malefic')),
    ('Rishabha', dict(Saturn='Yogakaraka', Venus='Benefic', Mercury='Benefic', Moon='Malefic', Sun='Neutral',
                      Jupiter='Malefic', Mars='Neutral')),
    ('Simha', dict(Mars='Yogakaraka', Sun='Benefic', Jupiter='Benefic', Saturn='Malefic', Venus='Malefic',
                   Mercury='Malefic', Moon='Neutral')),
    ('Dhanu', dict(Moon='Neutral')),                    # the Moon owns only the 8th: exempt from the 8th-lord rule
    ('Makara', dict(Venus='Yogakaraka', Sun='Neutral')),
])
def test_functional_nature(lagna, expect):
    c = ctx(lagna)
    for planet, label in expect.items():
        assert c.nature[planet] == label, (lagna, planet, c.owned[planet])


def test_lordship_badhaka_and_aspects():
    c = ctx('Kataka', Mars='Dhanu', Jupiter='Dhanu', Saturn='Thula', Sun='Thula', Moon='Kumbha')
    assert c.owned['Mars'] == [5, 10] and c.owned['Saturn'] == [7, 8] and c.owned['Rahu'] == []
    assert sorted(h for p in E.SEVEN for h in c.owned[p]) == list(range(1, 13))
    assert c.asp_houses['Mars'] == [1, 9, 12] and c.asp_houses['Jupiter'] == [2, 10, 12]
    assert c.asp_houses['Saturn'] == [1, 6, 10] and c.asp_houses['Sun'] == [10]
    assert c.asp_houses['Rahu'] == [6] and c.asp_houses['Ketu'] == [12]            # 7th only
    assert ctx('Mesha').badhaka_house == 11 and ctx('Rishabha').badhaka_house == 9 and ctx('Mithuna').badhaka_house == 7
    assert ctx('Kataka').badhaka_house == 11 and ctx('Meena').badhaka_house == 7


def test_pairs_conjunction_aspect_exchange_and_war():
    c = ctx('Mesha', Mars=('Thula', 10.0), Venus=('Mesha', 10.6), Saturn=('Mesha', 11.5), Sun='Kanya',
            Mercury='Kanya', Jupiter='Kataka', Moon='Makara')
    pairs = {(r['planet_a'], r['planet_b']): r for r in E._pairs(c)}
    assert len(pairs) == 36
    mv = pairs[('Mars', 'Venus')]
    assert mv['exchange_of_signs'] == 'yes' and mv['mutual_aspect'] == 'yes' and mv['same_sign'] == 'no'
    assert mv['separation_deg'] == pytest.approx(179.4) and mv['planetary_war'] == 'no'
    vs = pairs[('Venus', 'Saturn')]
    assert vs['same_sign'] == 'yes' and vs['planetary_war'] == 'yes' and vs['separation_deg'] == pytest.approx(0.9)
    assert pairs[('Sun', 'Mercury')]['planetary_war'] == 'no'             # the Sun takes no part in a war
    jm = pairs[('Moon', 'Jupiter')]
    assert jm['mutual_aspect'] == 'yes' and jm['b_counted_from_a'] == 7
    ms = pairs[('Mars', 'Saturn')]                                        # Mars 7th on Saturn, Saturn 7th on Mars
    assert ms['a_aspects_b'] == 'yes' and ms['b_aspects_a'] == 'yes'
    rk = pairs[('Rahu', 'Ketu')]
    assert rk['natural_a_to_b'] is None and rk['compound_a_to_b'] is None and rk['mutual_aspect'] == 'yes'
    sm = pairs[('Sun', 'Moon')]                                           # Makara is 5th from Kanya
    assert (sm['natural_a_to_b'], sm['temporary_a_to_b'], sm['compound_a_to_b']) == ('Friend', 'Enemy', 'Neutral')


# ── chart S: values known from outside the code ───────────────────────────────

def test_chart_s_positions(res):
    want = {'Lagna': ('Kataka', '23°37\'28"'), 'Sun': ('Thula', '17°53\'41"'), 'Moon': ('Kumbha', '23°45\'59"'),
            'Mars': ('Dhanu', '27°28\'11"'), 'Mercury': ('Vrischika', '02°20\'04"'),
            'Jupiter': ('Dhanu', '15°39\'56"'), 'Venus': ('Vrischika', '24°01\'29"'),
            'Saturn': ('Thula', '24°31\'55"'), 'Rahu': ('Rishabha', '04°35\'21"')}
    got = {p['name']: (A.RASIS[p['sign']], p['dms']) for p in res['vedic']['planets']}
    for name, value in want.items():
        assert got[name] == value
    md = res['vedic']['mandi']
    assert (A.RASIS[md['sign']], md['dms']) == ('Kataka', '11°29\'31"')


def test_chart_s_vimsopaka(T):
    got = {r['planet']: r['vimsopaka_shodasavarga'] for r in T['AI_Strength']}
    assert got == {'Jupiter': 17.45, 'Venus': 14.55, 'Saturn': 14.375, 'Mercury': 14.075, 'Moon': 13.10,
                   'Sun': 12.575, 'Mars': 12.55}
    for row in T['AI_Strength']:
        pts = {k: row[f'points_{k}'] for k in A.VARGA_KEYS}
        assert set(pts.values()) <= {20, 18, 15, 10, 7, 5}
        for scheme, weights in E.VIMSOPAKA_WEIGHTS.items():
            assert sum(weights.values()) == 20
            assert row[f'vimsopaka_{scheme}'] == pytest.approx(sum(w * pts[k] for k, w in weights.items()) / 20, abs=1e-9)
    planets = by(T['AI_Planets'], 'point')
    assert planets['Jupiter']['vimsopaka_shodasavarga'] == 17.45 and planets['Rahu']['vimsopaka_shodasavarga'] is None


def test_chart_s_lagna_sensitivity(T):
    rows = by(T['AI_VargaLagnas'], 'chart')
    known = {'D9': (74, 773), 'D10': (671, 95), 'D24': (287, 32), 'D60': (31, 95), 'D7': (561, 532)}
    for key, (earlier, later) in known.items():
        assert abs(rows[key]['seconds_earlier'] - earlier) <= 5, key
        assert abs(rows[key]['seconds_later'] - later) <= 5, key
    assert abs(rows['D1']['seconds_later'] - 1616) <= 5
    uncertain = [k for k in A.VARGA_KEYS if rows[k]['lagna_uncertain'] == 'yes']
    assert uncertain == ['D9', 'D10', 'D20', 'D24', 'D27', 'D40', 'D45', 'D60']
    for k in A.VARGA_KEYS:
        r = rows[k]
        assert (r['lagna_uncertain'] == 'yes') == (min(r['seconds_earlier'], r['seconds_later']) < 120)
        assert r['lagna_if_earlier'] != r['lagna_sign'] and r['lagna_if_later'] != r['lagna_sign']
    assert rows['D1']['lagna_if_earlier'] == 'Mithuna' and rows['D1']['lagna_if_later'] == 'Simha'


def test_lagna_margins_are_exact_to_the_second(res):
    """At the margin the real lagna is unchanged; one second further it is in the sign reported."""
    c = E.Ctx(res)
    margins = E.lagna_margins(c)
    sid = swe.SIDM_LAHIRI
    for key in A.VARGA_KEYS:
        now_sign = c.vargas[key]['lagna_sign']
        for direction in (-1, 1):
            secs, new_sign = margins[key][direction]
            at = lambda s: A.varga_sign(A._lagna_at(c.birth + timedelta(seconds=direction * s), c.offset, c.lat, c.lon, sid), key)
            assert at(secs) == now_sign and at(secs + 1) == new_sign != now_sign, (key, direction)
            # and it did not change anywhere before that (checked every 7 seconds)
            assert all(at(s) == now_sign for s in range(0, secs, 7)), (key, direction)


def test_chart_s_yogas(T):
    y = by(T['AI_Yogas'], 'yoga')
    found = sorted(n for n, r in y.items() if r['status'] == 'found')
    assert found == sorted([
        'Sasa yoga', 'Neecha bhanga of Sun - condition (a)', 'Neecha bhanga of Sun - condition (b)',
        'Raja yoga', 'Dhana yoga', 'Harsha yoga (Viparita raja yoga)', 'Vesi yoga', 'Amala yoga',
        'Kuja dosha from Venus'])
    assert [n for n, r in y.items() if r['status'] == 'cancelled'] == ['Kemadruma yoga']
    assert 'Formed: yes' in y['Kemadruma yoga']['note'] and 'Cancelled: yes' in y['Kemadruma yoga']['note']
    assert y['Sasa yoga']['planets'] == 'Saturn' and y['Sasa yoga']['houses'] == '4'
    assert 'its lord Venus is in a kendra from the Moon' in y['Neecha bhanga of Sun - condition (a)']['note']
    assert 'Saturn, which is exalted in Thula, is in a kendra from the lagna' in y['Neecha bhanga of Sun - condition (b)']['note']
    assert y['Raja yoga']['planets'] == 'Mars, Jupiter' and 'conjunction' in y['Raja yoga']['note']
    assert 'Mars (lord of 5, 10)' in y['Raja yoga']['note'] and 'Jupiter (lord of 6, 9)' in y['Raja yoga']['note']
    assert y['Harsha yoga (Viparita raja yoga)']['planets'] == 'Jupiter'
    for name in ('Gajakesari yoga', 'Budha-Aditya yoga', 'Kuja dosha from the lagna', 'Kuja dosha from the Moon',
                 'Sasa yoga - counted from the Moon (variant)', 'Kala Sarpa yoga (popular, not in BPHS)',
                 'Parivartana yoga', 'Adhi yoga', 'Sunapha yoga', 'Anapha yoga', 'Durudhara yoga'):
        assert y[name]['status'] == 'not found', name
    assert all(r['status'] in ('found', 'cancelled', 'not found') and r['rule'] for r in T['AI_Yogas'])
    assert len(y) == len(T['AI_Yogas'])                 # every row has its own name


SATURN_S = [('Makara', '24-01-2020 09:56'), ('Kumbha', '29-04-2022 07:53'), ('Makara', '12-07-2022 14:47'),
            ('Kumbha', '17-01-2023 18:04'), ('Meena', '29-03-2025 21:44'), ('Mesha', '03-06-2027 05:28'),
            ('Meena', '20-10-2027 07:12'), ('Mesha', '23-02-2028 19:24'), ('Rishabha', '08-08-2029 12:34')]
JUPITER_S = [('Rishabha', '01-05-2024 12:59'), ('Mithuna', '14-05-2025 22:36'), ('Kataka', '18-10-2025 19:47'),
             ('Mithuna', '05-12-2025 17:25'), ('Kataka', '02-06-2026 01:50')]
RAHU_S = [('Meena', '30-10-2023 16:36'), ('Kumbha', '18-05-2025 19:34'), ('Makara', '05-12-2026 22:32')]


def _dt(text):
    return datetime.strptime(text, '%d-%m-%Y %H:%M')


def test_chart_s_transit_dates(T):
    for planet, known in (('Saturn', SATURN_S), ('Jupiter', JUPITER_S), ('Rahu', RAHU_S)):
        rows = [r for r in T['AI_Transits'] if r['planet'] == planet]
        entries = {(r['sign'], r['entry_local']) for r in rows}
        for sign, when in known:
            assert any(s == sign and abs(e - _dt(when)) <= timedelta(minutes=1) for s, e in entries), (planet, sign, when)
        # nothing missing in between: the known list is the whole run of entries for its span
        span = [r for r in rows if _dt(known[0][1]) - timedelta(minutes=2) <= r['entry_local'] <= _dt(known[-1][1]) + timedelta(minutes=2)]
        assert [r['sign'] for r in span] == [s for s, _ in known], planet
    # Ketu mirrors Rahu
    rahu = [r for r in T['AI_Transits'] if r['planet'] == 'Rahu']
    ketu = [r for r in T['AI_Transits'] if r['planet'] == 'Ketu']
    assert len(rahu) == len(ketu)
    for a, b in zip(rahu, ketu):
        assert (a['sign_no'] - b['sign_no']) % 12 == 6 and a['entry_utc'] == b['entry_utc'] and a['exit_utc'] == b['exit_utc']


def test_transit_table_is_complete_and_consistent(T, res):
    birth, end = res['meta']['local_dt'], NOW + timedelta(days=10 * 365.25)
    houses = by(T['AI_Houses'], 'house')
    for planet in ('Saturn', 'Jupiter', 'Rahu', 'Ketu'):
        rows = [r for r in T['AI_Transits'] if r['planet'] == planet]
        assert rows[0]['entry_local'] <= birth < rows[0]['exit_local']           # the stay running at birth
        assert rows[-1]['entry_local'] < end <= rows[-1]['exit_local'] + timedelta(minutes=1)
        for a, b in zip(rows, rows[1:]):
            assert a['exit_local'] == b['entry_local'] and a['exit_utc'] == b['entry_utc']
            step = (b['sign_no'] - a['sign_no']) % 12
            assert step in (1, 11) and b['motion_at_entry'] == ('direct' if step == 1 else 'retrograde')
        for r in rows:
            assert r['entry_local'] - r['entry_utc'] == IST and r['entry_utc'].second == 0
            assert r['house_from_lagna'] == (r['sign_no'] - 4) % 12 + 1          # Kataka lagna
            assert r['house_from_moon'] == (r['sign_no'] - 11) % 12 + 1          # Kumbha Moon
            assert r['sav_bindus'] == houses[r['house_from_lagna']]['sav_bindus']
            good = r['house_from_moon'] in E.TRANSIT_GOOD_FROM_MOON[planet]
            assert r['result_from_moon'] == ('favourable' if good else 'unfavourable')
            assert (r['own_bav_bindus'] is not None) == (planet in ('Saturn', 'Jupiter'))
            label = {12: 'sade sati - 1st phase', 1: 'sade sati - 2nd phase', 2: 'sade sati - 3rd phase',
                     8: 'ashtama', 4: 'ardhashtama', 7: 'kantaka', 10: 'kantaka'}.get(r['house_from_moon'])
            assert r['saturn_from_moon'] == (label if planet == 'Saturn' else None)
    assert all(r['motion_at_entry'] == 'retrograde' for r in T['AI_Transits'] if r['planet'] in ('Rahu', 'Ketu'))


def test_chart_s_sade_sati(T):
    rows = T['AI_SadeSati']
    whole = [r for r in rows if r['phase'] == 'sade sati - whole' and r['row_type'] == 'span']
    assert [r['cycle'] for r in whole] == [1, 2]
    now_cycle = whole[1]
    assert abs(now_cycle['start_local'] - _dt('24-01-2020 09:56')) <= timedelta(minutes=1)
    assert abs(now_cycle['end_local'] - _dt('23-02-2028 19:24')) <= timedelta(minutes=1)
    assert now_cycle['saturn_sign'] == 'Makara, Kumbha, Meena' and now_cycle['break_count'] == 1
    brk = [r for r in rows if r['cycle'] == 2 and r['phase'] == 'sade sati - whole' and r['row_type'] == 'break']
    assert len(brk) == 1 and brk[0]['saturn_sign'] == 'Mesha'
    assert abs(brk[0]['start_local'] - _dt('03-06-2027 05:28')) <= timedelta(minutes=1)
    assert abs(brk[0]['end_local'] - _dt('20-10-2027 07:12')) <= timedelta(minutes=1)
    third = next(r for r in rows if r['cycle'] == 2 and r['phase'] == 'sade sati - 3rd phase' and r['row_type'] == 'span')
    in_break = any(r['start_local'] <= NOW < r['end_local'] for r in rows
                   if r['cycle'] == 2 and r['phase'] == 'sade sati - 3rd phase' and r['row_type'] == 'break')
    assert third['start_local'] <= NOW < third['end_local'] and not in_break       # the report date is in the 3rd phase
    for r in rows:
        assert r['start_local'] < r['end_local'] and (r['break_count'] is None) == (r['row_type'] == 'break')
    for cyc in (1, 2):                                   # the three phases start in order
        starts = [next(r['start_local'] for r in rows if r['cycle'] == cyc and r['phase'].endswith(p) and r['row_type'] == 'span')
                  for p in ('1st phase', '2nd phase', '3rd phase')]
        assert starts == sorted(starts)


def test_sade_sati_cycles_merge_only_short_breaks():
    t0 = datetime(2000, 1, 1)
    def stay(sign, start_days, days):
        s = t0 + timedelta(days=start_days)
        return {'sign': sign, 'entry_utc': s, 'exit_utc': s + timedelta(days=days), 'entry_local': s,
                'exit_local': s + timedelta(days=days), 'motion': 'direct'}
    # Moon in Mesha (0): sade sati signs are Meena (11), Mesha (0), Rishabha (1)
    seq = [stay(10, 0, 100), stay(11, 100, 800), stay(0, 900, 100), stay(11, 1000, 150), stay(0, 1150, 800),
           stay(1, 1950, 700), stay(2, 2650, 120), stay(1, 2770, 100), stay(2, 2870, 900), stay(3, 3770, 900),
           stay(11, 12000, 900)]
    cycles = E.sade_sati_cycles(seq, 0)
    assert len(cycles) == 2
    first = cycles[0]
    assert [s['sign'] for s in first] == [11, 0, 11, 0, 1, 2, 1]
    assert [s['phase'] for s in first] == [1, 2, 1, 2, 3, None, 3]         # the stay in Mithuna is a break
    assert [s['sign'] for s in cycles[1]] == [11]


def test_double_transit(T):
    rows = T['AI_DoubleTransit']
    assert rows and all(r['method'] == 'modern method' for r in rows)
    window_end = NOW + timedelta(days=10 * 365.25)
    reach = {'occupies': 0, 'aspects (3rd)': 2, 'aspects (5th)': 4, 'aspects (7th)': 6, 'aspects (9th)': 8,
             'aspects (10th)': 9}
    for r in rows:
        assert r['start_local'] < r['end_local'] and r['start_local'] < window_end and r['end_local'] > NOW
        target = (R[r['sign']] - 3) % 12 + 1
        assert r['house'] == target
        assert (R[r['jupiter_sign']] + reach[r['jupiter_reaches_by']]) % 12 == R[r['sign']]
        assert (R[r['saturn_sign']] + reach[r['saturn_reaches_by']]) % 12 == R[r['sign']]
        assert r['jupiter_reaches_by'] in ('occupies', 'aspects (5th)', 'aspects (7th)', 'aspects (9th)')
        assert r['saturn_reaches_by'] in ('occupies', 'aspects (3rd)', 'aspects (7th)', 'aspects (10th)')
        # both planets really are in those signs for the whole span, by the transit table
        for planet, sign in (('Jupiter', r['jupiter_sign']), ('Saturn', r['saturn_sign'])):
            assert any(t['planet'] == planet and t['sign'] == sign
                       and t['entry_local'] <= r['start_local'] + timedelta(minutes=1)
                       and r['end_local'] <= t['exit_local'] + timedelta(minutes=1) for t in T['AI_Transits'])
    # on the report date Jupiter is in Kataka (reaching Kataka, Vrischika, Makara, Meena) and Saturn in Meena
    # (reaching Meena, Rishabha, Kanya, Dhanu): only Meena, the 9th house, has both
    running = [r for r in rows if r['start_local'] <= NOW < r['end_local']]
    assert [r['house'] for r in running] == [9]
    # the span has its real start (Jupiter's entry into Kataka) and end, not the edges of the 10-year range
    assert abs(running[0]['start_local'] - _dt('02-06-2026 01:50')) <= timedelta(minutes=1)
    assert abs(running[0]['end_local'] - _dt('31-10-2026 12:02')) <= timedelta(minutes=1)
    assert rows == sorted(rows, key=lambda r: (r['house'], r['start_local']))


def test_transit_now(T, res):
    rows = by(T['AI_TransitNow'], 'planet')
    assert list(rows) == A.PLANETS
    assert rows['Saturn']['sign'] == 'Meena' and rows['Jupiter']['sign'] == 'Kataka' and rows['Rahu']['sign'] == 'Kumbha'
    assert rows['Saturn']['house_from_moon'] == 2 and rows['Saturn']['house_from_lagna'] == 9
    assert (rows['Ketu']['longitude_deg'] - rows['Rahu']['longitude_deg']) % 360 == pytest.approx(180, abs=1e-5)
    assert rows['Rahu']['retrograde'] is None and rows['Rahu']['speed_deg_per_day'] < 0
    # against the Swiss Ephemeris directly
    jd = A.julian_day(NOW - IST)
    swe.set_sid_mode(swe.SIDM_LAHIRI)
    for name, pid in A.SWE_IDS.items():
        lon = swe.calc_ut(jd, pid, swe.FLG_MOSEPH | swe.FLG_SIDEREAL | swe.FLG_SPEED)[0][0] % 360
        assert rows[name]['longitude_deg'] == pytest.approx(lon, abs=1e-5)
        assert rows[name]['as_of_local'] == NOW and rows[name]['as_of_utc'] == NOW - IST
    # transit table and positions agree on the sign occupied now
    for planet in ('Saturn', 'Jupiter', 'Rahu', 'Ketu'):
        cur = [r for r in T['AI_Transits'] if r['planet'] == planet and r['entry_local'] <= NOW < r['exit_local']]
        assert len(cur) == 1 and cur[0]['sign'] == rows[planet]['sign']


# ── transit search against a plain one-day search ─────────────────────────────

def _plain_entries(pid, jd0, jd1, sid_mode, step=1.0):
    """Sign entries found with fixed steps and bisection, written without enrich.py."""
    flags = swe.FLG_MOSEPH | swe.FLG_SIDEREAL
    out = []
    with A._SWE_LOCK:
        swe.set_sid_mode(sid_mode)
        sign_at = lambda t: int((swe.calc_ut(t, pid, flags)[0][0] % 360) // 30) % 12
        t, s = jd0, None
        s = sign_at(t)
        while t < jd1:
            t2 = min(t + step, jd1)
            s2 = sign_at(t2)
            if s2 != s:
                lo, hi = t, t2
                while (hi - lo) * 86400 > 0.5:
                    mid = (lo + hi) / 2
                    if sign_at(mid) == s:
                        lo = mid
                    else:
                        hi = mid
                out.append((hi, s, sign_at(hi)))
                t, s = hi, sign_at(hi)
            else:
                t = t2
    return out


@pytest.mark.parametrize('body,node,pid,tolerance', [
    ('Saturn', 'mean', swe.SATURN, 2), ('Jupiter', 'mean', swe.JUPITER, 2), ('Rahu', 'mean', swe.MEAN_NODE, 2),
    # the true node can hang at a boundary, where the moment found depends on the bracket: to the minute
    ('Rahu', 'true', swe.TRUE_NODE, 60),
])
def test_transit_search_equals_one_day_step_search_1900_to_2100(body, node, pid, tolerance):
    jd0, jd1 = swe.julday(1900, 1, 1, 0.0), swe.julday(2100, 1, 1, 0.0)
    fast = E.sign_entries(body, jd0, jd1, swe.SIDM_LAHIRI, node)
    slow = _plain_entries(pid, jd0, jd1, swe.SIDM_LAHIRI)
    assert len(fast) == len(slow) > 100
    for a, b in zip(fast, slow):
        assert a[1:] == b[1:] and abs(a[0] - b[0]) * 86400 <= tolerance


def test_transit_speed_limits_hold():
    """The step rule relies on no body ever moving faster than its limit."""
    jd0 = swe.julday(1900, 1, 1, 0.0)
    swe.set_sid_mode(swe.SIDM_LAHIRI)
    for body, node in (('Saturn', 'mean'), ('Jupiter', 'mean'), ('Rahu', 'mean'), ('Rahu', 'true')):
        pid, vmax = E._body(body, node)
        top = max(abs(swe.calc_ut(jd0 + 2.0 * i, pid, swe.FLG_MOSEPH | swe.FLG_SIDEREAL | swe.FLG_SPEED)[0][3])
                  for i in range(36525))
        assert top < vmax * 0.9, (body, node, top)


def test_transits_follow_the_ayanamsha_and_node_setting():
    lahiri = E.enrich(A.compute(now=NOW, **S))['tables']['AI_Transits']
    raman = E.enrich(A.compute(now=NOW, ayanamsha='raman', **S))['tables']['AI_Transits']
    true = E.enrich(A.compute(now=NOW, node='true', **S))['tables']['AI_Transits']
    pick = lambda rows, planet, sign, year: next(r['entry_utc'] for r in rows if r['planet'] == planet
                                                 and r['sign'] == sign and r['entry_utc'].year == year)
    # Raman's ayanamsha is about 1.45 degrees smaller, so Saturn reaches sidereal Meena some weeks earlier
    gap = pick(lahiri, 'Saturn', 'Meena', 2025) - pick(raman, 'Saturn', 'Meena', 2025)
    assert timedelta(days=5) < gap < timedelta(days=40)
    assert pick(lahiri, 'Rahu', 'Kumbha', 2025) != pick(true, 'Rahu', 'Kumbha', 2025)
    assert [r for r in lahiri if r['planet'] == 'Saturn'] == [r for r in true if r['planet'] == 'Saturn']


# ── dasa ──────────────────────────────────────────────────────────────────────

def test_bhukti_table_matches_the_dasa_tree(T, res):
    tree = [(d['lord'], b['lord'], b['start'], b['end']) for d in res['vedic']['dasa']['dasas'] for b in d['sub']]
    rows = T['AI_Dasa']
    assert len(rows) == len(tree) == 80
    sec = timedelta(seconds=1)
    for row, (dl, bl, start, end) in zip(rows, tree):
        assert (row['dasa_lord'], row['bhukti_lord']) == (dl, bl)
        assert abs(row['start_local'] - start) < sec and abs(row['end_local'] - end) < sec
        assert row['start_local'].microsecond == 0
    for a, b in zip(rows, rows[1:]):
        assert a['end_local'] == b['start_local']
    assert rows[0]['start_local'] == datetime(1984, 11, 4, 0, 20) and rows[0]['age_at_start_years'] == 0
    planets = by(T['AI_Planets'], 'point')
    for row in rows:
        for who in ('dasa_lord', 'bhukti_lord'):
            p = planets[row[who]]
            assert row[f'{who}_house'] == p['house_whole_sign'] and row[f'{who}_houses_owned'] == p['houses_owned']
            assert row[f'{who}_dignity'] == p['dignity'] and row[f'{who}_vimsopaka_shodasavarga'] == p['vimsopaka_shodasavarga']
    cur = [r for r in rows if r['start_local'] <= NOW < r['end_local']]
    assert [(r['dasa_lord'], r['bhukti_lord']) for r in cur] == [('Mercury', 'Rahu')]
    assert cur[0]['start_local'].date() == date(2024, 10, 24)          # the Lahiri date named in the review


def test_pratyantar_parts_add_up_exactly_to_each_antaram(res):
    birth = res['meta']['local_dt']
    count = 0
    for d in res['vedic']['dasa']['dasas']:
        for b in d['sub']:
            for a in b['sub']:
                parts = E.pratyantars(a, birth)
                if a['start'] == birth:                 # clipped at birth: only the parts still to run
                    assert parts[0][1] == birth and parts[-1][2] == a['end']
                    continue
                count += 1
                assert len(parts) == 9
                assert parts[0][1] == a['start'] and parts[-1][2] == a['end']
                assert sum((e - s for _, s, e in parts), timedelta()) == a['end'] - a['start']
                first = A.DASA_ORDER.index(a['lord'])
                assert [p[0] for p in parts] == [A.DASA_ORDER[(first + i) % 9] for i in range(9)]
                span = (a['end'] - a['start']).total_seconds()
                for (lord, s, e), nxt in zip(parts, parts[1:] + [None]):
                    assert (e - s).total_seconds() == pytest.approx(span * A.DASA_YRS[lord] / 120, abs=2e-6)
                    if nxt:
                        assert e == nxt[1]
    assert count > 700


def test_pratyantar_of_a_known_antaram():
    # Venus dasa, Venus bhukti, Venus antaram: 20 x 20/120 x 20/120 years; its Venus part is a sixth of that
    start = datetime(2000, 1, 1)
    years = 20 * 20 / 120 * 20 / 120
    a = {'lord': 'Venus', 'start': start, 'end': start + timedelta(days=years * 365.25), 'years': years}
    parts = E.pratyantars(a, datetime(1990, 1, 1))
    assert [p[0] for p in parts] == ['Venus', 'Sun', 'Moon', 'Mars', 'Rahu', 'Jupiter', 'Saturn', 'Mercury', 'Ketu']
    assert (parts[0][2] - parts[0][1]).total_seconds() == pytest.approx(years * 365.25 * 86400 / 6, abs=1e-3)
    assert (parts[1][2] - parts[1][1]).total_seconds() == pytest.approx(years * 365.25 * 86400 * 6 / 120, abs=1e-3)
    # born half way through: earlier parts are dropped and the running one starts at birth
    birth = start + (a['end'] - start) / 2
    clipped = E.pratyantars(dict(a, start=birth), birth)
    assert clipped[0][1] == birth and clipped[-1][2] == a['end'] and clipped[0][0] == 'Rahu'
    assert [p[0] for p in clipped] == ['Rahu', 'Jupiter', 'Saturn', 'Mercury', 'Ketu']


def test_pratyantar_table_covers_one_year_back_to_five_ahead(T):
    rows = T['AI_Pratyantar']
    lo, hi = NOW - timedelta(days=365.25), NOW + timedelta(days=5 * 365.25)
    assert rows[0]['start_local'] <= lo < rows[0]['end_local'] and rows[-1]['start_local'] < hi <= rows[-1]['end_local']
    for a, b in zip(rows, rows[1:]):
        assert a['end_local'] == b['start_local']
    cur = [r for r in rows if r['start_local'] <= NOW < r['end_local']]
    assert len(cur) == 1
    assert (cur[0]['dasa_lord'], cur[0]['bhukti_lord'], cur[0]['antaram_lord']) == ('Mercury', 'Rahu', 'Venus')
    assert all(r['days'] == pytest.approx((r['end_local'] - r['start_local']).total_seconds() / 86400, abs=1e-3) for r in rows)


# ── divisional charts ─────────────────────────────────────────────────────────

@pytest.mark.parametrize('key,div', [(v[0], v[1]) for v in A.VARGAS if v[0] != 'D30'])
def test_divisional_degree_runs_from_0_to_30_inside_each_part(key, div):
    size = 30.0 / div
    exact = (size * 1024).is_integer()            # part edges that a float holds exactly (15, 7.5, 3, 2.5 ...)
    for sign in (0, 1, 5, 11):
        for part in sorted({0, min(1, div - 1), div // 2, div - 1}):
            start = sign * 30 + part * size
            if exact:
                assert E.varga_degree(start, key) == 0.0 and A.varga_part(start, key) == part
            else:
                start += 1e-9                     # just inside the part: the edge itself is not a float
                assert E.varga_degree(start, key) == pytest.approx(0, abs=1e-6)
            end = sign * 30 + (part + 1) * size - 1e-7
            assert 29.99 < E.varga_degree(end, key) < 30
            assert A.varga_sign(start, key) == A.varga_sign(end, key) == A.varga_sign(start + size / 2, key)
            assert E.varga_degree(start + size / 2, key) == pytest.approx(15, abs=1e-5)
            assert E.varga_degree(start + size / 4, key) == pytest.approx(7.5, abs=1e-5)


def test_divisional_degree_in_trimsamsa():
    for base, edges in ((0, [0, 5, 10, 18, 25, 30]), (30, [0, 5, 12, 20, 25, 30])):          # odd sign, even sign
        for lo, hi in zip(edges, edges[1:]):
            assert E.varga_degree(base + lo, 'D30') == pytest.approx(0, abs=1e-9)
            assert 29.99 < E.varga_degree(base + hi - 1e-7, 'D30') < 30
            assert E.varga_degree(base + (lo + hi) / 2, 'D30') == pytest.approx(15)
            assert A.varga_sign(base + lo, 'D30') == A.varga_sign(base + hi - 1e-7, 'D30')


def test_divisional_degree_never_leaves_its_sign():
    rnd = random.Random(7)
    for _ in range(4000):
        lon = rnd.uniform(0, 360)
        for key in A.VARGA_KEYS:
            assert 0 <= E.varga_degree(lon, key) < 30
    assert E.varga_degree(123.456789, 'D1') == pytest.approx(3.456789)
    assert E.fmt_deg(29.99999) == '29°59\'59"' and E.fmt_deg(0) == '00°00\'00"' and E.fmt_deg(15.5) == '15°30\'00"'


def test_varga_sheet_matches_the_engine(T, res):
    rows = T['AI_Vargas']
    assert len(rows) == 176
    v = {x['key']: x for x in res['vedic']['vargas']}
    lon = {p['name']: p['lon'] for p in res['vedic']['planets'] + [res['vedic']['mandi']]}
    for r in rows:
        vg = v[r['chart']]
        assert r['sign_no'] == vg['signs'][r['point']] + 1 and r['sign'] == A.RASIS[r['sign_no'] - 1]
        assert r['sign_no'] == A.varga_sign(lon[r['point']], r['chart']) + 1
        assert r['house'] == (vg['signs'][r['point']] - vg['lagna_sign']) % 12 + 1
        assert 0 <= r['deg_in_chart'] < 30 and r['sign_lord'] == A.SIGN_LORDS[r['sign_no'] - 1]
        assert (r['dignity'] is not None) == (r['point'] in E.SEVEN) == (r['debilitated'] is not None)
    d1 = {r['point']: r for r in rows if r['chart'] == 'D1'}
    planets = by(T['AI_Planets'], 'point')
    for name, r in d1.items():
        assert r['deg_in_chart'] == pytest.approx(planets[name]['deg_in_sign'], abs=1e-6)
        assert r['dignity'] == planets[name]['dignity'] and r['house'] == planets[name]['house_whole_sign']
    lagnas = by(T['AI_VargaLagnas'], 'chart')
    for key in A.VARGA_KEYS:
        lg = next(r for r in rows if r['chart'] == key and r['point'] == 'Lagna')
        assert lg['house'] == 1 and lagnas[key]['lagna_sign'] == lg['sign']
        lord = next(r for r in rows if r['chart'] == key and r['point'] == lagnas[key]['lagna_lord'])
        assert lagnas[key]['lagna_lord_house'] == lord['house'] and lagnas[key]['lagna_lord_dignity'] == lord['dignity']


# ── planets and houses ────────────────────────────────────────────────────────

def test_planet_sheet_chart_s(T, res):
    p = by(T['AI_Planets'], 'point')
    assert list(p) == ['Lagna'] + A.PLANETS + ['Mandi']
    # the three house systems side by side, as in the review
    assert (p['Sun']['house_whole_sign'], p['Sun']['bhava_sripati'], p['Sun']['kp_house_placidus']) == (4, 4, 3)
    assert (p['Mercury']['house_whole_sign'], p['Mercury']['bhava_sripati'], p['Mercury']['kp_house_placidus']) == (5, 4, 4)
    assert (p['Rahu']['house_whole_sign'], p['Rahu']['bhava_sripati'], p['Rahu']['kp_house_placidus']) == (11, 10, 10)
    assert p['Sun']['dignity'] == 'Neutral' and p['Sun']['debilitated'] == 'yes'      # Venus: natural enemy, temporary friend
    assert p['Sun']['debilitation_cancelled_by'] == 'a, b' and p['Saturn']['debilitation_cancelled_by'] is None
    assert p['Saturn']['dignity'] == 'Exalted' and p['Saturn']['combust'] == 'yes'
    assert p['Saturn']['distance_from_sun_deg'] == pytest.approx(6.637, abs=1e-3) and p['Saturn']['combust_orb_deg'] == 15
    assert p['Saturn']['distance_from_deep_exaltation_deg'] == pytest.approx(4.5318, abs=1e-3)
    assert p['Jupiter']['dignity'] == 'Own sign' and p['Jupiter']['moolatrikona'] == 'no'       # 15°40' is past 10°
    assert p['Moon']['dignity'] == 'Enemy' and p['Mars']['functional_nature'] == 'Yogakaraka'
    assert p['Mars']['houses_owned'] == '5, 10' and p['Mars']['houses_aspected'] == '1, 9, 12'
    assert p['Saturn']['planets_aspected'] == 'Mars, Jupiter' and p['Mars']['aspected_by'] == 'Saturn'
    assert p['Lagna']['aspected_by'] == 'Mars, Saturn' and p['Mercury']['conjunct_with'] == 'Venus, Ketu'
    assert p['Mars']['vargottama'] == 'yes' and p['Sun']['vargottama'] == 'no'
    assert p['Rahu']['houses_owned'] == 'none' and p['Rahu']['dignity'] is None and p['Lagna']['houses_owned'] is None
    assert p['Mandi']['kp_house_placidus'] is None and p['Mandi']['house_whole_sign'] == 1
    for name, row in p.items():
        src = res['vedic']['mandi'] if name == 'Mandi' else next(x for x in res['vedic']['planets'] if x['name'] == name)
        assert row['longitude_deg'] == pytest.approx(src['lon'], abs=1e-6) and row['deg_in_sign_text'] == src['dms']
        assert row['sign_no'] == int(row['longitude_deg'] // 30) + 1 and 0 <= row['deg_in_sign'] < 30
        assert row['longitude_deg'] == pytest.approx((row['sign_no'] - 1) * 30 + row['deg_in_sign'], abs=1e-5)
        assert row['nakshatra'] == A.NAKS[src['nak']] and row['pada'] == src['pada']
        assert row['house_from_moon'] == (row['sign_no'] - p['Moon']['sign_no']) % 12 + 1
    # vargottama is the same fact the Divisional Charts sheet prints
    import charts
    _, table = charts.varga_table(res, 'en')
    assert {r['name']: r['vargottama'] for r in table} == {n: row['vargottama'] == 'yes' for n, row in p.items()}


def test_house_sheet_chart_s(T, res):
    h = by(T['AI_Houses'], 'house')
    assert [h[i]['sign'] for i in (1, 4, 7, 10)] == ['Kataka', 'Thula', 'Makara', 'Mesha']
    assert h[4]['occupants_whole_sign'] == 'Sun, Saturn' and h[1]['occupants_whole_sign'] == 'Mandi'
    assert h[10]['aspected_by'] == 'Sun, Jupiter, Saturn' and h[7]['aspected_by'] == 'none'
    assert h[6]['lord'] == 'Jupiter' and h[6]['lord_house'] == 6 and h[6]['lord_dignity'] == 'Own sign'
    assert h[1]['house_groups'] == 'kendra, trikona' and h[6]['house_groups'] == 'dusthana, upachaya' and h[2]['house_groups'] == 'none'
    assert h[8]['house_from_moon'] == 1 and h[8]['occupants_whole_sign'] == 'Moon'
    for i in range(1, 13):
        cusp = res['kp']['cusps'][i - 1]
        assert h[i]['kp_cusp_sub_lord'] == cusp['sub_lord'] and h[i]['kp_cusp_sign'] == A.RASIS[cusp['sign']]
    everyone = [n for i in range(1, 13) for n in h[i]['occupants_whole_sign'].split(', ') if n != 'none']
    assert sorted(everyone) == sorted(A.PLANETS + ['Mandi'])
    assert sorted(n for i in range(1, 13) for n in h[i]['kp_occupants'].split(', ') if n != 'none') == sorted(A.PLANETS)


def test_facts_sheet_chart_s(T):
    f = {r['key']: r['value'] for r in T['AI_Facts']}
    assert f['birth_date'] == date(1984, 11, 4) and f['birth_time'] == '00:20:00'
    assert f['birth_datetime_local'] == datetime(1984, 11, 4, 0, 20) and f['birth_datetime_utc'] == datetime(1984, 11, 3, 18, 50)
    assert f['utc_offset_hours'] == 5.5 and f['time_zone'] == 'Asia/Kolkata' and f['node_type'] == 'mean node'
    assert f['lagna_sign'] == 'Kataka' and f['moon_sign'] == 'Kumbha' and f['moon_nakshatra'] == 'Purva Bhadrapada'
    assert f['weekday'] == 'Saturday'                   # 00:20 on Sunday the 4th is before sunrise
    assert f['badhaka_house'] == 11 and f['badhaka_sign'] == 'Rishabha' and f['badhaka_lord'] == 'Venus'
    assert f['report_datetime_local'] == NOW and f['ayanamsha_deg'] == pytest.approx(23.645, abs=0.01)
    assert f['default_house_system'].startswith('whole sign') and 'AI_Dasa' in f['default_dasa_table']
    assert len(f) == len(T['AI_Facts'])


# ── yogas on built charts ─────────────────────────────────────────────────────

def test_mahapurusha_and_simple_yogas():
    y = yogas('Mesha', Mars='Makara', Jupiter='Kataka', Venus='Thula', Saturn='Mesha', Mercury='Mithuna',
              Sun='Mithuna', Moon='Thula')
    assert y['Ruchaka yoga']['status'] == 'found' and y['Hamsa yoga']['status'] == 'found'
    assert y['Malavya yoga']['status'] == 'found' and y['Sasa yoga']['status'] == 'not found'      # debilitated
    assert y['Bhadra yoga']['status'] == 'not found'                                                # own sign, 3rd house
    assert y['Bhadra yoga - counted from the Moon (variant)']['status'] == 'not found'              # 9th from the Moon
    assert y['Hamsa yoga - counted from the Moon (variant)']['status'] == 'found'                   # 10th from the Moon
    assert y['Gajakesari yoga']['status'] == 'found' and y['Budha-Aditya yoga']['status'] == 'found'
    assert y['Chandra-Mangala yoga']['status'] == 'not found'
    assert y['Guru-Chandala yoga']['status'] == 'not found'
    assert yogas('Mesha', Jupiter='Dhanu')['Guru-Chandala yoga']['status'] == 'found'              # with Ketu
    assert yogas('Mesha', Moon='Kanya', Mars='Kanya')['Chandra-Mangala yoga']['status'] == 'found'


def test_neecha_bhanga_conditions():
    # Saturn debilitated in Mesha (lagna). (a) its lord Mars in a kendra; (b) the Sun, exalted in Mesha, in a kendra
    y = yogas('Mesha', Saturn=('Mesha', 2.0), Mars='Kataka', Sun='Makara', Moon='Rishabha', Mercury='Dhanu',
              Jupiter='Dhanu', Venus='Dhanu')
    assert y['Neecha bhanga of Saturn - condition (a)']['status'] == 'found'
    assert y['Neecha bhanga of Saturn - condition (b)']['status'] == 'found'
    assert 'Neecha bhanga of Saturn - condition (c)' not in y              # Mesha 2° is Mesha in D9, not Thula
    # (c): Saturn at 20°-23°20' of Mesha is in Thula in D9
    y = yogas('Rishabha', Saturn=('Mesha', 21.0), Mars='Mithuna', Sun='Mithuna', Moon='Rishabha', Mercury='Mithuna',
              Jupiter='Mithuna', Venus='Mithuna')
    assert y['Neecha bhanga of Saturn - condition (c)']['status'] == 'found'
    assert 'Neecha bhanga of Saturn - condition (a)' not in y and 'Neecha bhanga of Saturn - condition (b)' not in y
    # none of the three
    y = yogas('Rishabha', Saturn=('Mesha', 2.0), Mars='Mithuna', Sun='Mithuna', Moon='Rishabha', Mercury='Mithuna',
              Jupiter='Mithuna', Venus='Mithuna')
    assert y['Neecha bhanga of Saturn']['status'] == 'not found'
    c = ctx('Rishabha', Saturn=('Mesha', 2.0), Mars='Mithuna', Sun='Mithuna', Moon='Rishabha', Mercury='Mithuna',
            Jupiter='Mithuna', Venus='Mithuna')
    E._yogas(c)
    c.margins = None
    saturn = next(r for r in E._planets(c) if r['point'] == 'Saturn')
    assert saturn['debilitated'] == 'yes' and saturn['debilitation_cancelled_by'] == 'none'
    # the Moon is debilitated in Vrischika, where no planet is exalted: (b) cannot apply
    y = yogas('Mesha', Moon='Vrischika', Mars='Mesha', Sun='Mesha', Mercury='Mesha', Jupiter='Mesha', Venus='Mesha',
              Saturn='Kumbha')
    assert y['Neecha bhanga of Moon - condition (a)']['status'] == 'found'
    assert 'Neecha bhanga of Moon - condition (b)' not in y
    assert yogas('Mesha', Saturn='Kumbha', Sun='Simha', Moon='Kataka', Mars='Mesha', Mercury='Mithuna',
                 Jupiter='Dhanu', Venus='Thula')['Neecha bhanga']['status'] == 'not found'


def test_raja_dhana_viparita_and_parivartana():
    # Mesha lagna: Mars 1/8, Venus 2/7, Mercury 3/6, Moon 4, Sun 5, Jupiter 9/12, Saturn 10/11
    y = E._yogas(ctx('Mesha', Moon='Simha', Sun='Kataka', Mars='Mithuna', Mercury='Kanya', Jupiter='Makara',
                     Saturn='Dhanu', Venus='Rishabha', Rahu='Mithuna', Ketu='Dhanu'))
    raja = [r for r in y if r['yoga'] == 'Raja yoga']
    # Moon (lord of 4) and Sun (5) exchange signs; Jupiter (9) and Saturn (10) exchange signs
    assert {r['planets'] for r in raja if 'exchange of signs' in r['note']} == {'Sun, Moon', 'Jupiter, Saturn'}
    assert all(r['status'] == 'found' for r in raja)
    swaps = {r['planets'] for r in y if r['yoga'] == 'Parivartana yoga'}
    assert swaps == {'Sun, Moon', 'Jupiter, Saturn'}
    dhana = {r['planets'] for r in y if r['yoga'] == 'Dhana yoga' and r['status'] == 'found'}
    assert 'Jupiter, Saturn' in dhana                   # lords of 9 and 11, exchange
    names = by(y, 'yoga')
    assert names['Harsha yoga (Viparita raja yoga)']['status'] == 'found'       # Mercury, lord of 6, in the 6th
    assert names['Sarala yoga (Viparita raja yoga)']['status'] == 'not found'   # Mars, lord of 8, in the 3rd
    assert names['Vimala yoga (Viparita raja yoga)']['status'] == 'not found'   # Jupiter, lord of 12, in the 10th
    # a lord of a dusthana in another dusthana also counts
    y2 = yogas('Mesha', Mars='Meena', Jupiter='Vrischika', Mercury='Meena')
    assert all(y2[f'{n} yoga (Viparita raja yoga)']['status'] == 'found' for n in ('Harsha', 'Sarala', 'Vimala'))
    # mutual aspect as the link: Moon (4) in Mesha and Jupiter (9) in Thula face each other
    y3 = E._yogas(ctx('Mesha', Moon='Mesha', Jupiter='Thula', Sun='Mithuna', Mars='Mithuna', Mercury='Mithuna',
                      Venus='Mithuna', Saturn='Mithuna'))
    assert any(r['yoga'] == 'Raja yoga' and r['planets'] == 'Moon, Jupiter' and 'mutual aspect' in r['note'] for r in y3)
    none = E._yogas(ctx('Mesha', Sun='Mesha', Moon='Rishabha', Mars='Mithuna', Mercury='Kataka', Jupiter='Simha',
                        Venus='Kanya', Saturn='Vrischika'))
    assert [r['status'] for r in none if r['yoga'] == 'Parivartana yoga'] == ['not found']


def test_yogas_from_the_moon_and_the_sun():
    base = dict(Sun='Thula', Moon='Mesha', Mars='Thula', Mercury='Thula', Jupiter='Thula', Venus='Thula', Saturn='Thula')
    y = yogas('Simha', **dict(base, Mars='Rishabha'))
    assert (y['Sunapha yoga']['status'], y['Anapha yoga']['status'], y['Durudhara yoga']['status']) == ('found', 'not found', 'not found')
    assert y['Kemadruma yoga']['status'] == 'not found'
    y = yogas('Simha', **dict(base, Mars='Rishabha', Venus='Meena'))
    assert (y['Sunapha yoga']['status'], y['Anapha yoga']['status'], y['Durudhara yoga']['status']) == ('not found', 'not found', 'found')
    # the Sun in the 2nd from the Moon does not count
    y = yogas('Simha', **dict(base, Sun='Rishabha', Mars='Simha', Mercury='Simha', Jupiter='Simha', Venus='Simha', Saturn='Simha'))
    assert y['Sunapha yoga']['status'] == 'not found' and y['Kemadruma yoga']['status'] == 'cancelled'   # planets in the lagna
    # formed and not cancelled: the five planets are 6th from the Moon and 3rd from the lagna
    y = yogas('Kataka', Sun='Mesha', Moon='Mesha', Mars='Kanya', Mercury='Kanya', Jupiter='Kanya', Venus='Kanya', Saturn='Kanya')
    assert y['Kemadruma yoga']['status'] == 'found' and 'Cancelled: no' in y['Kemadruma yoga']['note']
    # Vesi, Vasi, Ubhayachari; the Moon does not count
    sun = dict(Sun='Mesha', Moon='Rishabha', Mars='Simha', Mercury='Simha', Jupiter='Simha', Venus='Simha', Saturn='Simha')
    y = yogas('Simha', **sun)
    assert [y[n]['status'] for n in ('Vesi yoga', 'Vasi yoga', 'Ubhayachari yoga')] == ['not found'] * 3
    y = yogas('Simha', **dict(sun, Venus='Meena'))
    assert [y[n]['status'] for n in ('Vesi yoga', 'Vasi yoga', 'Ubhayachari yoga')] == ['not found', 'found', 'not found']
    y = yogas('Simha', **dict(sun, Venus='Meena', Mercury='Rishabha'))
    assert [y[n]['status'] for n in ('Vesi yoga', 'Vasi yoga', 'Ubhayachari yoga')] == ['not found', 'not found', 'found']


def test_adhi_amala_kuja_and_kala_sarpa():
    y = yogas('Mesha', Moon='Mesha', Mercury='Kanya', Jupiter='Thula', Venus='Vrischika', Sun='Kanya', Mars='Mithuna', Saturn='Mithuna')
    assert y['Adhi yoga']['status'] == 'found' and '3 of 3' in y['Adhi yoga']['note']
    y = yogas('Mesha', Moon='Mesha', Mercury='Kanya', Jupiter='Makara', Venus='Mithuna', Sun='Kanya', Mars='Mithuna', Saturn='Mithuna')
    assert y['Adhi yoga']['status'] == 'found' and '1 of 3' in y['Adhi yoga']['note'] and 'partial' in y['Adhi yoga']['note']
    assert y['Amala yoga']['status'] == 'found' and y['Amala yoga']['planets'] == 'Jupiter'       # 10th from lagna and Moon
    assert yogas('Mesha', Moon='Mesha', Saturn='Makara', Jupiter='Mithuna', Mercury='Mithuna', Venus='Mithuna')['Amala yoga']['status'] == 'not found'
    for place, flagged in zip(range(1, 13), [1, 1, 0, 1, 0, 0, 1, 1, 0, 0, 0, 1]):
        mars = A.RASIS[(R['Simha'] + place - 1) % 12]
        y = yogas('Simha', Mars=mars, Moon='Simha', Venus='Simha')
        for ref in ('the lagna', 'the Moon', 'Venus'):
            assert (y[f'Kuja dosha from {ref}']['status'] == 'found') == bool(flagged), (place, ref)
    # Kala Sarpa by longitude: Rahu at 15° Mithuna, Ketu at 15° Dhanu
    inside = dict(Sun='Kataka', Moon='Simha', Mars='Kanya', Mercury='Thula', Jupiter='Vrischika', Venus='Kataka', Saturn='Simha')
    name = 'Kala Sarpa yoga (popular, not in BPHS)'
    assert yogas('Mesha', **inside)[name]['status'] == 'found'
    assert yogas('Mesha', **dict(inside, Saturn='Makara'))[name]['status'] == 'not found'
    # same sign as Rahu: before it (10°) is outside the arc, after it (20°) is inside
    assert yogas('Mesha', **dict(inside, Saturn=('Mithuna', 10.0)))[name]['status'] == 'not found'
    assert yogas('Mesha', **dict(inside, Saturn=('Mithuna', 20.0)))[name]['status'] == 'found'
    other = {k: A.RASIS[(R[v] + 6) % 12] for k, v in inside.items()}
    assert yogas('Mesha', **other)[name]['status'] == 'found'


# ── topics ────────────────────────────────────────────────────────────────────

def test_topic_map_is_the_fixed_table(T):
    rows = by(T['AI_TopicMap'], 'topic')
    assert list(rows) == ['Job and career', 'Business', 'Marriage', 'Children', 'Abroad / onsite', 'Education',
                          'Home and vehicle', 'Money', 'Strain (general)', 'Parents', 'Siblings', 'Spiritual life']
    job = rows['Job and career']
    assert (job['d1_houses'], job['significators'], job['divisional_charts'], job['kp_houses']) == (
        '10, 6, 2, 11', 'Sun, Saturn, Mercury', 'D10', '2, 6, 10, 11')
    assert rows['Abroad / onsite']['divisional_charts'] == 'D9, D10' and rows['Abroad / onsite']['kp_houses'] == '3, 9, 12'
    assert rows['Home and vehicle']['d1_houses'] == '4' and rows['Parents']['kp_houses'] == 'none'
    assert rows['Marriage']['significators'] == 'Venus; Jupiter also for a woman'
    assert rows['Marriage']['significators_for_this_chart'] == 'Venus'               # no gender given
    woman = E.enrich(A.compute(now=NOW, gender='Female', **S))['tables']
    assert by(woman['AI_TopicMap'], 'topic')['Marriage']['significators_for_this_chart'] == 'Venus, Jupiter'
    assert any(r['topic'] == 'Marriage' and r['subject'] == 'significator Jupiter' for r in woman['AI_TopicFacts'])
    assert not any(r['topic'] == 'Marriage' and r['subject'] == 'significator Jupiter' for r in T['AI_TopicFacts'])


def test_topic_facts_repeat_the_other_sheets(T):
    facts = {(r['topic'], r['subject'], r['factor']): r['value'] for r in T['AI_TopicFacts']}
    assert len(facts) == len(T['AI_TopicFacts'])
    houses, planets, lagnas = by(T['AI_Houses'], 'house'), by(T['AI_Planets'], 'point'), by(T['AI_VargaLagnas'], 'chart')
    for topic, key_houses, sig, vargas, kp in E.TOPICS:
        for h in key_houses:
            assert facts[(topic, f'house {h}', 'sign')] == houses[h]['sign']
            assert facts[(topic, f'house {h}', 'lord')] == houses[h]['lord']
            assert facts[(topic, f'house {h}', 'lord_house')] == houses[h]['lord_house']
            assert facts[(topic, f'house {h}', 'occupants')] == houses[h]['occupants_whole_sign']
            assert facts[(topic, f'house {h}', 'aspected_by')] == houses[h]['aspected_by']
            assert facts[(topic, f'house {h}', 'sav_bindus')] == houses[h]['sav_bindus']
        for p in sig:
            assert facts[(topic, f'significator {p}', 'house')] == planets[p]['house_whole_sign']
        for key in vargas:
            assert facts[(topic, f'chart {key}', 'lagna_sign')] == lagnas[key]['lagna_sign']
            assert facts[(topic, f'chart {key}', 'lagna_uncertain')] == lagnas[key]['lagna_uncertain']
    assert facts[('Job and career', 'chart D10', 'lagna_uncertain')] == 'yes'
    assert facts[('Children', 'chart D7', 'lagna_uncertain')] == 'no'
    assert facts[('Job and career', 'house 10', 'aspected_by')] == 'Sun, Jupiter, Saturn'


def test_topic_scores_follow_the_rule(T, res):
    rows = T['AI_TopicScores']
    assert len(rows) == 80 * 12
    c = E.Ctx(res)
    E._yogas(c)
    planets = by(T['AI_Planets'], 'point')
    for row in rows:
        lb = '' if row['bhukti_lord_letters'] == 'none' else row['bhukti_lord_letters']
        ld = '' if row['dasa_lord_letters'] == 'none' else row['dasa_lord_letters']
        assert set(lb) <= set('abcdef') and list(lb) == sorted(lb)
        expect = len(lb) + len(ld) / 2 - (0 if row['deduction'] == 'none' else 1)
        assert row['relevance_score_not_a_prediction'] == expect
    # letters worked out by hand for Mars and the job topic (houses 10, 6, 2, 11; Sun, Saturn, Mercury; D10; KP 2, 6, 10, 11)
    # a: Mars owns 10. b: Mars sits in 6. c: Mars aspects 12, 1, 9 - no key house. d: not a significator.
    # e: in D10 Mars is in Kanya, 12th from the Thula lagna, not own or exalted - no. f: KP - Mars signifies 6.
    mars = next(r for r in rows if r['bhukti_lord'] == 'Mars' and r['topic'] == 'Job and career')
    d10 = next(r for r in T['AI_Vargas'] if r['chart'] == 'D10' and r['point'] == 'Mars')
    assert (d10['sign'], d10['house']) == ('Kanya', 12)
    assert 6 in c.kp_signifies['Mars']
    assert mars['bhukti_lord_letters'] == 'abf'
    # deductions: Saturn is combust; the Sun is debilitated but its debilitation is cancelled, and nobody is under 10
    ded = {r['bhukti_lord']: r['deduction'] for r in rows}
    assert ded['Saturn'] == 'combust' and ded['Sun'] == 'none' and ded['Rahu'] == 'none'
    assert all(planets[p]['vimsopaka_shodasavarga'] >= 10 for p in E.SEVEN)
    assert E.penalty(c, 'Saturn') == ['combust'] and E.penalty(c, 'Sun') == []
    c.neecha_bhanga['Sun'] = []
    assert E.penalty(c, 'Sun') == ['debilitated, not cancelled']
    c.vimsopaka['Sun']['shodasavarga'] = 9.9
    assert E.penalty(c, 'Sun') == ['vimsopaka under 10', 'debilitated, not cancelled']
    one = next(r for r in rows if (r['dasa_lord'], r['bhukti_lord'], r['topic']) == ('Mercury', 'Rahu', 'Abroad / onsite'))
    assert one['start_local'].date() == date(2024, 10, 24) and 'd' in one['bhukti_lord_letters']       # Rahu is a significator


# ── the block as a whole ──────────────────────────────────────────────────────

def test_block_shape(ai, T):
    assert ai['schema'] == 'horoscopegen-ai/1' and ai['report_datetime_local'] == NOW
    assert ai['sheets'] == list(T) == [
        'AI_ReadMe', 'AI_Facts', 'AI_Planets', 'AI_Houses', 'AI_Pairs', 'AI_Vargas', 'AI_VargaLagnas', 'AI_Strength',
        'AI_Ashtakavarga', 'AI_Yogas', 'AI_Dasa', 'AI_Pratyantar', 'AI_Transits', 'AI_TransitNow', 'AI_SadeSati',
        'AI_DoubleTransit', 'AI_TopicMap', 'AI_TopicFacts', 'AI_TopicScores']
    for sheet, rows in T.items():
        cols = [c['name'] for c in ai['columns'][sheet]]
        assert rows and len(cols) == len(set(cols)) and all(c and c == c.strip() for c in cols)
        assert all(list(r) == cols for r in rows), sheet
        assert all(c['meaning'] for c in ai['columns'][sheet])
        for c in cols:                                   # no column is empty from top to bottom
            assert any(r[c] is not None for r in rows), (sheet, c)
        for r in rows:
            for v in r.values():
                assert v is None or isinstance(v, (str, int, float, datetime, date)) and not isinstance(v, bool)
                assert not (isinstance(v, str) and (v.startswith('=') or v == '' or 'None' in v))
    assert len(T['AI_Planets']) == 11 and len(T['AI_Houses']) == 12 and len(T['AI_Pairs']) == 36
    assert len(T['AI_VargaLagnas']) == 16 and len(T['AI_Strength']) == 7 and len(T['AI_TopicMap']) == 12
    assert [r['rule'] for r in ai['rules']] == [n for n, _ in E.RULES]


def test_readme_says_what_an_ai_must_know(ai, T):
    lines = T['AI_ReadMe']
    text = ' '.join(r['text'] for r in lines)
    assert '2026-10-07 17:00:00' in text and 'UTC+05:30' in text
    assert 'Whole sign from the lagna' in text and 'AI_Dasa' in text
    assert 'Krishnamurti ayanamsha, Placidus houses' in text and 'Do not mix' in text
    assert 'go stale' in text and "today's date" in text
    assert 'rule-based relevance, not a prediction' in text and 'modern method' in text
    documented = {r['item'] for r in lines if r['section'] == 'Columns'}
    assert documented == {f'{s}.{c["name"]}' for s in ai['sheets'] for c in ai['columns'][s]}
    assert {r['item'] for r in lines if r['section'] == 'Rules'} == {n for n, _ in E.RULES}
    assert [r['text'] for r in lines if r['section'] == 'Not computed for this chart'] == [
        'Nothing: every value could be computed for this chart.']
    missing = ' '.join(r['text'] for r in lines if r['section'] == 'Not in this file')
    for word in ('Shadbala', 'Chara karakas', 'Yogini', 'Gulika', 'Remedies'):
        assert word in missing


def test_nothing_about_length_of_life(ai, T):
    banned = ('marak', 'ayur', 'longevity', 'lifespan', 'life span', 'death', 'mrityu', 'alpayu')
    allowed = 'Length of life: left out on purpose; this file gives no figure and no analysis of it'
    for sheet, rows in T.items():
        for r in rows:
            for v in list(r.values()) + list(r):
                if isinstance(v, str) and v != allowed:
                    low = v.lower()
                    assert not any(b in low for b in banned), (sheet, v)
    assert allowed in [r['text'] for r in T['AI_ReadMe']]
    for c in (x for cols in ai['columns'].values() for x in cols):
        assert not any(b in (c['name'] + ' ' + c['meaning']).lower() for b in banned)


def test_enrich_leaves_the_result_alone_and_is_repeatable(res):
    before = json.dumps(A.to_jsonable(res), sort_keys=True, default=str)
    first = E.enrich(res)
    second = E.enrich(res)
    assert json.dumps(A.to_jsonable(res), sort_keys=True, default=str) == before
    assert A.to_jsonable(first) == A.to_jsonable(second)
    json.dumps(A.to_jsonable(first))                    # everything in the block can be sent as JSON
    later = E.enrich(res, now=NOW + timedelta(days=400))
    assert later['report_datetime_local'] == NOW + timedelta(days=400)
    assert later['tables']['AI_Pratyantar'][0]['start_local'] > first['tables']['AI_Pratyantar'][0]['start_local']
    assert later['tables']['AI_Planets'] == first['tables']['AI_Planets']


def test_language_and_chart_style_do_not_change_the_block():
    base = A.to_jsonable(E.enrich(A.compute(now=NOW, **S)))
    for lang, style in (('ta', 'south'), ('hi', 'north'), ('bi', 'south')):
        other = A.to_jsonable(E.enrich(A.compute(now=NOW, lang=lang, chart_style=style, **S)))
        assert other == base


def test_a_place_without_sunrise_leaves_cells_empty_and_says_why():
    r = A.compute('X', '2001-06-21', '12:00', 'Tromso', lat=69.65, lon=18.96, tz='Europe/Oslo', now=NOW)
    assert r['vedic']['mandi'] is None
    block = E.enrich(r)
    t = block['tables']
    mandi = t['AI_Planets'][-1]
    assert mandi['point'] == 'Mandi' and all(v is None for k, v in mandi.items() if k != 'point')
    rows = [x for x in t['AI_Vargas'] if x['point'] == 'Mandi']
    assert len(t['AI_Vargas']) == 176 and len(rows) == 16 and all(x['sign'] is None for x in rows)
    why = [x['text'] for x in t['AI_ReadMe'] if x['section'] == 'Not computed for this chart']
    assert any(x.startswith('Mandi: the Sun does not rise or set') for x in why)
    facts = {x['key']: x['value'] for x in t['AI_Facts']}
    assert facts['sunrise_local'] is None and facts['birth_by_day_or_night'] is None
    assert 'Mandi' not in ' '.join(h['occupants_whole_sign'] for h in t['AI_Houses'])
    from excel_generator import generate_excel
    assert generate_excel(r, 'en', ai=block)[:2] == b'PK'


@pytest.mark.parametrize('b', [CHART_A, CHART_B])
def test_other_reference_charts_are_consistent(b):
    r = A.compute(now=NOW, **b)
    t = E.enrich(r)['tables']
    p = by(t['AI_Planets'], 'point')
    h = by(t['AI_Houses'], 'house')
    for name in A.PLANETS:
        house = p[name]['house_whole_sign']
        assert name in h[house]['occupants_whole_sign'].split(', ')
        for q in (x for x in p[name]['planets_aspected'].split(', ') if x != 'none'):
            assert name in p[q]['aspected_by'].split(', ')
        for hh in (int(x) for x in p[name]['houses_aspected'].split(', ')):
            assert name in h[hh]['aspected_by'].split(', ')
    for a, c in combinations(E.SEVEN, 2):
        pair = next(x for x in t['AI_Pairs'] if (x['planet_a'], x['planet_b']) == (a, c))
        assert (pair['same_sign'] == 'yes') == (p[a]['sign'] == p[c]['sign'])
        assert pair['temporary_a_to_b'] == pair['temporary_b_to_a']             # temporary friendship is mutual
    totals = {x['planet']: x['bindus'] for x in t['AI_Ashtakavarga'] if x['row_type'] == 'total'}
    assert totals['Sarvashtakavarga'] == 337
    for planet in ('Saturn', 'Jupiter', 'Rahu', 'Ketu'):
        rows = [x for x in t['AI_Transits'] if x['planet'] == planet]
        assert all(x['exit_local'] == y['entry_local'] for x, y in zip(rows, rows[1:]))
    assert all(x['lagna_if_later'] != x['lagna_sign'] for x in t['AI_VargaLagnas'])


# ── the workbook ──────────────────────────────────────────────────────────────

def test_ai_sheets_follow_notes_in_order(workbooks, ai):
    for lang in ('en', 'ta'):
        wb = workbooks[lang][1]
        assert len(wb.sheetnames) == 7 + 19
        assert wb.sheetnames[7:] == ai['sheets']
    assert workbooks['en'][1].sheetnames[:7] == ['Summary', 'Vedic', 'Divisional Charts', 'KP', 'ALP', 'Dasa', 'Notes']


def test_ai_sheets_are_flat_tables(workbooks, ai):
    wb = workbooks['en'][1]
    for sheet in ai['sheets']:
        ws = wb[sheet]
        cols = [c['name'] for c in ai['columns'][sheet]]
        rows = ai['tables'][sheet]
        assert not ws.merged_cells.ranges, sheet
        assert [c.value for c in ws[1]] == cols                          # header in row 1, from A1, none blank
        assert ws.max_column == len(cols) and ws.max_row == len(rows) + 1
        assert ws.freeze_panes == 'A2'
        for row in ws.iter_rows(min_row=2):
            for c in row:
                assert c.data_type != 'f' and not (isinstance(c.value, str) and c.value.startswith('='))
                if isinstance(c.value, datetime):
                    assert c.number_format in ('yyyy-mm-dd hh:mm:ss', 'yyyy-mm-dd')
                elif isinstance(c.value, (int, float)):
                    assert c.number_format in ('General', '0.000000')      # a number is never shown as a date
                assert c.fill.fill_type is None                             # no colour carries meaning
        # same values as the block that also goes into the JSON file
        for r, src in zip(ws.iter_rows(min_row=2, values_only=True), rows):
            for got, want in zip(r, src.values()):
                if isinstance(want, datetime):
                    assert abs(got - want) < timedelta(seconds=1)
                elif isinstance(want, date):
                    assert got.date() == want
                elif isinstance(want, float):
                    assert got == pytest.approx(want, abs=1e-9)
                else:
                    assert got == want


def test_ai_sheets_are_the_same_in_every_language(workbooks):
    en, ta = workbooks['en'][1], workbooks['ta'][1]
    for name in en.sheetnames[7:]:
        a = list(en[name].iter_rows(values_only=True))
        b = list(ta[name].iter_rows(values_only=True))
        assert a == b, name
    assert en.sheetnames[:7] != ta.sheetnames[:7]


def test_ai_sheets_load_with_pandas_as_one_rectangle(workbooks, ai):
    pd = pytest.importorskip('pandas')
    frames = pd.read_excel(io.BytesIO(workbooks['en'][0]), sheet_name=None)
    for sheet in ai['sheets']:
        df = frames[sheet]
        assert list(df.columns) == [c['name'] for c in ai['columns'][sheet]]
        assert not any(str(c).startswith('Unnamed') for c in df.columns)
        assert len(df) == len(ai['tables'][sheet])
        assert not df.isna().all(axis=0).any() and not df.isna().all(axis=1).any()
    t = frames['AI_Transits']
    assert str(t['entry_local'].dtype).startswith('datetime64') and str(t['entry_utc'].dtype).startswith('datetime64')
    assert str(frames['AI_Dasa']['start_local'].dtype).startswith('datetime64')
    assert str(frames['AI_Planets']['longitude_deg'].dtype) == 'float64'
    saturn = frames['AI_Planets'].set_index('point').loc['Saturn']
    assert saturn['sign'] == 'Thula' and saturn['house_whole_sign'] == 4 and saturn['dignity'] == 'Exalted'


def test_round_trip_d1_houses_from_the_planet_sheet_alone(workbooks, res):
    ws = workbooks['ta'][1]['AI_Planets']
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    lagna = next(r for r in rows if r['point'] == 'Lagna')
    rebuilt = {r['point']: (int(r['longitude_deg'] // 30) - int(lagna['longitude_deg'] // 30)) % 12 + 1 for r in rows}
    engine = {p['name']: p['house'] for p in res['vedic']['planets']}
    engine['Mandi'] = res['vedic']['mandi']['house']
    assert rebuilt == engine
    assert {r['point']: r['house_whole_sign'] for r in rows} == engine
    assert {r['point']: (r['sign_no'] - lagna['sign_no']) % 12 + 1 for r in rows} == engine
    # the D1 chart as drawn on the Vedic sheet: sign by sign, the same planets
    by_sign = {}
    for r in rows:
        by_sign.setdefault(r['sign'], []).append(r['point'])
    assert by_sign == {'Kataka': ['Lagna', 'Mandi'], 'Thula': ['Sun', 'Saturn'], 'Kumbha': ['Moon'],
                       'Dhanu': ['Mars', 'Jupiter'], 'Vrischika': ['Mercury', 'Venus', 'Ketu'], 'Rishabha': ['Rahu']}


def test_notes_sheet_names_every_rule(workbooks, ai):
    for lang in ('en', 'ta'):
        ws = workbooks[lang][1].worksheets[6]
        text = {ws.cell(r, 2).value: ws.cell(r, 3).value for r in range(1, ws.max_row + 1)}
        assert 'AI sheets: rules and variants' in text
        for rule in ai['rules']:
            assert text[rule['rule']] == rule['variant']
        first_new = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == 'AI sheets: rules and variants')
        assert first_new > GOLDEN['workbooks'][lang][6]['rows']             # added below what was there


# ── old behaviour is unchanged ────────────────────────────────────────────────

def _first_difference(got, want):
    for key in ('title', 'columns', 'merged', 'merged_count', 'heights', 'widths', 'freeze', 'breaks'):
        if got[key] != want[key]:
            return f"{want['title']}: {key} differs"
    for i, (a, b) in enumerate(zip(got['row_digests'], want['row_digests']), start=1):
        if a != b:
            return f"{want['title']}: row {i} differs"
    return None


@pytest.mark.parametrize('lang', ['en', 'ta'])
def test_seven_old_sheets_are_cell_for_cell_what_main_produced(lang, workbooks):
    from excel_generator import generate_excel
    want = GOLDEN['workbooks'][lang]
    # 1. without the AI block the workbook is exactly the old one
    plain = G.workbook_print(generate_excel(A.compute(lang=lang, now=NOW, **S), lang))
    assert plain == want, [_first_difference(a, b) for a, b in zip(plain, want)]
    # 2. with the AI block: six sheets are identical, and Notes is identical down to its old last row
    notes_rows = want[6]['rows']
    with_ai = G.workbook_print(workbooks[lang][0], limit={6: notes_rows})
    assert with_ai == want, [_first_difference(a, b) for a, b in zip(with_ai, want)]
    assert len(want) == 7 and sum(s['rows'] for s in want) > 1300


# ── API ───────────────────────────────────────────────────────────────────────

@pytest.fixture()
def client(monkeypatch):
    import app as app_module
    monkeypatch.setattr(app_module, 'compute', functools.partial(A.compute, now=NOW))
    return app_module.app.test_client()


def test_api_horoscope_without_ai_is_byte_for_byte_what_main_returned(client):
    rv = client.post('/api/horoscope', json=G.API_BODY)
    assert rv.status_code == 200
    assert {'sha256': hashlib.sha256(rv.data).hexdigest(), 'bytes': len(rv.data)} == GOLDEN['api_horoscope']
    assert 'ai' not in rv.get_json()
    for off in (False, 'false', 0, None, 'no'):
        again = client.post('/api/horoscope', json=dict(G.API_BODY, ai=off))
        assert again.data == rv.data


def test_api_horoscope_with_ai_adds_only_the_ai_block(client, ai):
    plain = client.post('/api/horoscope', json=G.API_BODY).get_json()
    rv = client.post('/api/horoscope', json=dict(G.API_BODY, ai=True))
    assert rv.status_code == 200
    full = rv.get_json()
    block = full.pop('ai')
    assert full == plain
    assert block == json.loads(json.dumps(A.to_jsonable(ai)))
    assert block['tables']['AI_Planets'][7]['point'] == 'Saturn' and block['tables']['AI_Planets'][7]['dignity'] == 'Exalted'
    assert block['tables']['AI_Transits'][0]['entry_utc'].count('T') == 1
    assert 'ai' in client.post('/api/horoscope', json=dict(G.API_BODY, ai='true')).get_json()


def test_api_download_json(client, ai):
    rv = client.post('/api/download/json', json=G.API_BODY)
    assert rv.status_code == 200
    assert 'HoroscopeGen_Chart_S.json' in rv.headers['Content-Disposition']
    assert rv.headers['Content-Type'].startswith('application/json')
    data = json.loads(rv.data.decode('utf-8'))
    assert set(data) == {'meta', 'vedic', 'kp', 'alp', 'ai'}
    plain = client.post('/api/horoscope', json=G.API_BODY).get_json()
    for key in ('meta', 'vedic', 'kp', 'alp'):
        assert data[key] == plain[key]
    assert data['ai'] == json.loads(json.dumps(A.to_jsonable(ai)))
    assert data['ai']['tables']['AI_Facts'][2] == {'key': 'birth_date', 'value': '1984-11-04',
                                                   'note': 'civil date at the birth place'}
    assert int(rv.headers['Content-Length']) == len(rv.data)


def test_api_download_excel_has_the_ai_sheets(client):
    import openpyxl
    rv = client.post('/api/download/excel', json=G.API_BODY)
    assert rv.status_code == 200 and rv.data[:2] == b'PK'
    wb = openpyxl.load_workbook(io.BytesIO(rv.data), read_only=True)
    assert len(wb.sheetnames) == 26 and wb.sheetnames[7] == 'AI_ReadMe' and wb.sheetnames[-1] == 'AI_TopicScores'


def test_api_json_errors(client):
    assert client.post('/api/download/json', json={'name': 'X'}).status_code == 400
    assert client.post('/api/download/json', json=dict(G.API_BODY, dob='bad')).status_code == 400
    assert client.get('/api/download/json').status_code == 405
