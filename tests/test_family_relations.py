"""
Tests for family_relations.py (section I of the family reading: relationships over time).

  * the example rule: a daughter's Surya dasa reads "needs care / space" with her
    father when her Sun is debilitated in the 8th, and "supportive" when it is exalted
  * I3 counts the 6th, 8th and 12th from each relation house; I5 reads natural
    relations both ways
  * roles: daughter-in-law, son-in-law and maternal uncle validate and read back;
    relations between members are derived through the anchor
  * policy: no row names a member as a cause, no banned or health words in any
    language, a minor gets no spouse or in-law rows
  * the same input gives the same rows; the API, Excel and PDF carry section I
"""
import io
import re
from datetime import datetime

import pytest

import astro_engine as A
import enrich as E
import family_engine as F
import family_relations as FR
import predictions as P
from rulebook import BANNED

NOW = datetime(2026, 10, 7, 11, 30)
LANGS = ('en', 'ta', 'hi', 'te', 'kn', 'ml', 'mr', 'bn')
FATHER = dict(name='Father R', dob='1975-01-26', tob='21:50', pob='Madurai', lat=9.9252, lon=78.1198, gender='male')
WEAK_SUN = dict(name='Daughter W', dob='2005-11-05', tob='15:10', pob='Chennai', lat=13.08, lon=80.27, gender='female')
STRONG_SUN = dict(name='Daughter S', dob='2008-04-25', tob='06:10', pob='Chennai', lat=13.08, lon=80.27, gender='female')


def _family(*members, years=15):
    ms = [{'relation': rel, 'res': A.compute(now=NOW, **kw)} for rel, kw in members]
    return F.family(ms, {'years': years}, now=NOW)


def _rows(result):
    return result['sections']['I']['rows']


@pytest.fixture(scope='module')
def weak():
    return _family(('Self', FATHER), ('Daughter', WEAK_SUN))


@pytest.fixture(scope='module')
def strong():
    return _family(('Self', FATHER), ('Daughter', STRONG_SUN))


@pytest.fixture(scope='module')
def big():
    return _family(('Self', dict(name='Ravi', dob='1984-11-04', tob='00:20', pob='Chennai', lat=13.08, lon=80.27,
                                 gender='male')),
                   ('Spouse', dict(name='Anu', dob='1987-03-12', tob='06:40', pob='Madurai', lat=9.93, lon=78.12)),
                   ('Daughter', dict(name='Meena', dob='2012-07-21', tob='14:05', pob='Chennai', lat=13.08, lon=80.27)),
                   ('Mother', dict(name='Kamala', dob='1958-01-30', tob='09:15', pob='Trichy', lat=10.79, lon=78.70)),
                   ('Maternal uncle', dict(name='Gopal', dob='1955-05-02', tob='18:30', pob='Trichy', lat=10.79,
                                           lon=78.70)), years=10)


def _sun_rows(result):
    return [r for r in _rows(result) if r['member_running'].startswith('Daughter') and r['other'] == 'Father R'
            and r['lord'] == 'Sun' and r['level'] == 'dasa' and r['rule_id'] in ('I1', 'I4')]


def test_example_rule_weak_sun_needs_space(weak):
    rows = _sun_rows(weak)
    i1 = [r for r in rows if r['rule_id'] == 'I1']
    assert i1 and all(r['tone'] == 'care' for r in i1)
    assert 'Debilitated' in i1[0]['basis'] and '8th from lagna' in i1[0]['basis']
    assert 'give each other space' in i1[0]['reading']
    assert i1[0]['relation'] == 'father' and i1[0]['start'].year == 2032
    i4 = [r for r in rows if r['rule_id'] == 'I4']            # the Sun separates: distance eases it
    assert i4 and i4[0]['tone'] == 'care' and 'distance' in i4[0]['reading']


def test_example_rule_strong_sun_supports(strong):
    i1 = [r for r in _sun_rows(strong) if r['rule_id'] == 'I1']
    assert i1 and all(r['tone'] == 'supportive' for r in i1)
    assert 'Exalted' in i1[0]['basis']
    i4 = [r for r in _sun_rows(strong) if r['rule_id'] == 'I4']
    assert all(r['tone'] == 'mixed' for r in i4)               # a strong separating planet: distance for good reasons


@pytest.mark.parametrize('house', range(1, 13))
def test_i3_counts_6_8_12_from_each_house(weak, house):
    m = weak['_people'][0]
    got = FR.trik_lords(m, house)
    want = {}
    for k in (6, 8, 12):
        h = (house + k - 2) % 12 + 1
        want.setdefault(A.SIGN_LORDS[m.house_sign(h)], []).append((k, h))
    assert got == want
    # father (9th): the lords of the 2nd, 4th and 8th
    if house == 9:
        assert sorted(h for v in got.values() for _, h in v) == [2, 4, 8]


def test_i5_natural_relations():
    assert E.natural_relation('Sun', 'Saturn') == E.natural_relation('Saturn', 'Sun') == 'Enemy'
    assert E.natural_relation('Sun', 'Jupiter') == E.natural_relation('Jupiter', 'Sun') == 'Friend'


def test_i5_rows_are_both_ways(big):
    for r in (x for x in _rows(big) if x['rule_id'] == 'I5'):
        assert ('enemies both ways' in r['basis']) == (r['tone'] == 'care')
        assert r['level'] == 'months' and r['end'] > r['start']


def test_new_relations_validate_and_read_back():
    F.check_relations(['Self', 'Daughter-in-law', 'Son-in-law', 'Maternal uncle'])
    assert F.relation_gender('Daughter-in-law') == 'female' and F.relation_gender('Son-in-law') == 'male'
    r = _family(('Self', dict(FATHER, name='Anchor')), ('Daughter-in-law', dict(WEAK_SUN, name='DIL')),
                ('Maternal uncle', dict(FATHER, name='Uncle', dob='1950-02-10')))
    rel = {x['member']: x for x in r['sections']['B']['relations']}
    assert rel['DIL']['house'] == 11 and rel['DIL']['karaka'] == 'Venus'
    assert rel['Uncle']['house'] == 6 and rel['Uncle']['karaka'] == 'Mercury'
    back = {(x['member'], x['house']) for x in r['sections']['B']['read_back']}
    assert ('DIL', 3) in back                                  # male anchor: father-in-law, 9th from the 7th
    assert ('Uncle', 7) in back and ('Uncle', 3) in back       # the anchor is his sister's child


def test_roles_derived_through_the_anchor(big):
    rel = {(r['member_running'], r['other']): r['relation'] for r in _rows(big) if r['relation']}
    assert rel.get(('Anu', 'Kamala')) == 'mother_in_law'
    assert rel.get(('Kamala', 'Anu')) == 'child_in_law'
    assert rel.get(('Kamala', 'Meena')) == 'grandchild'
    assert rel.get(('Ravi', 'Gopal')) == 'maternal_uncle'
    assert rel.get(('Gopal', 'Ravi')) == 'sister_child'
    assert rel.get(('Kamala', 'Gopal')) in ('younger_sibling', 'elder_sibling')


def test_minor_gets_no_spouse_or_in_law_rows(big):
    # Meena is 14: only parent, child, sibling, grandparent and uncle rows
    for r in _rows(big):
        if 'Meena' in (r['member_running'], r['other']):
            assert r['relation'] not in FR.ADULT_ROLES


CAUSE = re.compile(r'\b(because of|caused by|harms?|troubles?|bad for|blame)\b', re.I)


def test_policy_every_language(big, weak):
    for result in (big, weak):
        for r in _rows(result):
            for lang in LANGS:
                text = FR.render(r, lang)
                assert text and '{' not in text
                assert not BANNED.search(text) and not P.HEALTH.search(text) and not P.NATIVE_BANNED.search(text), text
            assert not CAUSE.search(r['reading']) and not CAUSE.search(r['basis']), r['reading']


def test_templates_complete():
    keys = FR.TEMPLATES
    ph = re.compile(r'\{(\w+)\}')
    for k, row in keys.items():
        for lang in LANGS:
            assert row.get(lang), (k, lang)
            assert set(ph.findall(row[lang])) == set(ph.findall(row['en'])), (k, lang)
            assert not P.NATIVE_BANNED.search(row[lang]) and not P.HEALTH.search(row[lang]), (k, lang)


def test_overall_rows(big):
    sums = [r for r in _rows(big) if r['rule_id'] == 'Σ']
    assert sums
    for s in sums:
        parts = [p.split() for p in s['basis'].split(' + ')]
        tones = [p[1] for p in parts]
        assert len(tones) >= 2
        if s['tone'] == 'care':
            assert tones.count('care') >= 2 and 'supportive' not in tones


def test_rows_inside_range_and_deterministic(big):
    again = _family(('Self', dict(name='Ravi', dob='1984-11-04', tob='00:20', pob='Chennai', lat=13.08, lon=80.27,
                                  gender='male')),
                    ('Spouse', dict(name='Anu', dob='1987-03-12', tob='06:40', pob='Madurai', lat=9.93, lon=78.12)),
                    ('Daughter', dict(name='Meena', dob='2012-07-21', tob='14:05', pob='Chennai', lat=13.08, lon=80.27)),
                    ('Mother', dict(name='Kamala', dob='1958-01-30', tob='09:15', pob='Trichy', lat=10.79, lon=78.70)),
                    ('Maternal uncle', dict(name='Gopal', dob='1955-05-02', tob='18:30', pob='Trichy', lat=10.79,
                                            lon=78.70)), years=10)
    strip = lambda rows: [{k: v for k, v in r.items()} for r in rows]
    assert strip(_rows(again)) == strip(_rows(big))
    hi = NOW.replace(year=NOW.year + 11)
    for r in _rows(big):
        assert r['end'] > NOW.replace(day=1) and r['start'] < hi


def test_rules_and_exclusions_printed(big):
    names = [x['name'] for x in big['rules']]
    for n in ('I1 Karaka period', 'I3 Trik from the relation house', 'I4 Separating planets', 'I overall'):
        assert n in names
    assert any('Moola' in x for x in big['not_included'])
    assert big['summary']['lines'][-1]['section'] == 'I'


def test_outputs_carry_section_i(big):
    import openpyxl
    from excel_generator import generate_family_excel
    from pdf_generator import generate_family_pdf
    for lang in ('en', 'ta', 'te'):
        wb = openpyxl.load_workbook(io.BytesIO(generate_family_excel(big, lang)))
        assert wb.sheetnames[1] == 'Relationships' and 'AI_FamilyRelationsOverTime' in wb.sheetnames
        ws = wb['Relationships']
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert FR.T('title', lang) in texts
        assert generate_family_pdf(big, lang)[:5] == b'%PDF-'
    ai = wb['AI_FamilyRelationsOverTime']
    assert ai.max_row - 1 == len(_rows(big))


def test_api_answer_has_localised_rows():
    import app as app_module
    app_module.limiter.enabled = False
    members = [dict(FATHER, relation='Self'), dict(WEAK_SUN, relation='Daughter')]
    data = app_module.app.test_client().post('/api/family', json={'members': members, 'lang': 'ta'}).get_json()
    rot = data['relations_over_time']
    assert rot['title'] == FR.T('title', 'ta') and rot['rows']
    row = rot['rows'][0]
    for k in ('reading', 'tone_label', 'relation_label', 'level_label', 'basis', 'start', 'end'):
        assert k in row
    assert 'key' not in row and 'params' not in row


def test_no_english_left_in_relationships_sheet(big):
    """The reading, tone, relation and period words of section I are in the report language (basis stays English)."""
    import openpyxl
    from excel_generator import generate_family_excel
    from tests.i18n_audit import english_runs
    names = re.compile(r'Ravi|Anu|Meena|Kamala|Gopal')
    for lang in LANGS[1:]:
        ws = openpyxl.load_workbook(io.BytesIO(generate_family_excel(big, lang)))['Relationships']
        for row in ws.iter_rows(min_row=1):
            for c in row[:-1]:                                   # every column but the basis
                if isinstance(c.value, str):
                    assert not english_runs(names.sub(' ', c.value)), (lang, c.coordinate, c.value)
