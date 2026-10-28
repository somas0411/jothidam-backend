"""
Tests for predictions.py: rule-based readings without any AI model.

  * the same chart and report moment always give the same rows
  * every row has text and a basis; no template placeholder is left over
  * every template has English and Tamil text
  * no banned word (death, length of life, source names) and no health word
    appears in any template or in the rows of the sample charts
  * below age 18 there are no career, business, marriage or children rows
  * the Excel workbook gains Predictions and AI_Predictions; the PDF builds in
    English and Tamil; /api/horoscope adds the block only when asked
"""
import io
import json
import os
import re
from datetime import datetime

import pytest

import astro_engine as A
import predictions as P
from enrich import enrich
from rulebook import BANNED

NOW = datetime(2026, 10, 7, 17, 0)
CHARTS = [
    dict(name='Chart S', dob='1984-11-04', tob='00:20', pob='Chennai', lat=13.0827, lon=80.2707),
    dict(name='Chart A', dob='1984-05-12', tob='06:30', pob='Chennai', lat=13.08, lon=80.27),
    dict(name='Chart B', dob='1998-03-06', tob='01:37', pob='Delhi', lat=28.6139, lon=77.2090),
    dict(name='Chart C', dob='1961-08-15', tob='14:05', pob='Mumbai', lat=19.0760, lon=72.8777),
    dict(name='Chart D', dob='1975-01-26', tob='21:50', pob='Madurai', lat=9.9252, lon=78.1198),
]
MINOR = dict(name='Child', dob='2016-06-01', tob='09:15', pob='Chennai', lat=13.0827, lon=80.2707)
SOURCE = ('chinnaraj', 'youtube', 'channel')


def _build(chart, lang='en'):
    res = A.compute(lang=lang, now=NOW, **chart)
    return res, P.build(res, enrich(res), lang)


@pytest.fixture(scope='module')
def blocks():
    return {(c['name'], lang): _build(c, lang)[1] for c in CHARTS + [MINOR] for lang in ('en', 'ta')}


def test_same_input_same_rows():
    a, b = _build(CHARTS[0])[1], _build(CHARTS[0])[1]
    assert a == b and a['rows']


def test_every_row_has_text_and_basis(blocks):
    for block in blocks.values():
        for r in block['rows']:
            assert r['text'].strip() and r['basis'].strip(), r
            assert '{' not in r['text'] and '}' not in r['text'], r['text']
            assert r['section'] in P.SECTIONS
            assert r['tone'] in ('favourable', 'mixed', 'challenging', 'neutral')
            assert r['status'] in ('running now', 'upcoming', 'past', 'dormant', 'always')


def test_rows_are_in_fixed_order(blocks):
    for block in blocks.values():
        keys = [r['sort_key'] for r in block['rows']]
        assert keys == sorted(keys) and len(set(keys)) == len(keys)
        assert block['rows'][0]['section'] == 'summary'


def test_every_template_has_english_and_tamil():
    for key, entry in P.texts().items():
        assert set(entry) >= {'en', 'ta'}, key
        if isinstance(entry['en'], dict):
            assert set(entry['en']) == set(entry['ta']), key


def test_every_template_used_exists():
    for chart in CHARTS + [MINOR]:
        res = A.compute(now=NOW, **chart)
        _, used = P.build_rows(res, enrich(res), 'ta')
        assert used <= set(P.texts())


def test_no_banned_or_health_words_in_templates():
    for name in os.listdir(P.DATA_DIR):
        with open(os.path.join(P.DATA_DIR, name), encoding='utf-8') as f:
            text = f.read()
        assert not BANNED.search(text.replace('Ayushman', '')), name   # the yoga name itself is never printed
        assert not P.HEALTH.search(text), name
        assert not any(word in text.lower() for word in SOURCE), name
    with open(P.__file__, encoding='utf-8') as f:
        assert 'chinnaraj' not in f.read().lower()


def test_no_banned_or_health_words_in_rows(blocks):
    for block in blocks.values():
        for r in block['rows']:
            blob = ' '.join(str(r[k]) for k in ('topic', 'item', 'text', 'basis', 'period'))
            assert not BANNED.search(blob), blob
            assert not P.HEALTH.search(blob), blob


def test_every_section_has_the_disclaimer(blocks):
    for block in blocks.values():
        for s in block['sections']:
            assert any(r['section'] == s['key'] and r['item'] == 'disclaimer' for r in block['rows'])


def test_summary_has_six_to_eight_lines(blocks):
    for block in blocks.values():
        assert 6 <= len(block['summary']) <= 8, block['summary']


def test_no_adult_topics_below_18(blocks):
    block = blocks[('Child', 'en')]
    topics = {r['topic'] for r in block['rows']}
    assert not topics & set(P.MINOR_TOPICS)
    assert not any(r['item'] in ('career fields', 'job or business', 'partner') for r in block['rows'])
    period = [r for r in block['rows'] if r['section'] == 'period' and r['item'] in ('dasa', 'bhukti', 'antaram')]
    assert period and all('studies' in r['text'] for r in period)
    adult = blocks[('Chart S', 'en')]
    assert {'Job and career', 'Marriage'} <= {r['topic'] for r in adult['rows']}


def test_tamil_text_is_tamil(blocks):
    block = blocks[('Chart S', 'ta')]
    assert block['title'] == 'பலன்கள்'
    lagna = next(r for r in block['rows'] if r['item'] == 'lagna')
    assert any('஀' <= ch <= '௿' for ch in lagna['text'])


def test_calendar_has_twelve_months_and_the_rule(blocks):
    rows = [r for r in blocks[('Chart S', 'en')]['rows'] if r['section'] == 'calendar']
    assert len([r for r in rows if r['item'] == 'month']) == P.CALENDAR_MONTHS
    assert any(r['item'] == 'rule' for r in rows)


def test_remedies_name_no_gemstone(blocks):
    for block in blocks.values():
        for r in block['rows']:
            if r['section'] == 'remedies':
                assert not re.search(r'\b(gems?|gemstones?|stones?|ruby|sapphire|pearl|emerald|diamond)\b', r['text'], re.I), r['text']


def test_excel_gains_the_two_sheets():
    import openpyxl
    from excel_generator import generate_excel
    for lang, title in (('en', 'Predictions'), ('ta', 'பலன்கள்')):
        res = A.compute(lang=lang, now=NOW, **CHARTS[0])
        ai = enrich(res)
        wb = openpyxl.load_workbook(io.BytesIO(generate_excel(res, lang, ai=ai, predictions=True)))
        assert wb.sheetnames[7] == title and wb.sheetnames[-1] == 'AI_Predictions'
        flat = wb['AI_Predictions']
        assert [c.value for c in flat[1]] == [c[0] for c in P.COLUMNS]
        assert flat.max_row == len(P.build(res, ai, 'en')['rows']) + 1
        readme = [row[0] for row in wb['AI_ReadMe'].iter_rows(values_only=True)]
        assert 'AI_Predictions' in readme
        # without the flag the workbook is the one it was
        plain = openpyxl.load_workbook(io.BytesIO(generate_excel(res, lang, ai=ai)))
        assert wb.sheetnames == plain.sheetnames[:7] + [title] + plain.sheetnames[7:] + ['AI_Predictions']


@pytest.mark.parametrize('lang', ['en', 'ta'])
def test_pdf_builds_with_predictions(lang):
    from pdf_generator import generate_pdf
    res = A.compute(lang=lang, now=NOW, **CHARTS[0])
    block = P.build(res, enrich(res), lang)
    with_p = generate_pdf(res, lang, predictions=block)
    plain = generate_pdf(res, lang)
    assert with_p[:5] == b'%PDF-' and len(with_p) > len(plain)


def test_api_horoscope_adds_predictions_only_when_asked():
    import app as app_module
    client = app_module.app.test_client()
    body = dict(CHARTS[0], lang='en', chartStyle='south', ayanamsha='lahiri')
    plain = client.post('/api/horoscope', json=body).get_json()
    assert 'predictions' not in plain
    full = client.post('/api/horoscope', json=dict(body, predictions=True)).get_json()
    block = full.pop('predictions')
    assert set(block) == {'title', 'disclaimer', 'columns', 'sections', 'labels', 'rows', 'summary', 'rule'}
    assert block['rows'] and 'ai' not in full
    full.pop('meta'), plain.pop('meta')                 # generated_at differs between the two calls
    assert full == plain
    json.dumps(block)                                   # plain JSON values only


def test_pdf_and_json_downloads_can_leave_predictions_out():
    import app as app_module
    client = app_module.app.test_client()
    body = dict(CHARTS[0], lang='en')
    on = client.post('/api/download/pdf', json=body)
    off = client.post('/api/download/pdf', json=dict(body, predictions=False))
    assert on.status_code == off.status_code == 200 and len(on.data) > len(off.data)
    data = json.loads(client.post('/api/download/json', json=dict(body, predictions='false')).data)
    assert 'predictions' not in data
