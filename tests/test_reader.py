"""
Tests for what schema horoscopegen-ai/2 adds for an AI reader: the dasa
levels, the faster transits and stations, AI_PeriodFacts, AI_TopicPromise,
AI_TopicWindows, the optional inputs, AI_Brief and the "ai": "brief" option.

Reference values that do not come from the code under test:
  * Values worked by hand for Chart S (04-11-1984 00:20 Chennai), with the
    working written beside each one.
  * Plain re-implementations written in this file, which read only the AI
    sheets: transit aspects, the antaram tree, score ranks, the promise
    tests of two topics, and the windows (a day-by-day scan).
  * A second way of finding a station: the extreme of the longitude itself,
    not the zero of the speed.
  * The retrograde of Mars in early 2027 as printed in published ephemerides:
    stations on 10 January and 1 April 2027 (UTC dates).
  * Dates printed for the same birth by another program (Om Tamil Calendar):
    Saturn-Rahu bhukti 09-12-2009 to 15-10-2012.
  * The calendar: 4 November 1984 was a Sunday.
"""
import functools
import io
import json
import math
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
import swisseph as swe

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import astro_engine as A          # noqa: E402
import enrich as E                # noqa: E402
import make_golden as G           # noqa: E402
from test_enrich import CHART_A, CHART_B, EXTRAS, IST, NOW, S, _plain_entries, by, chart  # noqa: E402

LIFE = {'marital_status': 'Married', 'marriage_year': 2012, 'children_count': 1, 'first_child_birth_year': 2014,
        'elder_siblings': 0, 'younger_siblings': 1, 'work_type': 'employed', 'lives_abroad': False,
        'events': [{'type': 'marriage', 'year': 2012}, {'type': 'job change', 'year': 2023},
                   {'type': 'home bought', 'year': 2019}]}
FULL = {'tob_uncertainty_min': 5, 'coordinates_source': 'entered', 'life': LIFE}

ORDER = ['Ketu', 'Venus', 'Sun', 'Moon', 'Mars', 'Rahu', 'Jupiter', 'Saturn', 'Mercury']      # Vimshottari order
YEARS = dict(zip(ORDER, (7, 20, 6, 10, 7, 18, 16, 19, 17)))
ASPECT = {'Mars': (4, 7, 8), 'Jupiter': (5, 7, 9), 'Saturn': (3, 7, 10)}                      # others: the 7th
SIGNS = ['Mesha', 'Rishabha', 'Mithuna', 'Kataka', 'Simha', 'Kanya', 'Thula', 'Vrischika', 'Dhanu', 'Makara',
         'Kumbha', 'Meena']


@pytest.fixture(scope='module')
def res():
    return A.compute(now=NOW, **S)


@pytest.fixture(scope='module')
def ai(res):
    return E.enrich(res, extras=EXTRAS)


@pytest.fixture(scope='module')
def T(ai):
    return ai['tables']


@pytest.fixture(scope='module')
def full(res):
    """The block when the user entered everything that can be entered."""
    return E.enrich(res, extras=FULL)


@pytest.fixture()
def client(monkeypatch):
    import app as app_module
    monkeypatch.setattr(app_module, 'compute', functools.partial(A.compute, now=NOW))
    return app_module.app.test_client()


def items(text):
    """'a, b' -> ['a', 'b']; 'none' -> []."""
    return [] if text in (None, 'none') else str(text).split(', ')


def holding(rows, moment, start='start_local', end='end_local'):
    return [r for r in rows if r[start] is not None and r[end] is not None and r[start] <= moment < r[end]]


# ── dasa levels ───────────────────────────────────────────────────────────────

def test_dasa_periods_follow_from_the_moon(T):
    rows = T['AI_DasaPeriods']
    facts = {r['key']: r['value'] for r in T['AI_Facts']}
    birth = facts['birth_datetime_local']
    assert [r['dasa_lord'] for r in rows] == ['Jupiter', 'Saturn', 'Mercury', 'Ketu', 'Venus', 'Sun', 'Moon', 'Mars', 'Rahu']
    assert rows[0]['start_local'] == birth and rows[0]['age_at_start_years'] == 0
    assert all(a['end_local'] == b['start_local'] for a, b in zip(rows, rows[1:]))
    # by hand: the Moon at 323.766268 has crossed (323.766268 - 320) / 13.333333 = 0.282470 of Purva
    # Bhadrapada, so 0.717530 of Jupiter's 16 years is left; Saturn's 19 years follow.
    moon = by(T['AI_Planets'], 'point')['Moon']['longitude_deg']
    left = 1 - (moon % (360 / 27)) / (360 / 27)
    assert left == pytest.approx(0.717530, abs=1e-6)
    saturn = birth + timedelta(days=16 * left * 365.25)
    mercury = saturn + timedelta(days=19 * 365.25)
    close = timedelta(minutes=2)                        # the longitude is rounded to 6 places
    assert abs(rows[1]['start_local'] - saturn) < close and abs(rows[2]['start_local'] - mercury) < close
    assert rows[2]['start_local'].date() == date(2015, 4, 29) and rows[2]['end_local'].date() == date(2032, 4, 28)
    # each dasa is the first to the last bhukti of AI_Dasa
    for r in rows:
        mine = [b for b in T['AI_Dasa'] if b['dasa_lord'] == r['dasa_lord']]
        assert (r['start_local'], r['end_local']) == (mine[0]['start_local'], mine[-1]['end_local'])
    lord = by(rows, 'dasa_lord')
    assert lord['Saturn']['dasa_lord_dignity'] == 'Exalted' and lord['Saturn']['dasa_lord_house'] == 4
    assert lord['Mercury']['dasa_lord_star_lord'] == 'Jupiter'                  # Mercury is in Vishakha


def test_antarams_tile_every_bhukti_in_vimshottari_proportion(T):
    rows = T['AI_Antaram']
    bhuktis = T['AI_Dasa']
    groups = {}
    for r in rows:
        groups.setdefault((r['dasa_lord'], r['bhukti_lord']), []).append(r)
    assert len(groups) == len(bhuktis) == 80
    second = timedelta(seconds=1)
    for b in bhuktis:
        mine = groups[(b['dasa_lord'], b['bhukti_lord'])]
        assert mine[0]['start_local'] == b['start_local'] and mine[-1]['end_local'] == b['end_local']
        assert all(x['end_local'] == y['start_local'] for x, y in zip(mine, mine[1:]))
        first = ORDER.index(b['bhukti_lord'])
        expected = [ORDER[(first + i) % 9] for i in range(9)]
        assert [x['antaram_lord'] for x in mine] == expected[9 - len(mine):]       # clipped only at birth
        if len(mine) == 9 and b['start_local'] > bhuktis[0]['end_local']:
            span = b['end_local'] - b['start_local']
            for x in mine:
                assert abs((x['end_local'] - x['start_local']) - span * (YEARS[x['antaram_lord']] / 120)) <= 2 * second
        for x in mine:
            assert x['days'] == pytest.approx((x['end_local'] - x['start_local']).total_seconds() / 86400, abs=1e-3)
    # by hand, Mercury-Rahu (24-10-2024 to 14-05-2027): the antarams run Rahu 18, Jupiter 16, Saturn 19,
    # Mercury 17, Ketu 7, then Venus 20: Venus starts after 77/120 of the bhukti and ends after 97/120.
    b = next(x for x in bhuktis if (x['dasa_lord'], x['bhukti_lord']) == ('Mercury', 'Rahu'))
    venus = next(x for x in groups[('Mercury', 'Rahu')] if x['antaram_lord'] == 'Venus')
    span = b['end_local'] - b['start_local']
    assert abs(venus['start_local'] - (b['start_local'] + span * (77 / 120))) <= 2 * second
    assert abs(venus['end_local'] - (b['start_local'] + span * (97 / 120))) <= 2 * second
    assert venus['start_local'].date() == date(2026, 6, 14) and venus['end_local'].date() == date(2026, 11, 16)


def test_pratyantars_tile_the_antarams_inside_their_range(T):
    groups = {}
    for r in T['AI_Pratyantar']:
        groups.setdefault((r['dasa_lord'], r['bhukti_lord'], r['antaram_lord']), []).append(r)
    lo, hi = T['AI_Pratyantar'][0]['start_local'], T['AI_Pratyantar'][-1]['end_local']
    whole = 0
    for a in T['AI_Antaram']:
        mine = groups.get((a['dasa_lord'], a['bhukti_lord'], a['antaram_lord']))
        if not mine or a['start_local'] < lo or a['end_local'] > hi:
            continue
        whole += 1
        assert len(mine) == 9 and mine[0]['pratyantar_lord'] == a['antaram_lord']
        assert mine[0]['start_local'] == a['start_local'] and mine[-1]['end_local'] == a['end_local']
    assert whole > 80                                   # 11 years of antarams


# ── transits ──────────────────────────────────────────────────────────────────

def _reach(T, planet, sign_no):
    """Aspected houses and natal points touched, worked out from AI_Planets alone."""
    points = [r for r in T['AI_Planets'] if r['sign_no'] is not None]
    lagna = next(r['sign_no'] for r in points if r['point'] == 'Lagna')
    aspected = [(sign_no - 1 + k - 1) % 12 + 1 for k in ASPECT.get(planet, (7,))]
    return (sorted((s - lagna) % 12 + 1 for s in aspected),
            [r['point'] for r in points if r['sign_no'] == sign_no],
            [r['point'] for r in points if r['sign_no'] in aspected])


@pytest.mark.parametrize('sheet', ['AI_Transits', 'AI_TransitsFast'])
def test_transit_rows_say_what_the_planet_touches(T, sheet):
    for r in T[sheet]:
        houses, inside, aspected = _reach(T, r['planet'], r['sign_no'])
        assert [int(x) for x in items(r['houses_aspected_from_lagna'])] == houses, r
        assert items(r['natal_points_in_sign']) == inside and items(r['natal_points_aspected']) == aspected
    if sheet == 'AI_Transits':
        # by hand: Saturn in Mesha (house 10) aspects the 3rd, 7th and 10th signs from it: Mithuna, Thula and
        # Makara = houses 12, 4, 7. Thula holds the natal Sun and Saturn; the other two are empty.
        saturn = next(r for r in T[sheet] if r['planet'] == 'Saturn' and r['sign'] == 'Mesha' and r['entry_local'].year == 2027)
        assert saturn['houses_aspected_from_lagna'] == '4, 7, 12' and saturn['natal_points_aspected'] == 'Sun, Saturn'
        assert saturn['natal_points_in_sign'] == 'none'
        # Jupiter in Simha (house 2) aspects Dhanu, Kumbha and Mesha = houses 6, 8, 10: natal Mars and Jupiter, and the Moon
        jupiter = next(r for r in T[sheet] if r['planet'] == 'Jupiter' and r['sign'] == 'Simha' and r['entry_local'].year == 2026)
        assert jupiter['houses_aspected_from_lagna'] == '6, 8, 10' and jupiter['natal_points_aspected'] == 'Moon, Mars, Jupiter'
        # Rahu in Kumbha sits on the natal Moon
        rahu = next(r for r in T[sheet] if r['planet'] == 'Rahu' and r['sign'] == 'Kumbha' and r['entry_local'].year == 2025)
        assert rahu['natal_points_in_sign'] == 'Moon' and rahu['houses_aspected_from_lagna'] == '2'


GOOD_FROM_MOON = {'Sun': (3, 6, 10, 11), 'Mars': (3, 6, 11), 'Mercury': (2, 4, 6, 8, 10, 11),
                  'Venus': (1, 2, 3, 4, 5, 8, 9, 11, 12)}
FAST_RANGE = {'Mars': (2, 10), 'Sun': (1, 2), 'Mercury': (1, 2), 'Venus': (1, 2)}


def test_faster_transits_cover_their_range_and_agree_with_the_other_sheets(T):
    rows = T['AI_TransitsFast']
    assert [p for p in dict.fromkeys(r['planet'] for r in rows)] == ['Mars', 'Sun', 'Mercury', 'Venus']
    bav = {(r['planet'], r['sign']): r['bindus'] for r in T['AI_Ashtakavarga'] if r['row_type'] == 'sign'}
    now = by(T['AI_TransitNow'], 'planet')
    year = timedelta(days=365.25)
    for planet, (back, ahead) in FAST_RANGE.items():
        mine = [r for r in rows if r['planet'] == planet]
        assert all(a['exit_local'] == b['entry_local'] for a, b in zip(mine, mine[1:]))
        assert all(r['entry_local'] is not None and r['exit_local'] is not None for r in mine)
        assert mine[0]['entry_local'] <= NOW - back * year < mine[0]['exit_local']
        assert mine[-1]['entry_local'] < NOW + ahead * year <= mine[-1]['exit_local']
        current = holding(mine, NOW, 'entry_local', 'exit_local')
        assert len(current) == 1 and current[0]['sign'] == now[planet]['sign']       # AI_TransitNow is a separate calculation
        for r in mine:
            assert r['own_bav_bindus'] == bav[(planet, r['sign'])] and r['sav_bindus'] == bav[('Sarvashtakavarga', r['sign'])]
            assert r['result_from_moon'] == ('favourable' if r['house_from_moon'] in GOOD_FROM_MOON[planet] else 'unfavourable')
            assert r['entry_utc'] == r['entry_local'] - IST and r['exit_utc'] == r['exit_local'] - IST
    sun = [r for r in rows if r['planet'] == 'Sun']
    assert all(b['sign_no'] == a['sign_no'] % 12 + 1 and a['motion_at_entry'] == 'direct' for a, b in zip(sun, sun[1:]))
    assert all(28.5 < (r['exit_local'] - r['entry_local']).total_seconds() / 86400 < 32.5 for r in sun)     # a solar month
    # Mars goes round in 1.881 years: 12 years are 6.4 rounds = 77 signs, plus a few re-entries when retrograde
    assert len(sun) == 37 and 76 <= len([r for r in rows if r['planet'] == 'Mars']) <= 90


@pytest.mark.parametrize('body,pid', [('Mars', swe.MARS), ('Sun', swe.SUN), ('Mercury', swe.MERCURY), ('Venus', swe.VENUS)])
def test_faster_transits_equal_a_plain_one_day_search(T, body, pid):
    mine = [r for r in T['AI_TransitsFast'] if r['planet'] == body]
    lo, hi = mine[0]['entry_utc'] - timedelta(days=2), mine[-1]['exit_utc'] + timedelta(days=2)
    plain = _plain_entries(pid, A.julian_day(lo), A.julian_day(hi), swe.SIDM_LAHIRI)
    assert len(plain) == len(mine) + 1                  # the entry of every row and the exit of the last
    for (jd, _, entered), row in zip(plain, mine):
        assert SIGNS[entered] == row['sign']
        assert abs(E.jd_to_utc(jd) - row['entry_utc']) <= timedelta(seconds=61)       # rows are cut to the minute


@pytest.mark.parametrize('body,pid,tolerance', [('Mars', swe.MARS, 2), ('Sun', swe.SUN, 2),
                                                ('Mercury', swe.MERCURY, 60), ('Venus', swe.VENUS, 60)])
def test_faster_transit_search_equals_one_day_step_search_1950_to_2080(body, pid, tolerance):
    jd0, jd1 = swe.julday(1950, 1, 1, 0.0), swe.julday(2080, 1, 1, 0.0)
    fast = E.sign_entries(body, jd0, jd1, swe.SIDM_LAHIRI)
    slow = _plain_entries(pid, jd0, jd1, swe.SIDM_LAHIRI)
    assert len(fast) == len(slow) > 700
    for a, b in zip(fast, slow):
        assert a[1:] == b[1:] and abs(a[0] - b[0]) * 86400 <= tolerance


def _longitude(pid, jd):
    with A._SWE_LOCK:
        swe.set_sid_mode(swe.SIDM_LAHIRI)
        return swe.calc_ut(jd, pid, swe.FLG_MOSEPH | swe.FLG_SIDEREAL)[0][0]


def _extreme(pid, jd, sign, half=4.0):
    """The moment the longitude is greatest (sign +1) or smallest (-1) near jd: a golden-section search."""
    base = _longitude(pid, jd)
    f = lambda t: sign * ((_longitude(pid, t) - base + 180.0) % 360.0 - 180.0)
    lo, hi = jd - half, jd + half
    g = (math.sqrt(5) - 1) / 2
    a, b = hi - g * (hi - lo), lo + g * (hi - lo)
    fa, fb = f(a), f(b)
    while hi - lo > 1e-4:
        if fa < fb:
            lo, a, fa = a, b, fb
            b = lo + g * (hi - lo)
            fb = f(b)
        else:
            hi, b, fb = b, a, fa
            a = hi - g * (hi - lo)
            fa = f(a)
    return (lo + hi) / 2


def test_stations_are_the_turning_points_of_the_longitude(T):
    rows = T['AI_Stations']
    ids = {'Mars': swe.MARS, 'Jupiter': swe.JUPITER, 'Saturn': swe.SATURN}
    assert set(r['planet'] for r in rows) == set(ids) and {r['station'] for r in rows} == {'retrograde', 'direct'}
    year = timedelta(days=365.25)
    for planet, pid in ids.items():
        mine = [r for r in rows if r['planet'] == planet]
        assert all(a['station'] != b['station'] and a['moment_utc'] < b['moment_utc'] for a, b in zip(mine, mine[1:]))
        assert NOW - IST - year <= mine[0]['moment_utc'] and mine[-1]['moment_utc'] <= NOW - IST + 10 * year
        for r in mine:
            jd = A.julian_day(r['moment_utc'])
            turn = _extreme(pid, jd, 1 if r['station'] == 'retrograde' else -1)
            assert abs(turn - jd) * 1440 < 45, r       # the longitude is flat here: minutes, not seconds
            assert r['longitude_deg'] == pytest.approx(_longitude(pid, jd), abs=1e-4)
            assert r['sign'] == SIGNS[int(r['longitude_deg'] // 30)] and r['moment_local'] == r['moment_utc'] + IST
        # a plain count: days on which the daily change of longitude changes its sign
        lo = A.julian_day(NOW - IST - year)
        steps = [(_longitude(pid, lo + d + 1) - _longitude(pid, lo + d) + 180) % 360 - 180 for d in range(int(11 * 365.25))]
        assert sum(1 for a, b in zip(steps, steps[1:]) if (a < 0) != (b < 0)) == len(mine)
        gaps = [(b['moment_utc'] - a['moment_utc']).days for a, b in zip(mine, mine[1:]) if a['station'] == 'retrograde']
        low, high = {'Mars': (55, 85), 'Jupiter': (110, 125), 'Saturn': (130, 145)}[planet]
        assert all(low <= g <= high for g in gaps), (planet, gaps)       # how long each stays retrograde
    # published dates of the Mars retrograde of early 2027
    mars = [r for r in rows if r['planet'] == 'Mars' and r['moment_utc'].year == 2027]
    assert [(r['station'], r['moment_utc'].date()) for r in mars] == [('retrograde', date(2027, 1, 10)),
                                                                      ('direct', date(2027, 4, 1))]


# ── one row per period ────────────────────────────────────────────────────────

def _tied_topics(T, lord):
    """Topics for which a planet owns a key house, sits in one or is a significator, from the sheets alone."""
    p = by(T['AI_Planets'], 'point')[lord]
    out = []
    for m in T['AI_TopicMap']:
        houses = [int(x) for x in items(m['d1_houses'])]
        owns = set(int(x) for x in items(p['houses_owned'])) & set(houses)
        if owns or p['house_whole_sign'] in houses or lord in items(m['significators_for_this_chart']):
            out.append(m['topic'])
    return out


def test_period_facts_repeat_the_other_sheets(T):
    rows = T['AI_PeriodFacts']
    planets, houses = by(T['AI_Planets'], 'point'), by(T['AI_Houses'], 'house')
    varga = {(r['chart'], r['point']): r for r in T['AI_Vargas']}
    pair = {}
    for r in T['AI_Pairs']:
        pair[(r['planet_a'], r['planet_b'])] = (r['b_counted_from_a'], r['compound_a_to_b'])
        pair[(r['planet_b'], r['planet_a'])] = ((14 - r['b_counted_from_a'] - 1) % 12 + 1, r['compound_b_to_a'])
    found = [(r['yoga'], items(r['planets'])) for r in T['AI_Yogas'] if r['status'] == 'found']
    top = {}
    for r in T['AI_TopicScores']:
        if r['in_top_quarter'] == 'yes':
            top.setdefault((r['dasa_lord'], r['bhukti_lord'], r['start_local']), []).append(r['topic'])
    bhukti_rows = [r for r in rows if r['level'] == 'bhukti']
    antaram_rows = [r for r in rows if r['level'] == 'antaram']
    assert [(r['dasa_lord'], r['bhukti_lord'], r['start_local'], r['end_local']) for r in bhukti_rows] == [
        (r['dasa_lord'], r['bhukti_lord'], r['start_local'], r['end_local']) for r in T['AI_Dasa']]
    lo, hi = NOW - timedelta(days=365.25), NOW + timedelta(days=10 * 365.25)
    expected = [a for a in T['AI_Antaram'] if a['end_local'] > lo and a['start_local'] < hi]
    assert [(r['antaram_lord'], r['start_local'], r['end_local']) for r in antaram_rows] == [
        (a['antaram_lord'], a['start_local'], a['end_local']) for a in expected]
    parent = None
    for r in rows:
        if r['level'] == 'bhukti':
            parent = r
            assert r['antaram_lord'] is None and r['period_lord'] == r['bhukti_lord'] and r['place_from_bhukti_lord'] is None
        else:
            assert r['period_lord'] == r['antaram_lord'] and (r['dasa_lord'], r['bhukti_lord']) == (parent['dasa_lord'], parent['bhukti_lord'])
        lord, p = r['period_lord'], planets[r['period_lord']]
        for col, src in (('houses_owned', 'houses_owned'), ('house', 'house_whole_sign'), ('sign', 'sign'),
                         ('dignity', 'dignity'), ('debilitated', 'debilitated'), ('combust', 'combust'),
                         ('debilitation_cancelled_by', 'debilitation_cancelled_by'), ('retrograde', 'retrograde'),
                         ('functional_nature', 'functional_nature'), ('vimsopaka_shodasavarga', 'vimsopaka_shodasavarga'),
                         ('own_bav_bindus_in_sign_occupied', 'own_bav_bindus_in_sign_occupied'),
                         ('sign_lord', 'sign_lord'), ('star_lord', 'star_lord'), ('d9_sign', 'd9_sign')):
            assert r[col] == p[src], (lord, col)
        for key in ('D9', 'D10'):
            v = varga[(key, lord)]
            k = key.lower()
            assert (r[f'{k}_sign'], r[f'{k}_dignity'], r[f'{k}_debilitated']) == (v['sign'], v['dignity'], v['debilitated'])
        assert r['sign_lord_house'] == planets[p['sign_lord']]['house_whole_sign']
        star = planets[p['star_lord']]
        assert (r['star_lord_houses_owned'], r['star_lord_house']) == (star['houses_owned'], star['house_whole_sign'])
        if lord == r['dasa_lord']:
            assert (r['place_from_dasa_lord'], r['relation_to_dasa_lord'], r['relation_of_dasa_lord']) == (1, 'same planet', 'same planet')
        else:
            place, relation = pair[(r['dasa_lord'], lord)]
            assert r['place_from_dasa_lord'] == place and r['relation_of_dasa_lord'] == relation
            assert r['relation_to_dasa_lord'] == pair[(lord, r['dasa_lord'])][1]
        if r['level'] == 'antaram':
            assert r['place_from_bhukti_lord'] == (1 if lord == r['bhukti_lord'] else pair[(r['bhukti_lord'], lord)][0])
        assert items(r['yogas']) == [name for name, members in found if lord in members]
        assert items(r['lord_tied_topics']) == _tied_topics(T, lord)
        assert items(r['bhukti_top_quarter_topics']) == top.get((parent['dasa_lord'], parent['bhukti_lord'], parent['start_local']), [])
    assert planets['Moon']['houses_owned'] == '1' and houses[1]['lord'] == 'Moon'


def test_taras_of_chart_s_by_hand(T):
    # The janma nakshatra is Purva Bhadrapada (25), ruled by Jupiter. Counting the dasa lords from Jupiter:
    # Jupiter 1 Janma, Saturn 2 Sampat, Mercury 3 Vipat, Ketu 4 Kshema, Venus 5 Pratyari, Sun 6 Sadhaka,
    # Moon 7 Vadha, Mars 8 Mitra, Rahu 9 Parama Mitra.
    lordship = {'Jupiter': (1, 'Janma', 'mixed'), 'Saturn': (2, 'Sampat', 'favourable'),
                'Mercury': (3, 'Vipat', 'unfavourable'), 'Ketu': (4, 'Kshema', 'favourable'),
                'Venus': (5, 'Pratyari', 'unfavourable'), 'Sun': (6, 'Sadhaka', 'favourable'),
                'Moon': (7, 'Vadha', 'unfavourable'), 'Mars': (8, 'Mitra', 'favourable'),
                'Rahu': (9, 'Parama Mitra', 'favourable')}
    # The tara of the nakshatra occupied: (its number - 25) mod 9 + 1. Sun in Swati 15: 9. Moon 25: 1.
    # Mars in Uttara Ashadha 21: 6. Mercury in Vishakha 16: 1. Jupiter in Purva Ashadha 20: 5. Venus in
    # Jyeshtha 18: 3. Saturn in Vishakha 16: 1. Rahu in Krittika 3: 6. Ketu in Anuradha 17: 2.
    position = {'Sun': 9, 'Moon': 1, 'Mars': 6, 'Mercury': 1, 'Jupiter': 5, 'Venus': 3, 'Saturn': 1, 'Rahu': 6, 'Ketu': 2}
    seen = set()
    for r in T['AI_PeriodFacts']:
        lord = r['period_lord']
        seen.add(lord)
        assert (r['tara_of_lordship_no'], r['tara_of_lordship'], r['tara_of_lordship_label']) == lordship[lord]
        assert r['tara_of_position_no'] == position[lord]
        assert r['tara_of_position'] == E.TARA_NAMES[position[lord] - 1]
    assert seen == set(lordship)
    for r in T['AI_DasaPeriods']:
        assert (r['dasa_lord_tara_of_lordship_no'], r['dasa_lord_tara_of_lordship'],
                r['dasa_lord_tara_of_lordship_label']) == lordship[r['dasa_lord']]


# ── topic scores, promise, windows ────────────────────────────────────────────

def test_rank_and_top_quarter_by_plain_counting(T):
    rows = T['AI_TopicScores']
    for topic in {r['topic'] for r in rows}:
        mine = [r for r in rows if r['topic'] == topic]
        assert len(mine) == 80
        distinct = sorted({r['relevance_score_not_a_prediction'] for r in mine}, reverse=True)
        for r in mine:
            score = r['relevance_score_not_a_prediction']
            assert r['rank_in_topic'] == distinct.index(score) + 1
            higher = len([x for x in mine if x['relevance_score_not_a_prediction'] > score])
            assert r['in_top_quarter'] == ('yes' if higher < 20 else 'no')
        assert any(r['rank_in_topic'] == 1 and r['in_top_quarter'] == 'yes' for r in mine)
    # by hand: Mercury-Jupiter scores 6 for the job topic and nothing scores higher
    one = next(r for r in rows if (r['dasa_lord'], r['bhukti_lord'], r['topic']) == ('Mercury', 'Jupiter', 'Job and career'))
    assert one['relevance_score_not_a_prediction'] == 6 and one['rank_in_topic'] == 1 and one['in_top_quarter'] == 'yes'


def test_score_ranks_at_the_edge_of_the_quarter():
    # eight scores, so a quarter is two: 9 and 8 have fewer than two above them; 7 has exactly two and is out
    assert E.score_ranks([9, 8, 7, 7, 5, 5, 5, 1]) == [(1, True), (2, True), (3, False), (3, False), (4, False),
                                                      (4, False), (4, False), (5, False)]
    # ties at the top all count: three share the best score, the next has three above it
    assert E.score_ranks([6, 6, 6, 4, 4, 2, 2, 2]) == [(1, True)] * 3 + [(2, False)] * 2 + [(3, False)] * 3
    assert E.score_ranks([3.5]) == [(1, True)] and E.score_ranks([2, 2, 2, 2]) == [(1, True)] * 4


def test_topic_promise_of_chart_s_by_hand(T):
    rows = T['AI_TopicPromise']
    assert 'Strain (general)' not in {r['topic'] for r in rows}
    assert {r['topic'] for r in rows} == {m['topic'] for m in T['AI_TopicMap']} - {'Strain (general)'}
    for topic in {r['topic'] for r in rows}:
        mine = [r for r in rows if r['topic'] == topic]
        total = mine[-1]
        assert total['row_type'] == 'total' and [r['row_type'] for r in mine[:-1]] == ['factor'] * (len(mine) - 1)
        assert total['supports_count'] == len([r for r in mine if r['side'] == 'supports'])
        assert total['weakens_count'] == len([r for r in mine if r['side'] == 'weakens'])
        assert all(r['side'] in ('supports', 'weakens') and r['fact'] and r['factor'] and r['subject'] for r in mine[:-1])
        assert all(r['supports_count'] is None and r['weakens_count'] is None for r in mine[:-1])
    got = lambda topic: [(r['subject'], r['side'], r['fact']) for r in rows if r['topic'] == topic and r['row_type'] == 'factor']
    # Home and vehicle: house 4 is Thula, lord Venus in house 5 (a trikona; a friend's sign, so no dignity
    # row). In the house: the Sun (debilitated: weakens) and Saturn (exalted: supports). No benefic is in
    # Thula or aspects it. 33 bindus. Significators: Venus is the lord and is tested once; Mars is in
    # house 6 and the Moon in house 8, both in neutral or enemy signs.
    assert got('Home and vehicle') == [
        ('lord of house 4', 'supports', 'Venus is in house 5'),
        ('house 4', 'weakens', 'Sun is in house 4'),
        ('house 4', 'supports', 'Saturn is exalted in house 4'),
        ('house 4', 'supports', 'house 4 has 33 bindus'),
        ('significator Mars', 'weakens', 'Mars is in house 6'),
        ('significator Moon', 'weakens', 'Moon is in house 8'),
    ]
    assert got('Mother') == got('Home and vehicle')[:4] + [('significator Moon', 'weakens', 'Moon is in house 8')]
    # Father: house 9 is Meena, lord Jupiter in its own sign in house 6; nobody in Meena; of the benefics
    # none aspects it (Mars does, and aspects of malefics are not tested); 27 bindus is in between. The Sun
    # is debilitated but cancelled, in a neutral sign, in house 4: no row.
    assert got('Father') == [
        ('lord of house 9', 'weakens', 'Jupiter is in house 6'),
        ('lord of house 9', 'supports', 'Jupiter is in its own sign in Dhanu'),
    ]
    job = got('Job and career')
    assert ('yoga', 'supports', 'Raja yoga between Mars and Jupiter') in job
    assert ('significator Saturn', 'weakens', 'Saturn is combust') in job
    assert ('house 6', 'supports', 'Mars is in house 6') in job and ('house 11', 'supports', 'Rahu is in house 11') in job
    assert not any(s == 'significator Sun' for s, _, _ in job)               # debilitation cancelled, neutral sign


def _promise(lagna, topic, **signs):
    c = E.Ctx(chart(lagna, **signs))
    E._yogas(c)
    return [(r['subject'], r['factor'], r['side']) for r in E._topic_promise(c)
            if r['topic'] == topic and r['row_type'] == 'factor']


def test_topic_promise_rules_on_built_charts():
    # Mesha lagna: the 7th is Thula, lord Venus. Venus debilitated in Kanya (house 6). Mercury, who owns Kanya
    # and is exalted there, is in Mithuna: the 3rd from the lagna and from the Moon in Mesha, so no kendra and
    # no cancellation; in D9 Venus at 15 degrees of Kanya is in Rishabha. Two weakening rows for the lord of 7.
    rows = _promise('Mesha', 'Marriage', Venus='Kanya', Mercury='Mithuna', Sun='Mithuna', Moon='Mesha',
                    Mars='Mithuna', Jupiter='Mithuna', Saturn='Mithuna', Rahu='Kumbha', Ketu='Simha')
    assert ('lord of house 7, 2', 'lord in house 6, 8 or 12', 'weakens') in rows
    assert ('lord of house 7, 2', 'lord debilitated, not cancelled', 'weakens') in rows
    assert not any(subject == 'significator Venus' for subject, _, _ in rows)      # Venus is tested once, as lord
    # Mars in the 7th in an enemy's sign weakens; Jupiter in the 7th supports; the lord in its own sign in a kendra
    rows = _promise('Mesha', 'Marriage', Venus='Thula', Jupiter='Thula', Mars='Thula', Sun='Kumbha', Moon='Kumbha',
                    Mercury='Kumbha', Saturn='Kumbha', Rahu='Kanya', Ketu='Meena')
    assert ('lord of house 7, 2', 'lord in a kendra or trikona', 'supports') in rows
    assert ('lord of house 7, 2', 'lord exalted, in moolatrikona or in its own sign', 'supports') in rows
    assert ('house 7', 'Jupiter, Venus or Mercury in the house', 'supports') in rows
    assert ('house 7', 'Sun, Mars, Saturn, Rahu or Ketu in the house', 'weakens') in rows
    # a combust lord: Venus 3 degrees from the Sun
    rows = _promise('Mesha', 'Marriage', Venus=('Simha', 12.0), Sun=('Simha', 15.0))
    assert ('lord of house 7, 2', 'lord combust', 'weakens') in rows
    # Saturn in its own sign in the 11th supports twice over; Rahu in the 3rd supports
    rows = _promise('Mesha', 'Siblings', Saturn='Kumbha', Rahu='Mithuna', Ketu='Dhanu', Mars='Simha')
    assert ('house 11', 'Sun, Mars or Saturn in the house in its own or exaltation sign', 'supports') in rows
    assert ('house 3', 'Sun, Mars, Saturn, Rahu or Ketu in house 3, 6 or 11', 'supports') in rows


def _layers(T, topic, moment, cache={}):
    """The layers of AI_TopicWindows that hold for a topic at one moment, worked out from the sheets alone."""
    m = by(T['AI_TopicMap'], 'topic')[topic]
    houses = [int(x) for x in items(m['d1_houses'])]
    lagna = by(T['AI_Planets'], 'point')['Lagna']['sign_no']
    key_signs = {(lagna - 1 + h - 1) % 12 + 1 for h in houses}
    bav = {(r['planet'], r['sign_no']): r['bindus'] for r in T['AI_Ashtakavarga'] if r['row_type'] == 'sign'}
    out = []
    bhukti = holding(T['AI_Dasa'], moment)
    antaram = holding(T['AI_Antaram'], moment)
    if not bhukti or not antaram:
        return None
    score = next(r for r in T['AI_TopicScores'] if r['topic'] == topic and r['start_local'] == bhukti[0]['start_local']
                 and r['bhukti_lord'] == bhukti[0]['bhukti_lord'])
    if score['in_top_quarter'] == 'yes':
        out.append('bhukti')
    if topic in _tied_topics(T, antaram[0]['antaram_lord']):
        out.append('antaram')
    for planet in ('Jupiter', 'Saturn'):
        stay = holding([r for r in T['AI_Transits'] if r['planet'] == planet], moment, 'entry_local', 'exit_local')[0]
        reach = {stay['sign_no']} | {(stay['sign_no'] - 1 + k - 1) % 12 + 1 for k in ASPECT[planet]}
        if reach & key_signs and bav[(planet, stay['sign_no'])] >= 4:
            out.append(planet.lower())
    return out


def test_topic_windows_equal_a_day_by_day_scan_of_the_sheets(T):
    rows = T['AI_TopicWindows']
    assert all(r['label'] == 'rule-based overlap, not a prediction' for r in rows)
    assert all(2 <= r['layers_agreeing'] == len(items(r['layers'])) <= 4 for r in rows)
    assert all(NOW <= r['start_local'] < r['end_local'] <= NOW + timedelta(days=10 * 365.25, seconds=1) for r in rows)
    topics = [m['topic'] for m in T['AI_TopicMap']]
    assert [t for t in dict.fromkeys(r['topic'] for r in rows)] == [t for t in topics if any(r['topic'] == t for r in rows)]
    for topic in topics:
        mine = [r for r in rows if r['topic'] == topic]
        assert all(a['end_local'] <= b['start_local'] for a, b in zip(mine, mine[1:]))       # they never overlap
        # windows that touch were not joined only because something differs
        for a, b in zip(mine, mine[1:]):
            if a['end_local'] == b['start_local']:
                assert (a['layers'], a['bhukti_lord'], a['dasa_lord']) != (b['layers'], b['bhukti_lord'], b['dasa_lord'])
        day = NOW + timedelta(hours=19)                 # noon, every fifth day for ten years
        checked = 0
        while day < NOW + timedelta(days=10 * 365.25):
            expect = _layers(T, topic, day)
            got = holding(mine, day)
            if expect is not None and len(expect) >= 2:
                assert len(got) == 1 and items(got[0]['layers']) == expect, (topic, day, expect, got)
            else:
                assert not got, (topic, day, expect, got)
            checked += 1
            day += timedelta(days=5)
        assert checked > 700
    # by hand, the job topic (houses 10, 6, 2, 11): from 26-06-2027 Jupiter is in Simha (5 bindus) and aspects
    # houses 6 and 10; Saturn is in Mesha, house 10 (7 bindus), until 20-10-2027; Mercury-Jupiter is the
    # top-scoring bhukti; the antaram lords Jupiter (owns 6) and Saturn (significator) are both tied.
    one = next(r for r in rows if r['topic'] == 'Job and career' and r['layers_agreeing'] == 4)
    assert (one['start_local'], one['end_local']) == (datetime(2027, 6, 26, 5, 18), datetime(2027, 10, 20, 7, 12))
    assert (one['bhukti_lord'], one['antaram_lords'], one['jupiter_signs'], one['saturn_signs']) == (
        'Jupiter', 'Jupiter, Saturn', 'Simha', 'Mesha')
    assert (one['jupiter_own_bindus'], one['saturn_own_bindus'], one['key_houses_reached_by_both']) == ('5', '7', '10')
    assert one['jupiter_key_houses_reached'] == '2, 6, 10' and one['saturn_key_houses_reached'] == '10'
    double = [r for r in T['AI_DoubleTransit'] if r['house'] == 10 and r['start_local'] == one['start_local']]
    assert double and double[0]['end_local'] == one['end_local']               # the same span in AI_DoubleTransit


# ── what only the user knows ──────────────────────────────────────────────────

def test_nothing_entered_means_not_given_and_no_event_sheets(ai, T):
    assert 'AI_LifeEvents' not in ai['sheets'] and 'AI_EventCheck' not in ai['sheets']
    assert [r['key'] for r in T['AI_LifeFacts']] == [k for k, _ in E.LIFE_FIELDS]
    assert all(r['value'] == 'not given' and r['note'] for r in T['AI_LifeFacts'])
    facts = {r['key']: r['value'] for r in T['AI_Facts']}
    assert facts['birth_time_uncertainty_minutes'] is None and facts['coordinates_source'] == 'entered'
    assert all(r['lagna_uncertain_for_stated_accuracy'] == 'not given' for r in T['AI_VargaLagnas'])
    plain = E.enrich(A.compute(now=NOW, **S))['tables']                         # called without extras
    assert {r['key']: r['value'] for r in plain['AI_Facts']}['coordinates_source'] is None
    text = ' '.join(r['text'] for r in T['AI_ReadMe'])
    assert 'present only when the user entered past events' in text and 'not computed' in text


def test_what_the_user_entered_is_repeated_and_checked(full, T):
    t = full['tables']
    assert full['sheets'] == [s for s in E.SHEETS] and len(full['sheets']) == 30
    life = {r['key']: r['value'] for r in t['AI_LifeFacts']}
    assert life == {'marital_status': 'married', 'marriage_year': 2012, 'children_count': 1,
                    'first_child_birth_year': 2014, 'elder_siblings': 0, 'younger_siblings': 1,
                    'work_type': 'employed', 'lives_abroad': 'no'}
    assert [(r['event_type'], r['year'], r['topic'], r['age_that_year']) for r in t['AI_LifeEvents']] == [
        ('marriage', 2012, 'Marriage', 28), ('job change', 2023, 'Job and career', 39),
        ('home bought', 2019, 'Home and vehicle', 35)]
    facts = {r['key']: r['value'] for r in t['AI_Facts']}
    assert facts['birth_time_uncertainty_minutes'] == 5
    # with 5 minutes: uncertain where a margin of AI_VargaLagnas is under 300 seconds
    for r in t['AI_VargaLagnas']:
        margins = [x for x in (r['seconds_earlier'], r['seconds_later']) if x is not None]
        assert r['lagna_uncertain_for_stated_accuracy'] == ('yes' if min(margins) < 300 else 'no')
    stated = [r['chart'] for r in t['AI_VargaLagnas'] if r['lagna_uncertain_for_stated_accuracy'] == 'yes']
    assert stated == ['D4', 'D9', 'D10', 'D12', 'D16', 'D20', 'D24', 'D27', 'D40', 'D45', 'D60']      # D4, D12, D16 are new
    # nothing else changes
    for sheet in ('AI_Planets', 'AI_Dasa', 'AI_TopicScores', 'AI_TopicWindows', 'AI_PeriodFacts', 'AI_Transits'):
        assert t[sheet] == T[sheet]


def test_event_check_lists_the_bhuktis_of_that_year(full):
    t = full['tables']
    rows = t['AI_EventCheck']
    score = {(r['dasa_lord'], r['bhukti_lord'], r['start_local'], r['topic']): r for r in t['AI_TopicScores']}
    for e in t['AI_LifeEvents']:
        mine = [r for r in rows if (r['event_type'], r['year']) == (e['event_type'], e['year'])]
        y0, y1 = datetime(e['year'], 1, 1), datetime(e['year'] + 1, 1, 1)
        overlapping = [b for b in t['AI_Dasa'] if b['end_local'] > y0 and b['start_local'] < y1]
        assert [(r['dasa_lord'], r['bhukti_lord']) for r in mine] == [(b['dasa_lord'], b['bhukti_lord']) for b in overlapping]
        assert mine[0]['overlap_start_local'] == y0 and mine[-1]['overlap_end_local'] == y1
        assert all(a['overlap_end_local'] == b['overlap_start_local'] for a, b in zip(mine, mine[1:]))
        for r, b in zip(mine, overlapping):
            s = score[(b['dasa_lord'], b['bhukti_lord'], b['start_local'], e['topic'])]
            assert r['topic'] == e['topic']
            assert (r['relevance_score_not_a_prediction'], r['rank_in_topic'], r['in_top_quarter']) == (
                s['relevance_score_not_a_prediction'], s['rank_in_topic'], s['in_top_quarter'])
            inside = [a['antaram_lord'] for a in t['AI_Antaram'] if (a['dasa_lord'], a['bhukti_lord']) == (b['dasa_lord'], b['bhukti_lord'])
                      and a['end_local'] > r['overlap_start_local'] and a['start_local'] < r['overlap_end_local']]
            assert items(r['antaram_lords_tied_to_topic']) == list(dict.fromkeys(
                x for x in inside if e['topic'] in _tied_topics(t, x)))
    # 2012: Saturn-Rahu ran until 15-10-2012 (the date another program prints for this birth), then
    # Saturn-Jupiter. Jupiter entered sidereal Rishabha in May 2012; Saturn was in Thula and, retrograde, Kanya.
    first, second = [r for r in rows if r['event_type'] == 'marriage']
    assert (first['dasa_lord'], first['bhukti_lord'], first['overlap_end_local'].date()) == ('Saturn', 'Rahu', date(2012, 10, 15))
    assert (second['dasa_lord'], second['bhukti_lord']) == ('Saturn', 'Jupiter')
    assert first['jupiter_signs'] == 'Mesha (house 10), Rishabha (house 11)'
    assert first['saturn_signs'] == 'Thula (house 4), Kanya (house 3)'
    # marriage houses are 7, 2, 11: Jupiter in Rishabha occupies the 11th; Saturn from Thula aspects
    # Dhanu, Mesha and Kataka (houses 6, 10, 1) and from Kanya Vrischika, Meena and Mithuna (5, 9, 12): none
    assert first['jupiter_reaches_key_house'] == 'yes' and first['saturn_reaches_key_house'] == 'no'


@pytest.mark.parametrize('bad', [
    {'tob_uncertainty_min': 'soon'}, {'tob_uncertainty_min': -1}, {'tob_uncertainty_min': 121},
    {'life': 'married'}, {'life': {'marital_status': 'engaged'}}, {'life': {'children_count': 2.5}},
    {'life': {'children_count': 21}}, {'life': {'work_type': 'astronaut'}}, {'life': {'lives_abroad': 'maybe'}},
    {'life': {'events': 'marriage'}}, {'life': {'events': [{'type': 'marriage'}]}},
    {'life': {'events': [{'type': 'lottery', 'year': 2010}]}}, {'life': {'events': ['marriage 2012']}},
    {'life': {'events': [{'type': 'marriage', 'year': 2012}] * 11}},
])
def test_bad_optional_input_is_refused(bad):
    with pytest.raises(A.InputError):
        E.clean_extras(bad)


def test_years_must_lie_between_birth_and_report(res):
    for life in ({'marriage_year': 1983}, {'first_child_birth_year': 2027}, {'events': [{'type': 'job change', 'year': 2030}]}):
        with pytest.raises(A.InputError):
            E.enrich(res, extras={'life': life})
    assert E.enrich(res, extras={'life': {'marriage_year': 1984, 'events': [{'type': 'higher study', 'year': 2026}]}})
    clean = E.clean_extras({'life': {'lives_abroad': True, 'work_type': ' Business '}})
    assert clean['life']['lives_abroad'] == 'yes' and clean['life']['work_type'] == 'business'
    assert E.clean_extras(None) == E.clean_extras({}) and E.clean_extras({})['events'] == []


# ── small additions ───────────────────────────────────────────────────────────

def test_new_facts_of_chart_s_by_hand(T):
    facts = {r['key']: r['value'] for r in T['AI_Facts']}
    assert [r['key'] for r in T['AI_Facts']][-4:] == ['civil_weekday', 'coordinates_source',
                                                     'birth_time_uncertainty_minutes', 'dasa_shift_days_per_minute']
    # 4 November 1984 was a Sunday; the birth at 00:20 is before sunrise, so the Vedic day is still Saturday
    assert (facts['civil_weekday'], facts['weekday']) == ('Sunday', 'Saturday')
    assert date(1984, 11, 4).isoweekday() == 7
    # The Moon moves 11.868969 degrees a day = 0.00824234 a minute = 0.000618176 of a nakshatra; of Jupiter's
    # 16 years that is 16 x 365.25 x 0.000618176 = 3.612617 days less to run, while the start itself moves
    # 60 seconds = 0.000694 days later: -3.612617 + 0.000694 = -3.611923.
    speed = by(T['AI_Planets'], 'point')['Moon']['speed_deg_per_day']
    assert facts['dasa_shift_days_per_minute'] == pytest.approx(60 / 86400 - 16 * 365.25 * speed / 1440 / (360 / 27), abs=2e-3)
    assert facts['dasa_shift_days_per_minute'] == -3.612


def test_new_planet_and_varga_columns(T):
    p = by(T['AI_Planets'], 'point')
    # Mercury is 14.439705 degrees from the Sun against an orb of 14: just outside. Saturn 6.637156 against 15.
    assert p['Mercury']['combust_margin_deg'] == pytest.approx(0.439705, abs=1e-6) and p['Mercury']['combust'] == 'no'
    assert p['Saturn']['combust_margin_deg'] == pytest.approx(-8.362844, abs=1e-6) and p['Saturn']['combust'] == 'yes'
    for r in T['AI_Planets']:
        if r['combust_orb_deg'] is None:
            assert r['combust_margin_deg'] is None
        else:
            assert r['combust_margin_deg'] == pytest.approx(r['distance_from_sun_deg'] - r['combust_orb_deg'], abs=1e-6)
            assert (r['combust_margin_deg'] < 0) == (r['combust'] == 'yes')
        assert r['bhava_differs_from_whole_sign'] == ('yes' if r['bhava_sripati'] != r['house_whole_sign'] else 'no')
    assert [r['point'] for r in T['AI_Planets'] if r['bhava_differs_from_whole_sign'] == 'yes'] == ['Mercury', 'Rahu', 'Ketu']
    lagnas = by(T['AI_VargaLagnas'], 'chart')
    number = {name: i + 1 for i, name in enumerate(SIGNS)}
    for r in T['AI_Vargas']:
        for col, src in (('house_if_lagna_earlier', 'lagna_if_earlier'), ('house_if_lagna_later', 'lagna_if_later')):
            other = lagnas[r['chart']][src]
            if r['point'] == 'Lagna' or other is None:
                assert r[col] is None
            else:
                assert r[col] == (r['sign_no'] - number[other]) % 12 + 1
    # by hand: the D9 lagna is Kumbha and becomes Makara 77 seconds earlier; the Sun is in Meena in D9:
    # 2nd from Kumbha, 3rd from Makara, and 1st from Meena, the lagna 772 seconds later.
    sun = next(r for r in T['AI_Vargas'] if (r['chart'], r['point']) == ('D9', 'Sun'))
    assert (lagnas['D9']['lagna_if_earlier'], lagnas['D9']['lagna_if_later']) == ('Makara', 'Meena')
    assert (sun['house'], sun['house_if_lagna_earlier'], sun['house_if_lagna_later']) == (2, 3, 1)


def test_weakening_factors_and_dharma_karmadhipati(T):
    y = by(T['AI_Yogas'], 'yoga')
    assert T['AI_Yogas'][-1]['yoga'] == 'Dharma-Karmadhipati yoga'               # added after the old rows
    # the 9th lord Jupiter and the 10th lord Mars are together in Dhanu, house 6
    dk = y['Dharma-Karmadhipati yoga']
    assert (dk['status'], dk['planets'], dk['houses']) == ('found', 'Jupiter, Mars', '6') and 'conjunction' in dk['note']
    assert dk['weakening_factors'] == 'Jupiter in house 6, Mars in house 6'
    assert y['Sasa yoga']['weakening_factors'] == 'Saturn combust'               # exalted, 6.6 degrees from the Sun
    assert y['Raja yoga']['weakening_factors'] == 'Mars in house 6, Jupiter in house 6'
    assert y['Harsha yoga (Viparita raja yoga)']['weakening_factors'] == 'none'   # house 6 is what that yoga needs
    assert y['Vesi yoga']['weakening_factors'] == 'none' and y['Amala yoga']['weakening_factors'] == 'none'
    for r in T['AI_Yogas']:
        if r['status'] != 'found' or 'dosha' in r['yoga'] or 'Neecha bhanga' in r['yoga']:
            assert r['weakening_factors'] is None, r['yoga']
    # built charts. Rishabha lagna: Saturn owns the 9th and the 10th.
    one = {r['yoga']: r for r in E._yogas(E.Ctx(chart('Rishabha')))}['Dharma-Karmadhipati yoga']
    assert one['status'] == 'found' and one['planets'] == 'Saturn' and 'owns both' in one['note']
    # Mesha lagna: Jupiter (9) in Kataka and Saturn (10) in Thula neither share a sign nor aspect each other both ways
    c = E.Ctx(chart('Mesha', Jupiter='Kataka', Saturn='Thula', Venus='Mithuna', Sun='Mithuna', Mercury='Mithuna'))
    rows = {r['yoga']: r for r in E._yogas(c)}
    assert rows['Dharma-Karmadhipati yoga']['status'] == 'not found'
    assert rows['Dharma-Karmadhipati yoga']['weakening_factors'] is None
    # Jupiter exalted in a kendra: Hamsa, and nothing reduces it. With the Sun 5 degrees away it is combust.
    assert rows['Hamsa yoga']['status'] == 'found' and rows['Hamsa yoga']['weakening_factors'] == 'none'
    c = E.Ctx(chart('Mesha', Jupiter=('Kataka', 15.0), Sun=('Kataka', 20.0)))
    assert {r['yoga']: r for r in E._yogas(c)}['Hamsa yoga']['weakening_factors'] == 'Jupiter combust'
    # a debilitated planet in a yoga: Sasa needs Saturn strong, so use Malavya... Venus in Thula with a
    # debilitated Sun beside it does not weaken Malavya (the Sun is no part of it); the Raja yoga of the
    # Sun (lord of 5) and Venus (lord of 7) does carry it.
    c = E.Ctx(chart('Mesha', Venus=('Thula', 25.0), Sun=('Thula', 5.0), Saturn='Mithuna', Moon='Mithuna',
                    Mars='Mithuna', Mercury='Mithuna', Jupiter='Mithuna'))
    rows = E._yogas(c)
    malavya = next(r for r in rows if r['yoga'] == 'Malavya yoga')
    raja = next(r for r in rows if r['yoga'] == 'Raja yoga' and r['planets'] == 'Sun, Venus')
    assert malavya['status'] == 'found' and malavya['weakening_factors'] == 'none'
    assert raja['weakening_factors'].startswith('Sun debilitated')


def test_readme_names_the_new_sheets_and_warnings(ai, T):
    lines = T['AI_ReadMe']
    text = ' '.join(r['text'] for r in lines)
    assert 'AI_Brief is a short copy' in text and 'AI_PeriodFacts' in text and 'AI_TopicWindows' in text
    assert 'every point of D2 falls in Kataka or Simha' in text
    assert 'not that an event will happen' in text
    assert {r['item'] for r in lines if r['section'] == 'Sheets'} == set(ai['sheets'])
    about = {r['item']: r['text'] for r in lines if r['section'] == 'Sheets'}
    for sheet in ai['sheets']:
        assert about[sheet].endswith(f"({len(T[sheet]) if sheet != 'AI_ReadMe' else 0} rows)") or sheet == 'AI_ReadMe'
    missing = ' '.join(r['text'] for r in lines if r['section'] == 'Not in this file')
    assert 'Transits of the Moon' in missing and 'Mercury and Venus' in missing and 'Tara' not in missing
    assert set(E.SHEET_ABOUT) == set(E.SHEETS) == set(E.COLUMNS)


@pytest.mark.parametrize('extras', [EXTRAS, FULL])
def test_nothing_about_length_of_life_in_the_new_sheets(res, extras):
    block = E.enrich(res, extras=extras)
    banned = ('marak', 'ayur', 'longevity', 'lifespan', 'life span', 'death', 'mrityu', 'alpayu', 'alive', 'died')
    allowed = 'Length of life: left out on purpose; this file gives no figure and no analysis of it'
    for sheet, rows in block['tables'].items():
        for r in rows:
            for v in list(r.values()) + list(r):
                if isinstance(v, str) and v != allowed:
                    assert not any(b in v.lower() for b in banned), (sheet, v)
    for rule in block['rules']:
        assert not any(b in (rule['rule'] + ' ' + rule['variant']).lower() for b in banned)


# ── the brief ─────────────────────────────────────────────────────────────────

def _cell(value):
    """A cell as the brief writes it, written here without enrich.py."""
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _pairs(text):
    return [tuple(part.split('=', 1)) for part in text.split('; ')]


@pytest.mark.parametrize('which', ['plain', 'full'])
def test_every_line_of_the_brief_is_a_copy_of_its_source(ai, full, which):
    block = ai if which == 'plain' else full
    t = block['tables']
    rows = t['AI_Brief']
    assert [s for s in dict.fromkeys(r['section'] for r in rows)] == [
        s for s in ('About', 'Birth', 'Warnings', 'Life facts (entered by the user)', 'Planets', 'Houses', 'Yogas',
                    'As of the report date', 'Topics') if which == 'full' or not s.startswith('Life')]
    facts = {r['key']: r['value'] for r in t['AI_Facts']}
    checked = 0
    for r in rows:
        assert r['source'] in block['sheets'] and r['value'] is not None
        if r['section'] == 'Birth':
            assert r['value'] == facts[r['key']]
        elif r['source'] == 'AI_LifeFacts':
            assert r['value'] == by(t['AI_LifeFacts'], 'key')[r['key']]['value'] != 'not given'
        elif r['source'] == 'AI_ReadMe' or r['key'] in ('as_of',) or r['section'] == 'Warnings' or r['value'] == 'none':
            continue
        else:
            pairs = _pairs(r['value'])
            assert pairs and all(len(p) == 2 and p[1] != '' for p in pairs), r
            # one row of the source sheet holds every cell named on the line
            matches = [src for src in t[r['source']] if all(col in src and src[col] is not None and _cell(src[col]) == val
                                                           for col, val in pairs)]
            assert matches, r
            checked += 1
    assert checked > 100
    assert len(rows) <= 250 and len(json.dumps(A.to_jsonable(rows))) < 40000
    warn = {r['key']: r['value'] for r in rows if r['section'] == 'Warnings'}
    assert warn['charts with lagna_uncertain = yes'] == 'D9, D10, D20, D24, D27, D40, D45, D60'
    if which == 'full':
        assert warn['charts with lagna_uncertain_for_stated_accuracy = yes'].startswith('D4, D9, D10, D12, D16')
        life = [r for r in rows if r['section'].startswith('Life')]
        assert len(life) == 8 + 3 and life[-1]['value'] == 'event_type=home bought; year=2019; topic=Home and vehicle'


def test_brief_of_chart_s_says_what_runs_on_the_report_date(T):
    rows = T['AI_Brief']
    line = {r['key']: r['value'] for r in rows}
    assert line['as_of'] == NOW
    # the same periods and dates the reading of this chart was built on
    assert line['running dasa'] == 'dasa_lord=Mercury; start_local=2015-04-29 00:12:31; end_local=2032-04-28 06:12:31'
    assert line['running bhukti'].startswith('dasa_lord=Mercury; bhukti_lord=Rahu; start_local=2024-10-24')
    assert line['running antaram'].startswith('dasa_lord=Mercury; bhukti_lord=Rahu; antaram_lord=Venus; start_local=2026-06-14')
    assert 'pratyantar_lord=Saturn; start_local=2026-09-22' in line['running pratyantar']
    assert line['next bhukti 1'].startswith('dasa_lord=Mercury; bhukti_lord=Jupiter; start_local=2027-05-14')
    assert line['transit Saturn'].startswith('sign=Meena; entry_local=2025-03-29 21:44:00; exit_local=2027-06-03 05:28:00; '
                                             'house_from_lagna=9; house_from_moon=2; own_bav_bindus=1')
    assert 'saturn_from_moon=sade sati - 3rd phase' in line['transit Saturn']
    assert line['next sign Jupiter'].startswith('sign=Simha; entry_local=2026-10-31 12:02:00')
    assert line['transit Rahu'].startswith('sign=Kumbha') and 'natal_points_in_sign=Moon' in line['transit Rahu']
    assert line['sade sati span holding the report date'].startswith('cycle=2; start_local=2020-01-24 09:56:00; end_local=2028-02-23 19:24:00')
    assert line['Sun'].startswith('sign=Thula;') and 'debilitated=yes; debilitation_cancelled_by=a, b' in line['Sun']
    assert line['house 10'] == ('sign=Mesha; lord=Mars; lord_house=6; lord_dignity=Neutral; lord_debilitated=no; '
                                'occupants_whole_sign=none; aspected_by=Sun, Jupiter, Saturn; sav_bindus=34')
    assert line['Sasa yoga'] == 'status=found; planets=Saturn; houses=4; weakening_factors=Saturn combust'
    assert line['Job and career: where to look'] == 'd1_houses=10, 6, 2, 11; significators_for_this_chart=Sun, Saturn, Mercury; divisional_charts=D10'
    assert 'Strain (general): promise' not in line and line['Home and vehicle: promise'] == 'supports_count=3; weakens_count=3'
    assert line['Job and career: running bhukti'].startswith('bhukti_lord=Rahu; relevance_score_not_a_prediction=5; rank_in_topic=3')
    stale = next(r['value'] for r in rows if r['key'] == 'warning')
    assert 'true at the report date only' in stale


def test_brief_windows_are_the_first_three_with_three_layers(T):
    for m in T['AI_TopicMap']:
        topic = m['topic']
        mine = [r for r in T['AI_TopicWindows'] if r['topic'] == topic]
        strong = [r for r in mine if r['layers_agreeing'] >= 3][:3]
        weak = [r for r in mine if r['layers_agreeing'] < 3][:3 - len(strong)]
        expect = sorted(strong + weak, key=lambda r: r['start_local'])
        got = [r['value'] for r in T['AI_Brief'] if r['key'].startswith(f'{topic}: window ')]
        assert len(got) == len(expect) <= 3
        for text, src in zip(got, expect):
            pairs = dict(_pairs(text))
            assert pairs['start_local'] == _cell(src['start_local']) and pairs['layers'] == src['layers']


def test_brief_survives_a_date_outside_every_table(res):
    """A report date long after the tables end: the as-of lines that cannot be given are left out."""
    later = E.enrich(res, now=datetime(2110, 1, 1, 12, 0))
    keys = [r['key'] for r in later['tables']['AI_Brief']]
    assert 'as_of' in keys and 'running dasa' not in keys and 'running bhukti' not in keys
    assert not any(k.endswith('running bhukti') for k in keys) and later['tables']['AI_TopicWindows'] == []


# ── other charts ──────────────────────────────────────────────────────────────

@pytest.mark.parametrize('b', [
    CHART_A, CHART_B,
    dict(name='Newborn', dob='2026-09-27', tob='06:10', pob='Chennai', lat=13.0827, lon=80.2707),
    dict(name='Elder', dob='1931-02-14', tob='18:45', pob='Madurai', lat=9.9252, lon=78.1198),
    dict(name='South', dob='1975-06-30', tob='23:58', pob='Sydney', lat=-33.8688, lon=151.2093, tz='Australia/Sydney'),
])
def test_other_charts_build_and_keep_the_invariants(b):
    block = E.enrich(A.compute(now=NOW, **b), extras=EXTRAS)
    t = block['tables']
    assert block['sheets'] == [s for s in E.SHEETS if s not in E.OPTIONAL_SHEETS]
    for sheet in block['sheets']:
        cols = [c['name'] for c in block['columns'][sheet]]
        assert t[sheet] and all(list(r) == cols for r in t[sheet]), sheet
    assert all(a['end_local'] == b2['start_local'] for a, b2 in zip(t['AI_Antaram'], t['AI_Antaram'][1:]))
    assert t['AI_Antaram'][0]['start_local'] == t['AI_Dasa'][0]['start_local']
    assert t['AI_Antaram'][-1]['end_local'] == t['AI_Dasa'][-1]['end_local']
    for topic in {r['topic'] for r in t['AI_TopicWindows']}:
        mine = [r for r in t['AI_TopicWindows'] if r['topic'] == topic]
        assert all(x['end_local'] <= y['start_local'] for x, y in zip(mine, mine[1:]))
        for r in mine[::7]:                             # a sample of windows against the plain scan
            assert _layers(t, topic, r['start_local'] + (r['end_local'] - r['start_local']) / 2) == items(r['layers'])
    for r in t['AI_Transits'] + t['AI_TransitsFast']:
        houses, inside, aspected = _reach(t, r['planet'], r['sign_no'])
        assert [int(x) for x in items(r['houses_aspected_from_lagna'])] == houses and items(r['natal_points_in_sign']) == inside
    totals = [r for r in t['AI_TopicPromise'] if r['row_type'] == 'total']
    assert len(totals) == 13 and all(r['supports_count'] + r['weakens_count'] > 0 for r in totals)
    assert len(json.dumps(A.to_jsonable(t['AI_Brief']))) < 40000
    json.dumps(A.to_jsonable(block))


def test_place_without_sunrise_still_builds_every_new_sheet():
    r = A.compute('X', '2001-06-21', '12:00', 'Tromso', lat=69.65, lon=18.96, tz='Europe/Oslo', now=NOW)
    t = E.enrich(r, extras=FULL)['tables']
    assert all('Mandi' not in items(x['natal_points_in_sign']) for x in t['AI_Transits'] + t['AI_TransitsFast'])
    assert 'Mandi' not in [x['key'] for x in t['AI_Brief'] if x['section'] == 'Planets']
    assert t['AI_PeriodFacts'] and t['AI_TopicWindows'] and t['AI_EventCheck'] and t['AI_Stations']


# ── API and workbook ──────────────────────────────────────────────────────────

def test_api_brief_returns_meta_and_the_brief_only(client, ai):
    for value in ('brief', 'Brief', ' BRIEF '):
        rv = client.post('/api/horoscope', json=dict(G.API_BODY, ai=value))
        assert rv.status_code == 200
        out = rv.get_json()
        assert set(out) == {'meta', 'ai'}
        assert set(out['ai']) == {'schema', 'report_datetime_local', 'utc_offset_hours', 'columns', 'brief'}
        assert out['ai']['schema'] == 'horoscopegen-ai/2' and out['ai']['report_datetime_local'] == '2026-10-07T17:00:00'
        assert [c['name'] for c in out['ai']['columns']] == ['section', 'key', 'value', 'source']
        assert out['ai']['brief'] == json.loads(json.dumps(A.to_jsonable(ai['tables']['AI_Brief'])))
        assert len(rv.data) < 45000                     # small enough to paste into a chat
    plain = client.post('/api/horoscope', json=G.API_BODY).get_json()
    assert out['meta'] == plain['meta']
    rv = client.post('/api/download/json', json=dict(G.API_BODY, ai='brief'))
    assert rv.status_code == 200 and 'HoroscopeGen_Chart_S.json' in rv.headers['Content-Disposition']
    data = json.loads(rv.data.decode('utf-8'))
    assert set(data) == {'meta', 'ai'} and data['ai']['brief'] == out['ai']['brief']
    assert int(rv.headers['Content-Length']) == len(rv.data) < 45000


def test_api_takes_the_optional_inputs_and_refuses_bad_ones(client, full):
    body = dict(G.API_BODY, ai=True, tobUncertaintyMin=5, life=LIFE)
    rv = client.post('/api/horoscope', json=body)
    assert rv.status_code == 200
    block = rv.get_json()['ai']
    assert block == json.loads(json.dumps(A.to_jsonable(full)))
    assert block['sheets'][-1] == 'AI_EventCheck' and 'AI_LifeEvents' in block['sheets']
    # without latitude and longitude the place is looked up (the built-in list when the network is not there)
    looked = dict(G.API_BODY, ai='brief')
    del looked['lat'], looked['lon']
    rv = client.post('/api/horoscope', json=looked)
    if rv.status_code == 200:
        assert not any(r['key'] == 'coordinates_source' for r in rv.get_json()['ai']['brief'])
        full_block = client.post('/api/horoscope', json=dict(looked, ai=True)).get_json()['ai']
        facts = {r['key']: r['value'] for r in full_block['tables']['AI_Facts']}
        assert facts['coordinates_source'] == 'looked up'
    for bad in (dict(tobUncertaintyMin='x'), dict(tobUncertaintyMin=500), dict(life={'marital_status': 'engaged'}),
                dict(life={'events': [{'type': 'marriage', 'year': 1900}]}), dict(life=[1, 2])):
        for route, extra in (('/api/horoscope', {'ai': True}), ('/api/horoscope', {'ai': 'brief'}),
                             ('/api/download/json', {}), ('/api/download/excel', {})):
            rv = client.post(route, json=dict(G.API_BODY, **extra, **bad))
            assert rv.status_code == 400 and rv.get_json()['error'], (route, bad)
    # the optional inputs do nothing to a request that does not ask for the AI block
    rv = client.post('/api/horoscope', json=dict(G.API_BODY, tobUncertaintyMin=5, life=LIFE))
    assert rv.data == client.post('/api/horoscope', json=G.API_BODY).data


def test_workbook_carries_every_new_sheet_as_a_flat_table(client, full):
    import openpyxl
    rv = client.post('/api/download/excel', json=dict(G.API_BODY, tobUncertaintyMin=5, life=LIFE))
    assert rv.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(rv.data))
    assert wb.sheetnames[7:] == full['sheets'] and len(wb.sheetnames) == 37
    for sheet in ('AI_Brief', 'AI_LifeFacts', 'AI_LifeEvents', 'AI_DasaPeriods', 'AI_Antaram', 'AI_PeriodFacts',
                  'AI_TransitsFast', 'AI_Stations', 'AI_TopicPromise', 'AI_TopicWindows', 'AI_EventCheck'):
        ws = wb[sheet]
        cols = [c['name'] for c in full['columns'][sheet]]
        rows = full['tables'][sheet]
        assert [c.value for c in ws[1]] == cols and ws.max_row == len(rows) + 1 and not ws.merged_cells.ranges
        for got, src in zip(ws.iter_rows(min_row=2, values_only=True), rows):
            for a, b in zip(got, src.values()):
                if isinstance(b, datetime):
                    assert abs(a - b) < timedelta(seconds=1)
                elif isinstance(b, date):
                    assert a.date() == b                # a date cell comes back as midnight of that day
                elif isinstance(b, float):
                    assert a == pytest.approx(b, abs=1e-9)
                else:
                    assert a == b
    pd = pytest.importorskip('pandas')
    frames = pd.read_excel(io.BytesIO(rv.data), sheet_name=None)
    for sheet in full['sheets']:
        df = frames[sheet]
        assert len(df) == len(full['tables'][sheet]) and not df.isna().all(axis=0).any(), sheet
    assert str(frames['AI_TopicWindows']['start_local'].dtype).startswith('datetime64')
    assert str(frames['AI_Stations']['moment_utc'].dtype).startswith('datetime64')
